"""Gera o conjunto "texto_dificil": documentos em texto corrido com os dados de cada pessoa espalhados em seções.

Para que serve: a prova de generalização (EXP-017) mostrou que o Leitor erra no texto corrido difícil, em que cada
pessoa aparece em várias partes do documento (a identificação numa seção, o contrato em outra, o endereço em outra).
Este conjunto é NOVO, escrito para medir a correção da leitura, e é congelado ANTES de ela existir (ADR-58).
Ninguém que corrige abre este conjunto: ele só serve para medir.

Analogia: é uma prova com questões novas do mesmo tipo da que o aluno errou. Se ele aprendeu a regra, acerta; se
decorou as questões antigas, não.

As 6 estruturas (4 documentos de cada, de 3 a 5 pessoas por documento):
- nome_curto: identificação com o nome inteiro; contrato e endereço com o nome curto (primeiro nome + último sobrenome);
- apelido: a identificação apresenta o apelido; contrato e endereço usam só o apelido;
- ordem_trocada: nome inteiro em todas as seções, cada seção numa ordem diferente;
- referencia_por_cpf: contrato e endereço citam a pessoa só pelo CPF (escrito de outro jeito);
- cpf_repetido: um parágrafo por pessoa e, no fim, uma seção de documentos conferidos que repete o CPF;
- dependentes: como nome_curto, com filhos que têm os mesmos sobrenomes do funcionário (não são funcionários).

Nada vem dos arquivos do QA nem do lote guardado: as pessoas saem do mesmo sorteio fictício do gerador dos documentos
de teste (scripts/gerar_documentos_de_teste.py), com uma semente própria.

A reserva (ADR-140): um 2º conjunto do mesmo tipo, com outra semente, guardado FORA do repositório. A semente é
sorteada com 64 bits (--semente-aleatoria) e nunca fica no Git: o gerador é público e o manifesto traz as impressões
digitais, então uma semente curta poderia ser achada testando uma por uma, e a reserva deixaria de ser "não vista".

Uso:
  python scripts/gerar_texto_corrido_dificil.py   (o conjunto congelado, em data/avaliacao/documentos_texto_dificil/)
  python scripts/gerar_texto_corrido_dificil.py --semente-aleatoria --pasta <fora do repositório> \\
      --conjunto texto_dificil_reserva --manifesto data/avaliacao/reserva_d39_manifesto.json   (a reserva)
"""
import argparse
import json
import random
import secrets
import sys
from datetime import date
from pathlib import Path

from docx import Document

# Pasta raiz do projeto, para importar os módulos ao rodar o script diretamente
RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from eval.congelamento import impressao_digital  # noqa: E402
from scripts.gerar_documentos_de_teste import (NOMES_FEMININOS, NOMES_MASCULINOS, cpf_variado,  # noqa: E402
                                               data_curta, data_variada, gabarito_da_pessoa, pasta_do_conjunto,
                                               salario_variado, sortear_data, sortear_pessoa)

# O nome do conjunto (a pasta é data/avaliacao/documentos_texto_dificil)
CONJUNTO = "texto_dificil"
# A semente própria deste conjunto (os mesmos documentos a cada geração)
SEMENTE = 20260929
# O nível no gabarito (depois dos N1 a N5 do conjunto de desenvolvimento)
NIVEL = "N6"
# Quantos documentos de cada estrutura
DOCUMENTOS_POR_ESTRUTURA = 4
# Quantas pessoas cada documento tem (sorteado entre os dois)
MENOS_PESSOAS = 3
MAIS_PESSOAS = 5
# Os campos que o gabarito espera de cada pessoa
CAMPOS = ["nome_completo", "cpf", "data_nascimento", "data_admissao", "cargo", "valor_renda",
          "logradouro_residencial", "numero_residencial", "bairro_residencial", "municipio_residencial",
          "uf_residencial", "cep_residencial"]
# Os apelidos possíveis (o texto sempre apresenta o apelido antes de usá-lo)
APELIDOS = ["Nina", "Tuca", "Bebel", "Dudu", "Lelê", "Zeca", "Cacá", "Juju", "Nando", "Tati", "Bito", "Mimi",
            "Guga", "Lili", "Pipo", "Rafa"]
# Títulos das seções, sorteados para os documentos não ficarem todos iguais
TITULOS_DA_IDENTIFICACAO = ["1. Identificação", "DADOS PESSOAIS", "Quem são os novos colaboradores",
                            "Parte A — documentos"]
TITULOS_DO_CONTRATO = ["2. Contratação", "CARGOS E SALÁRIOS", "Vaga, início e remuneração", "Parte B — contrato"]
TITULOS_DO_ENDERECO = ["3. Endereços", "ONDE MORAM", "Endereço residencial", "Parte C — endereço"]
ABERTURAS = ["Segue o material das admissões deste mês, organizado por assunto.",
             "Bom dia! Juntei aqui o que o RH mandou, separado em partes.",
             "Relação das contratações recentes (informações reunidas de várias planilhas).",
             "Conforme combinado, os dados vão divididos por tema para facilitar a conferência."]


def endereco_em_texto(pessoa: dict) -> str:
    """O endereço completo numa frase. Ex.: "Rua das Acácias, 120, Vila Mariana, Campinas/SP, CEP 13045-120"."""
    return (f"{pessoa['rua']}, {pessoa['numero_da_casa']}, {pessoa['bairro']}, {pessoa['cidade']}/{pessoa['uf']}, "
            f"CEP {pessoa['cep']}")


def nome_curto(pessoa: dict) -> str:
    """O primeiro nome e o último sobrenome. Ex.: "Helena Duarte Monteiro" → "Helena Monteiro"."""
    partes = pessoa["nome"].split()
    return f"{partes[0]} {partes[-1]}"


def sortear_pessoas(sorteio: random.Random, quantas: int, numero_inicial: int) -> list[dict]:
    """Pessoas com o primeiro nome e o nome curto diferentes entre si (a referência nunca fica ambígua)."""
    pessoas = []
    primeiros_nomes = set()
    numero = numero_inicial
    while len(pessoas) < quantas:
        numero += 1
        pessoa = sortear_pessoa(sorteio, numero)
        primeiro_nome = pessoa["nome"].split()[0]
        # Um primeiro nome repetido deixaria o nome curto ou o apelido ambíguo: sorteia outra
        if primeiro_nome in primeiros_nomes:
            continue
        primeiros_nomes.add(primeiro_nome)
        pessoas.append(pessoa)
    return pessoas


def embaralhada(sorteio: random.Random, pessoas: list[dict]) -> list[dict]:
    """Uma cópia da lista em outra ordem (a seção seguinte não segue a ordem da identificação)."""
    copia = list(pessoas)
    sorteio.shuffle(copia)
    return copia


def linha_de_identificacao(sorteio: random.Random, pessoa: dict) -> str:
    """A identificação da pessoa: nome inteiro, CPF e nascimento, em formatos sorteados."""
    return (f"{pessoa['nome']} — CPF {cpf_variado(sorteio, pessoa['cpf'])} — nascimento "
            f"{data_variada(sorteio, pessoa['nascimento'])}")


def linha_de_contrato(sorteio: random.Random, pessoa: dict, referencia: str) -> str:
    """O contrato da pessoa, apresentado pela referência dada (nome curto, apelido, nome inteiro ou CPF)."""
    return (f"{referencia}: {pessoa['cargo']}, admissão em {data_variada(sorteio, pessoa['admissao'])}, salário "
            f"{salario_variado(sorteio, pessoa['salario'])}.")


def linha_de_endereco(pessoa: dict, referencia: str) -> str:
    """O endereço da pessoa, apresentado pela referência dada."""
    return f"{referencia} — {endereco_em_texto(pessoa)}."


def documento_em_secoes(sorteio: random.Random, pessoas: list[dict], referencias: dict,
                        apresentacao: dict | None = None) -> Document:
    """O documento em três seções (identificação, contrato, endereço), cada uma numa ordem.

    referencias: {número da pessoa: como o contrato e o endereço a chamam}. apresentacao: {número: texto extra na
    identificação} (ex.: o apelido).
    """
    documento = Document()
    documento.add_paragraph(sorteio.choice(ABERTURAS))
    # A identificação, na ordem sorteada das pessoas
    documento.add_paragraph(sorteio.choice(TITULOS_DA_IDENTIFICACAO))
    for pessoa in pessoas:
        linha = linha_de_identificacao(sorteio, pessoa)
        if apresentacao and pessoa["numero"] in apresentacao:
            linha += apresentacao[pessoa["numero"]]
        documento.add_paragraph(linha)
    # O contrato, em outra ordem
    documento.add_paragraph(sorteio.choice(TITULOS_DO_CONTRATO))
    for pessoa in embaralhada(sorteio, pessoas):
        documento.add_paragraph(linha_de_contrato(sorteio, pessoa, referencias[pessoa["numero"]]))
    # O endereço, em mais outra ordem
    documento.add_paragraph(sorteio.choice(TITULOS_DO_ENDERECO))
    for pessoa in embaralhada(sorteio, pessoas):
        documento.add_paragraph(linha_de_endereco(pessoa, referencias[pessoa["numero"]]))
    documento.add_paragraph("Qualquer dúvida, é só chamar.")
    return documento


def estrutura_nome_curto(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """Contrato e endereço pelo nome curto. Devolve (documento, terceiros por pessoa)."""
    referencias = {}
    for pessoa in pessoas:
        referencias[pessoa["numero"]] = nome_curto(pessoa)
    return documento_em_secoes(sorteio, pessoas, referencias), []


def estrutura_apelido(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """A identificação apresenta o apelido; contrato e endereço usam só o apelido."""
    apelidos = sorteio.sample(APELIDOS, len(pessoas))
    referencias = {}
    apresentacao = {}
    for pessoa, apelido in zip(pessoas, apelidos):
        referencias[pessoa["numero"]] = apelido
        apresentacao[pessoa["numero"]] = f" (todos chamam de {apelido})"
    return documento_em_secoes(sorteio, pessoas, referencias, apresentacao), []


def estrutura_ordem_trocada(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """Nome inteiro em todas as seções, às vezes em maiúsculas, cada seção numa ordem."""
    referencias = {}
    for pessoa in pessoas:
        # Em 40% das pessoas, o nome aparece em maiúsculas (como em planilhas coladas no texto)
        if sorteio.random() < 0.4:
            referencias[pessoa["numero"]] = pessoa["nome"].upper()
        else:
            referencias[pessoa["numero"]] = pessoa["nome"]
    return documento_em_secoes(sorteio, pessoas, referencias), []


def estrutura_referencia_por_cpf(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """Contrato e endereço citam a pessoa só pelo CPF (escrito de outro jeito que na identificação)."""
    referencias = {}
    for pessoa in pessoas:
        referencias[pessoa["numero"]] = f"CPF {cpf_variado(sorteio, pessoa['cpf'])}"
    return documento_em_secoes(sorteio, pessoas, referencias), []


def estrutura_cpf_repetido(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """Um parágrafo por pessoa (sem o endereço) e, depois, os documentos conferidos (o CPF de novo) e os endereços."""
    documento = Document()
    documento.add_paragraph(sorteio.choice(ABERTURAS))
    for pessoa in pessoas:
        documento.add_paragraph(
            f"{pessoa['nome']}, {pessoa['cargo']}, começa em {data_variada(sorteio, pessoa['admissao'])} com salário "
            f"de {salario_variado(sorteio, pessoa['salario'])}. Nascimento: "
            f"{data_variada(sorteio, pessoa['nascimento'])}. CPF {cpf_variado(sorteio, pessoa['cpf'])}.")
    # A seção que repete o CPF (o mesmo CPF duas vezes no documento)
    documento.add_paragraph("Documentos conferidos pelo RH")
    for pessoa in embaralhada(sorteio, pessoas):
        documento.add_paragraph(f"{nome_curto(pessoa)} — CPF {cpf_variado(sorteio, pessoa['cpf'])} — identidade "
                                f"{pessoa['rg']} — conferido.")
    # Os endereços, pelo nome curto
    documento.add_paragraph(sorteio.choice(TITULOS_DO_ENDERECO))
    for pessoa in embaralhada(sorteio, pessoas):
        documento.add_paragraph(linha_de_endereco(pessoa, nome_curto(pessoa)))
    return documento, []


def estrutura_dependentes(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict]]:
    """Como nome_curto, com uma seção de dependentes que têm os mesmos sobrenomes (não são funcionários)."""
    documento, _ = estrutura_nome_curto(sorteio, pessoas)
    documento.add_paragraph("Dependentes informados (para o plano de saúde, não para a conta)")
    primeiros_nomes_usados = set()
    for pessoa in pessoas:
        primeiros_nomes_usados.add(pessoa["nome"].split()[0])
    terceiros = []
    for pessoa in pessoas:
        # Um filho ou filha com os dois sobrenomes do funcionário e outro primeiro nome
        candidatos = []
        for nome in NOMES_FEMININOS + NOMES_MASCULINOS:
            if nome not in primeiros_nomes_usados:
                candidatos.append(nome)
        primeiro_nome = sorteio.choice(candidatos)
        primeiros_nomes_usados.add(primeiro_nome)
        sobrenomes = pessoa["nome"].split()[1:]
        nome_do_dependente = f"{primeiro_nome} {' '.join(sobrenomes)}"
        # Um filho nascido entre 2010 e 2021 (nunca confundido com a data de nascimento de um funcionário adulto)
        nascimento = sortear_data(sorteio, date(2010, 1, 1), date(2021, 12, 31))
        documento.add_paragraph(f"{nome_do_dependente}, dependente de {nome_curto(pessoa)}, nascido(a) em "
                                f"{data_curta(nascimento)}.")
        terceiros.append({"nome": nome_do_dependente, "data": data_curta(nascimento)})
    return documento, terceiros


# As estruturas, na ordem em que os documentos são gerados
ESTRUTURAS = [("nome_curto", estrutura_nome_curto), ("apelido", estrutura_apelido),
              ("ordem_trocada", estrutura_ordem_trocada), ("referencia_por_cpf", estrutura_referencia_por_cpf),
              ("cpf_repetido", estrutura_cpf_repetido), ("dependentes", estrutura_dependentes)]


def gerar(pasta: Path, semente: int = SEMENTE, conjunto: str = CONJUNTO) -> dict:
    """Gera os documentos e o gabarito na pasta. Devolve o gabarito (o mesmo gravado em gabarito.json).

    Recebe: a pasta, a semente do sorteio e o nome do conjunto. Sem informar, gera o conjunto congelado de sempre
    (texto_dificil, semente 20260929). A reserva usa outra semente, que só fica na pasta dela (fora do repositório).
    """
    sorteio = random.Random(semente)
    pasta.mkdir(parents=True, exist_ok=True)
    gabarito = {"conjunto": conjunto, "semente": semente, "documentos": []}
    numero_da_pessoa = 0
    for nome_da_estrutura, montar in ESTRUTURAS:
        for indice in range(1, DOCUMENTOS_POR_ESTRUTURA + 1):
            # As pessoas deste documento
            quantas = sorteio.randint(MENOS_PESSOAS, MAIS_PESSOAS)
            pessoas = sortear_pessoas(sorteio, quantas, numero_da_pessoa)
            numero_da_pessoa = pessoas[-1]["numero"]
            documento, terceiros = montar(sorteio, pessoas)
            arquivo = f"{NIVEL}_{nome_da_estrutura}_{indice}.docx"
            documento.save(pasta / arquivo)
            # O gabarito: os campos de cada pessoa e, nos dependentes, os dados que não podem virar funcionário
            gabaritos = []
            for posicao, pessoa in enumerate(pessoas):
                gabarito_da_pessoa_atual = gabarito_da_pessoa(pessoa, CAMPOS)
                if terceiros:
                    gabarito_da_pessoa_atual["terceiros"] = [terceiros[posicao]]
                gabaritos.append(gabarito_da_pessoa_atual)
            gabarito["documentos"].append({"arquivo": arquivo, "nivel": NIVEL, "estrutura": nome_da_estrutura,
                                           "pessoas": gabaritos, "cabecalho_para_campo": {}})
            print(f"{arquivo}: {len(pessoas)} pessoas")
    (pasta / "gabarito.json").write_text(json.dumps(gabarito, ensure_ascii=False, indent=1), encoding="utf-8")
    return gabarito


def manifesto_da_pasta(pasta: Path, conjunto: str) -> dict:
    """As impressões digitais (SHA-256) de todos os arquivos da pasta, para congelar um conjunto guardado FORA do Git.

    Recebe: a pasta e o nome do conjunto. Devolve: {"conjunto", "arquivos": {nome: impressão digital}}. Não guarda a
    semente: com ela, qualquer um regeraria o conjunto, e ele deixaria de ser "não visto".
    """
    arquivos = {}
    for arquivo in sorted(pasta.glob("*")):
        # Só arquivos (a mesma impressão digital da foto congelada, que ignora CRLF × LF)
        if arquivo.is_file():
            arquivos[arquivo.name] = impressao_digital(arquivo)
    return {"conjunto": conjunto, "arquivos": arquivos}


def diferencas_do_manifesto(pasta: Path, manifesto: dict) -> list[str]:
    """O que mudou na pasta desde o manifesto (vazio = igual). Usado antes de medir com a reserva.

    Ex.: ["faltando: N6_apelido_1.docx", "alterado: gabarito.json", "sobrando: rascunho.txt"].
    """
    atual = manifesto_da_pasta(pasta, manifesto["conjunto"])["arquivos"]
    diferencas = []
    for nome, impressao in manifesto["arquivos"].items():
        # Arquivo congelado que sumiu ou mudou
        if nome not in atual:
            diferencas.append(f"faltando: {nome}")
        elif atual[nome] != impressao:
            diferencas.append(f"alterado: {nome}")
    for nome in atual:
        # Arquivo que não estava no congelamento
        if nome not in manifesto["arquivos"]:
            diferencas.append(f"sobrando: {nome}")
    return diferencas


def ler_opcoes() -> argparse.Namespace:
    """As opções da linha de comando (sem nenhuma, gera o conjunto congelado de sempre)."""
    leitor = argparse.ArgumentParser(description="Gera o conjunto de texto corrido difícil (EXP-017).")
    leitor.add_argument("--semente", type=int, default=SEMENTE, help="a semente do sorteio (a da reserva não fica no Git)")
    leitor.add_argument("--semente-aleatoria", action="store_true", dest="semente_aleatoria",
                        help="sorteia uma semente de 64 bits (a da reserva), que não aparece no comando nem na tela")
    leitor.add_argument("--pasta", default="", help="onde gravar (padrão: data/avaliacao/documentos_texto_dificil)")
    leitor.add_argument("--conjunto", default=CONJUNTO, help="o nome do conjunto no gabarito")
    leitor.add_argument("--manifesto", default="", help="onde gravar o manifesto das impressões digitais (opcional)")
    return leitor.parse_args()


def conferir_destino(semente: int, pasta_informada: str) -> Path:
    """A pasta onde gravar, com a trava da reserva.

    Recebe: a semente e a pasta informada ("" = nenhuma). Devolve: a pasta.
    Uma semente diferente da do conjunto congelado é uma reserva: ela precisa de --pasta, e a pasta precisa ficar
    FORA do repositório. Sem isso, a reserva gravaria por cima do conjunto congelado ou entraria no Git (e deixaria de
    ser "não vista"). Levanta ValueError com o motivo. Ex.: conferir_destino(20260929, "") → a pasta de sempre.
    """
    # O conjunto congelado de sempre: a pasta dele (ou a informada, para testes)
    if semente == SEMENTE:
        if pasta_informada:
            return Path(pasta_informada)
        return pasta_do_conjunto(CONJUNTO)
    # Uma reserva sem pasta gravaria por cima do conjunto congelado
    if not pasta_informada:
        raise ValueError("Uma semente diferente da do conjunto congelado precisa de --pasta, fora do repositório.")
    pasta = Path(pasta_informada).resolve()
    # Uma reserva dentro de QUALQUER repositório Git (o principal ou uma cópia isolada, que tem um arquivo .git na
    # raiz) entraria no Git e ficaria visível para quem corrige: a pasta e cada pasta acima dela são conferidas
    for lugar in [pasta] + list(pasta.parents):
        if (lugar / ".git").exists():
            raise ValueError("A reserva não pode ficar dentro de um repositório Git: use uma pasta fora dele "
                             "(ex.: D:\\AI_Payroll_Hub\\Bases_de_Teste\\reserva_d39\\...).")
    return pasta


def main() -> None:
    """Gera o conjunto e, se pedido, o manifesto."""
    opcoes = ler_opcoes()
    # A semente: sorteada com 64 bits (a reserva; difícil de adivinhar pelo manifesto) ou a informada
    if opcoes.semente_aleatoria:
        semente = secrets.randbits(64)
    else:
        semente = opcoes.semente
    # A pasta, com a trava da reserva
    pasta = conferir_destino(semente, opcoes.pasta)
    gabarito = gerar(pasta, semente, opcoes.conjunto)
    pessoas = 0
    for documento in gabarito["documentos"]:
        pessoas += len(documento["pessoas"])
    print(f"{len(gabarito['documentos'])} documentos, {pessoas} pessoas")
    # O manifesto (para um conjunto guardado fora do repositório)
    if opcoes.manifesto:
        manifesto = manifesto_da_pasta(pasta, opcoes.conjunto)
        Path(opcoes.manifesto).write_text(json.dumps(manifesto, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"Manifesto com {len(manifesto['arquivos'])} arquivos em {opcoes.manifesto}")


if __name__ == "__main__":
    main()
