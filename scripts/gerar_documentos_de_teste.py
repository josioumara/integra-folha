"""Gera os documentos Word fictícios do experimento EXP-010, em 5 níveis de complexidade, com o gabarito (ADR-73).

Para que serve: medir o quanto a leitura acerta em cada tipo de documento e comparar os dois modos do Leitor de
Documentos (a IA vendo só etiquetas × a IA vendo os dados). Tudo é inventado: nomes, CPFs (com dígito verificador
certo), endereços e e-mails de exemplo.

Os níveis:
    N1 · Word com TABELA limpa                          (lido por regra, sem IA)
    N2 · Word com FICHAS "Rótulo: valor"                (lido por regra, sem IA)
    N3 · TEXTO SIMPLES: uma frase igual para cada pessoa (IA)
    N4 · TEXTO VARIADO: frases, rótulos e formatos diferentes (datas por extenso, CPF sem pontos...) (IA)
    N5 · TEXTO MISTURADO: blocos com títulos variados, dados de outras pessoas (dependentes, contato de emergência),
         extras no salário e ARMADILHAS: data em conflito, CPF que "vem depois", PIS ilegível (IA)

Há DOIS conjuntos, com sementes diferentes (mesmos modelos de frase, pessoas e formatos sorteados de novo):
- desenvolvimento (data/avaliacao/documentos_desenvolvimento/): onde os erros são estudados e o leitor é ajustado;
- prova (data/avaliacao/documentos_prova/): só para medir. Nada é ajustado olhando para ela; senão o número engana.

As funções documento_n1 a documento_n5 também escrevem os documentos da prova por tipo de arquivo
(scripts/gerar_prova_por_tipo.py): quando a pessoa tem o código da profissão ("cbo"), ele entra logo depois do cargo.
As pessoas destes dois conjuntos não têm o código, e os documentos deles saem iguais aos gravados.

O gabarito (gabarito.json de cada conjunto) diz, para cada pessoa: os campos esperados, as "armadilhas"
(campos que devem ficar EM BRANCO, porque o documento não permite decidir) e se uma pergunta é esperada. Também guarda
os dados das outras pessoas (terceiros), para medir se algum deles "vazou" para um campo.

Uso:  python scripts/gerar_documentos_de_teste.py      (as sementes fixas geram sempre os mesmos documentos)
"""
import json
import random
import sys
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from docx import Document

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from services.documentos import gerar_cpf  # noqa: E402

# Os dois conjuntos e a semente de cada um (semente fixa: os mesmos documentos a cada geração)
SEMENTE_POR_CONJUNTO = {"desenvolvimento": 20260925, "prova": 20260926}


def pasta_do_conjunto(conjunto: str) -> Path:
    """Onde ficam os documentos e o gabarito do conjunto (versionados: a prova fica congelada no Git)."""
    return RAIZ / "data" / "avaliacao" / f"documentos_{conjunto}"
# Pessoas por documento
PESSOAS_POR_DOCUMENTO = 8

# Listas próprias deste gerador (diferentes das da base sintética, para a prova não repetir a base)
NOMES_FEMININOS = ["Helena", "Luíza", "Marta", "Renata", "Sílvia", "Tatiane", "Vanessa", "Yasmin", "Priscila",
                   "Olívia", "Noemi", "Lorena", "Jéssica", "Irene", "Gláucia", "Fernanda"]
NOMES_MASCULINOS = ["Anderson", "Caio", "Everton", "Geraldo", "Hugo", "Igor", "Júlio", "Leandro", "Mauro",
                    "Otávio", "Rafael", "Sérgio", "Tiago", "Valter", "Wesley", "Nelson"]
SOBRENOMES = ["Duarte", "Monteiro", "Siqueira", "Brandão", "Nogueira", "Pacheco", "Quintela", "Rezende", "Tavares",
              "Valadares", "Xavier", "Assunção", "Bittencourt", "Coutinho", "Fontes", "Guimarães", "Leal", "Moura",
              "Pimentel", "Sampaio"]
CARGOS = ["analista financeiro", "assistente de logística", "vendedor externo", "recepcionista",
          "técnico de manutenção", "auxiliar de cozinha", "coordenador de projetos", "desenvolvedor back-end",
          "enfermeira do trabalho", "operador de empilhadeira", "supervisor de loja", "designer gráfico"]
CIDADES = [("Campinas", "SP", "130"), ("Curitiba", "PR", "807"), ("Recife", "PE", "520"), ("Goiânia", "GO", "741"),
           ("Salvador", "BA", "419"), ("Porto Alegre", "RS", "904"), ("Belo Horizonte", "MG", "301"),
           ("Fortaleza", "CE", "601"), ("Manaus", "AM", "690"), ("Florianópolis", "SC", "880")]
DDD_DO_ESTADO = {"SP": "19", "PR": "41", "PE": "81", "GO": "62", "BA": "71", "RS": "51", "MG": "31", "CE": "85",
                 "AM": "92", "SC": "48"}
RUAS = ["Rua das Acácias", "Avenida Getúlio Vargas", "Rua Padre Anchieta", "Rua Dona Laura", "Alameda dos Ipês",
        "Rua Marechal Deodoro", "Avenida Beira Mar", "Rua Sete de Setembro", "Travessa São Bento", "Rua do Sol"]
BAIRROS = ["Vila Mariana", "Bigorrilho", "Graças", "Setor Oeste", "Rio Vermelho", "Moinhos de Vento", "Savassi",
           "Aldeota", "Chapada", "Trindade", "Jardim Europa", "Centro"]
ESTADOS_CIVIS = ["solteiro", "casado", "divorciado", "viúvo", "união estável"]
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro",
         "novembro", "dezembro"]


# ============================== Pessoas ==============================

def sortear_data(sorteio: random.Random, inicio: date, fim: date) -> date:
    """Uma data qualquer entre início e fim."""
    return inicio + timedelta(days=sorteio.randint(0, (fim - inicio).days))


def sortear_pessoa(sorteio: random.Random, numero: int) -> dict:
    """Uma pessoa fictícia completa (os valores "de verdade", antes de virar texto)."""
    feminino = sorteio.random() < 0.5
    primeiro_nome = sorteio.choice(NOMES_FEMININOS if feminino else NOMES_MASCULINOS)
    sobrenomes = sorteio.sample(SOBRENOMES, 2)
    nome = f"{primeiro_nome} {sobrenomes[0]} {sobrenomes[1]}"
    cidade, uf, prefixo_do_cep = sorteio.choice(CIDADES)
    mae = f"{sorteio.choice(NOMES_FEMININOS)} {sobrenomes[0]} {sorteio.choice(SOBRENOMES)}"
    celular = f"9{sorteio.randint(8000, 9999)}{sorteio.randint(1000, 9999)}"
    return {
        "numero": numero, "feminino": feminino, "nome": nome, "cpf": gerar_cpf(sorteio),
        "nascimento": sortear_data(sorteio, date(1970, 1, 1), date(2003, 12, 31)),
        "admissao": sortear_data(sorteio, date(2024, 1, 2), date(2025, 12, 20)),
        "cargo": sorteio.choice(CARGOS), "salario": Decimal(sorteio.randint(1800, 9500)) + Decimal("0.00"),
        "mae": mae, "rua": sorteio.choice(RUAS), "numero_da_casa": str(sorteio.randint(10, 2999)),
        "bairro": sorteio.choice(BAIRROS), "cidade": cidade, "uf": uf,
        "cep": f"{prefixo_do_cep}{sorteio.randint(10, 99)}-{sorteio.randint(0, 999):03d}",
        "ddd": DDD_DO_ESTADO[uf], "celular": celular,
        "email": f"{primeiro_nome.lower()}.{sobrenomes[0].lower()}@exemplo.com".replace("í", "i").replace("é", "e")
                 .replace("á", "a").replace("ú", "u").replace("ã", "a").replace("â", "a").replace("ç", "c")
                 .replace("ó", "o"),
        "estado_civil": sorteio.choice(ESTADOS_CIVIS),
        "rg": f"{sorteio.randint(10, 59)}.{sorteio.randint(100, 999)}.{sorteio.randint(100, 999)}-{sorteio.randint(0, 9)}",
        "pis": f"{sorteio.randint(100, 299)}.{sorteio.randint(10000, 99999)}.{sorteio.randint(10, 99)}-{sorteio.randint(0, 9)}",
    }


# ============================== Formatos de escrita ==============================

def dinheiro(valor: Decimal) -> str:
    """4350.00 → "4.350,00" (padrão brasileiro)."""
    inteiro = f"{int(valor):,}".replace(",", ".")
    return f"{inteiro},{int(valor * 100) % 100:02d}"


def data_curta(dia: date) -> str:
    """14/09/1991."""
    return dia.strftime("%d/%m/%Y")


def data_variada(sorteio: random.Random, dia: date) -> str:
    """A mesma data escrita de um jeito sorteado entre os que aparecem em documentos de RH."""
    jeitos = [dia.strftime("%d/%m/%Y"), dia.strftime("%d-%m-%Y"), dia.strftime("%d.%m.%Y"),
              f"{dia.day} de {MESES[dia.month - 1]} de {dia.year}", dia.isoformat(), dia.strftime("%d/%m/%y")]
    return sorteio.choice(jeitos)


def cpf_variado(sorteio: random.Random, cpf: str) -> str:
    """O CPF com pontuação, só com dígitos ou com espaços."""
    digitos = cpf.replace(".", "").replace("-", "")
    jeitos = [cpf, digitos, f"{digitos[:3]} {digitos[3:6]} {digitos[6:9]} {digitos[9:]}"]
    return sorteio.choice(jeitos)


def salario_variado(sorteio: random.Random, valor: Decimal) -> str:
    """O salário com R$, sem R$ ou sem centavos."""
    jeitos = [f"R$ {dinheiro(valor)}", dinheiro(valor), f"R$ {dinheiro(valor)[:-3]}"]
    return sorteio.choice(jeitos)


def celular_variado(sorteio: random.Random, pessoa: dict) -> str:
    """O celular com parênteses, com espaço, só dígitos ou com +55."""
    numero = pessoa["celular"]
    jeitos = [f"({pessoa['ddd']}) {numero[:5]}-{numero[5:]}", f"{pessoa['ddd']} {numero[:5]}-{numero[5:]}",
              f"{pessoa['ddd']}{numero}", f"+55 {pessoa['ddd']} {numero[:5]}-{numero[5:]}"]
    return sorteio.choice(jeitos)


def ele_ou_ela(pessoa: dict, feminino: str, masculino: str) -> str:
    """A palavra no gênero da pessoa (ex.: "admitida" ou "admitido")."""
    return feminino if pessoa["feminino"] else masculino


def trecho_do_cbo(pessoa: dict) -> str:
    """O código da profissão (CBO) logo depois do cargo, quando a pessoa tem um; senão, nada.

    Para que serve: a prova por tipo de arquivo (scripts/gerar_prova_por_tipo.py) precisa dos 4 obrigatórios em todo
    documento, e o CBO é um deles (ADR-143). As pessoas dos conjuntos antigos não têm "cbo": para elas o texto sai igual
    ao de antes, e a prova congelada não muda.
    Ex.: {"cbo": "4110-10", ...} → " (CBO 4110-10)"; sem "cbo" → "".
    """
    # Sem o código, o texto fica exatamente como era
    if not pessoa.get("cbo"):
        return ""
    return f" (CBO {pessoa['cbo']})"


def campos_com_o_cbo(pessoa: dict, campos: list[str]) -> list[str]:
    """Os campos do gabarito da pessoa, com o codigo_cbo logo depois do cargo quando ela tem um código.

    Ex.: ["nome_completo", "cargo", "valor_renda"] e uma pessoa com "cbo" → ["nome_completo", "cargo", "codigo_cbo",
    "valor_renda"]. Sem "cbo", a lista volta igual.
    """
    # Sem o código, os campos são os de sempre
    if not pessoa.get("cbo"):
        return list(campos)
    campos_da_pessoa = []
    for campo in campos:
        campos_da_pessoa.append(campo)
        # O código da profissão entra logo depois do cargo
        if campo == "cargo":
            campos_da_pessoa.append("codigo_cbo")
    return campos_da_pessoa


# ============================== Gabarito ==============================

def gabarito_da_pessoa(pessoa: dict, campos: list[str]) -> dict:
    """Os valores esperados de cada campo pedido, num formato simples (a avaliação padroniza os dois lados)."""
    todos = {
        "nome_completo": pessoa["nome"], "cpf": pessoa["cpf"], "data_nascimento": data_curta(pessoa["nascimento"]),
        "data_admissao": data_curta(pessoa["admissao"]), "cargo": pessoa["cargo"],
        "valor_renda": dinheiro(pessoa["salario"]), "nome_mae": pessoa["mae"],
        "logradouro_residencial": pessoa["rua"], "numero_residencial": pessoa["numero_da_casa"],
        "bairro_residencial": pessoa["bairro"], "municipio_residencial": pessoa["cidade"],
        "uf_residencial": pessoa["uf"], "cep_residencial": pessoa["cep"],
        "telefone_celular": f"{pessoa['ddd']}{pessoa['celular']}", "email_pessoal": pessoa["email"],
        "estado_civil": pessoa["estado_civil"], "numero_documento": pessoa["rg"], "nis_pis": pessoa["pis"],
        # O código da profissão, só nas pessoas que têm um (as da prova por tipo de arquivo)
        "codigo_cbo": pessoa.get("cbo", ""),
    }
    esperado = {}
    for campo in campos:
        esperado[campo] = todos[campo]
    return {"campos": esperado, "armadilhas": [], "espera_duvida": False, "terceiros": []}


# ============================== N1 · Tabela ==============================

COLUNAS_DA_TABELA = [("Nome", "nome_completo"), ("CPF", "cpf"), ("Nascimento", "data_nascimento"),
                     ("Cargo", "cargo"), ("Admissão", "data_admissao"), ("Salário", "valor_renda"),
                     ("Celular", "telefone_celular"), ("E-mail", "email_pessoal"), ("CEP", "cep_residencial"),
                     ("Cidade", "municipio_residencial"), ("UF", "uf_residencial")]


def colunas_da_tabela(pessoas: list[dict]) -> list[tuple[str, str]]:
    """As colunas da tabela do N1: as de sempre e, quando as pessoas têm o código da profissão, a coluna "CBO".

    Ex.: pessoas sem "cbo" → COLUNAS_DA_TABELA; com "cbo" → as mesmas, com ("CBO", "codigo_cbo") depois do cargo.
    """
    # As pessoas da prova por tipo de arquivo têm o código; as dos conjuntos antigos, não
    tem_cbo = bool(pessoas) and bool(pessoas[0].get("cbo"))
    colunas = []
    for cabecalho, campo in COLUNAS_DA_TABELA:
        colunas.append((cabecalho, campo))
        # O código da profissão entra logo depois do cargo
        if campo == "cargo" and tem_cbo:
            colunas.append(("CBO", "codigo_cbo"))
    return colunas


def valores_da_linha_da_tabela(pessoa: dict) -> dict:
    """O que cada campo mostra na linha da pessoa, na tabela do N1 (campo → texto da célula)."""
    return {"nome_completo": pessoa["nome"], "cpf": pessoa["cpf"], "data_nascimento": data_curta(pessoa["nascimento"]),
            "cargo": pessoa["cargo"], "codigo_cbo": pessoa.get("cbo", ""),
            "data_admissao": data_curta(pessoa["admissao"]), "valor_renda": f"R$ {dinheiro(pessoa['salario'])}",
            "telefone_celular": f"({pessoa['ddd']}) {pessoa['celular'][:5]}-{pessoa['celular'][5:]}",
            "email_pessoal": pessoa["email"], "cep_residencial": pessoa["cep"],
            "municipio_residencial": pessoa["cidade"], "uf_residencial": pessoa["uf"]}


def documento_n1(pessoas: list[dict]) -> tuple[Document, list[dict], dict]:
    """Uma tabela limpa, com cabeçalho na primeira linha."""
    documento = Document()
    documento.add_paragraph("Relação de funcionários admitidos para cadastro da conta-salário.")
    # As colunas deste documento (com a do CBO quando as pessoas têm o código)
    colunas = colunas_da_tabela(pessoas)
    tabela = documento.add_table(rows=len(pessoas) + 1, cols=len(colunas))
    for posicao, (cabecalho, _) in enumerate(colunas):
        tabela.cell(0, posicao).text = cabecalho
    # Os campos do gabarito são os das colunas, na mesma ordem
    campos = []
    for _, campo in colunas:
        campos.append(campo)
    gabaritos = []
    for linha, pessoa in enumerate(pessoas, start=1):
        valores = valores_da_linha_da_tabela(pessoa)
        # Cada célula da linha recebe o valor do campo da coluna
        for posicao, (_, campo) in enumerate(colunas):
            tabela.cell(linha, posicao).text = valores[campo]
        gabaritos.append(gabarito_da_pessoa(pessoa, campos))
    cabecalho_para_campo = {}
    for cabecalho, campo in colunas:
        cabecalho_para_campo[cabecalho] = campo
    return documento, gabaritos, cabecalho_para_campo


# ============================== N2 · Fichas ==============================

ROTULOS_DAS_FICHAS = [("Nome", "nome_completo"), ("CPF", "cpf"), ("Data de nascimento", "data_nascimento"),
                      ("Nome da mãe", "nome_mae"), ("Cargo", "cargo"), ("Data de admissão", "data_admissao"),
                      ("Salário", "valor_renda"), ("Endereço", "logradouro_residencial"),
                      ("Número", "numero_residencial"), ("Bairro", "bairro_residencial"),
                      ("Cidade", "municipio_residencial"), ("UF", "uf_residencial"), ("CEP", "cep_residencial"),
                      ("Celular", "telefone_celular"), ("E-mail", "email_pessoal")]


def rotulos_das_fichas(pessoa: dict) -> list[tuple[str, str]]:
    """Os rótulos da ficha da pessoa: os de sempre e, quando ela tem o código da profissão, a linha "CBO".

    Ex.: pessoa sem "cbo" → ROTULOS_DAS_FICHAS; com "cbo" → os mesmos, com ("CBO", "codigo_cbo") depois do cargo.
    """
    rotulos = []
    for rotulo, campo in ROTULOS_DAS_FICHAS:
        rotulos.append((rotulo, campo))
        # O código da profissão entra logo depois do cargo
        if campo == "cargo" and pessoa.get("cbo"):
            rotulos.append(("CBO", "codigo_cbo"))
    return rotulos


def documento_n2(pessoas: list[dict]) -> tuple[Document, list[dict], dict]:
    """Uma ficha "Rótulo: valor" por pessoa."""
    documento = Document()
    gabaritos = []
    for pessoa in pessoas:
        valores = {"nome_completo": pessoa["nome"], "cpf": pessoa["cpf"],
                   "data_nascimento": data_curta(pessoa["nascimento"]), "nome_mae": pessoa["mae"],
                   "cargo": pessoa["cargo"], "codigo_cbo": pessoa.get("cbo", ""),
                   "data_admissao": data_curta(pessoa["admissao"]),
                   "valor_renda": f"R$ {dinheiro(pessoa['salario'])}", "logradouro_residencial": pessoa["rua"],
                   "numero_residencial": pessoa["numero_da_casa"], "bairro_residencial": pessoa["bairro"],
                   "municipio_residencial": pessoa["cidade"], "uf_residencial": pessoa["uf"],
                   "cep_residencial": pessoa["cep"],
                   "telefone_celular": f"({pessoa['ddd']}) {pessoa['celular'][:5]}-{pessoa['celular'][5:]}",
                   "email_pessoal": pessoa["email"]}
        campos = []
        for rotulo, campo in rotulos_das_fichas(pessoa):
            documento.add_paragraph(f"{rotulo}: {valores[campo]}")
            campos.append(campo)
        documento.add_paragraph("")
        gabaritos.append(gabarito_da_pessoa(pessoa, campos))
    # Os rótulos que o documento usa (com o do CBO quando as pessoas têm o código)
    rotulos_do_documento = ROTULOS_DAS_FICHAS
    if pessoas:
        rotulos_do_documento = rotulos_das_fichas(pessoas[0])
    cabecalho_para_campo = {}
    for rotulo, campo in rotulos_do_documento:
        cabecalho_para_campo[rotulo] = campo
    return documento, gabaritos, cabecalho_para_campo


# ============================== N3 · Texto simples ==============================

def documento_n3(pessoas: list[dict]) -> tuple[Document, list[dict], dict]:
    """Uma frase igual para cada pessoa, entre uma saudação e uma assinatura."""
    documento = Document()
    documento.add_paragraph("Bom dia! Seguem os funcionários contratados neste mês.")
    gabaritos = []
    for pessoa in pessoas:
        admitido = ele_ou_ela(pessoa, "admitida", "admitido")
        nascido = ele_ou_ela(pessoa, "nascida", "nascido")
        documento.add_paragraph(
            f"{pessoa['nome']}, CPF {pessoa['cpf']}, {nascido} em {data_curta(pessoa['nascimento'])}, foi {admitido} em "
            f"{data_curta(pessoa['admissao'])} como {pessoa['cargo']}{trecho_do_cbo(pessoa)}, com salário de "
            f"R$ {dinheiro(pessoa['salario'])}.")
        campos = campos_com_o_cbo(pessoa, ["nome_completo", "cpf", "data_nascimento", "data_admissao", "cargo",
                                           "valor_renda"])
        gabaritos.append(gabarito_da_pessoa(pessoa, campos))
    documento.add_paragraph("Obrigada! Equipe de RH.")
    return documento, gabaritos, {}


# ============================== N4 · Texto variado ==============================

CAMPOS_DO_N4 = ["nome_completo", "cpf", "data_nascimento", "data_admissao", "cargo", "valor_renda",
                "telefone_celular", "email_pessoal", "logradouro_residencial", "numero_residencial",
                "bairro_residencial", "municipio_residencial", "uf_residencial", "cep_residencial"]


def paragrafos_n4(sorteio: random.Random, pessoa: dict) -> list[str]:
    """Um ou dois parágrafos sobre a pessoa, com um dos modelos de frase e formatos sorteados."""
    rotulo_do_cpf = sorteio.choice(["CPF", "cpf", "documento fiscal", "cadastro de pessoa física", "CPF/MF"])
    cpf = cpf_variado(sorteio, pessoa["cpf"])
    nascimento = data_variada(sorteio, pessoa["nascimento"])
    admissao = data_variada(sorteio, pessoa["admissao"])
    salario = salario_variado(sorteio, pessoa["salario"])
    celular = celular_variado(sorteio, pessoa)
    endereco = sorteio.choice([
        f"Mora na {pessoa['rua']}, {pessoa['numero_da_casa']}, {pessoa['bairro']}, {pessoa['cidade']} - {pessoa['uf']}, "
        f"CEP {pessoa['cep']}.",
        f"Endereço: {pessoa['rua']} nº {pessoa['numero_da_casa']}, bairro {pessoa['bairro']}, "
        f"{pessoa['cidade']}/{pessoa['uf']}, {pessoa['cep'].replace('-', '')}.",
    ])
    contato = f"Contato {celular}, e-mail {pessoa['email']}."
    # O cargo, com o código da profissão logo depois quando a pessoa tem um (sem ele, só o cargo, como antes)
    cargo = f"{pessoa['cargo']}{trecho_do_cbo(pessoa)}"
    modelos = [
        f"{pessoa['nome']} ({rotulo_do_cpf} {cpf}) começou em {admissao} como {cargo}. "
        f"Nasceu em {nascimento}. Salário de {salario} por mês.",
        f"Contratamos {pessoa['nome']} para a vaga de {cargo}, com início em {admissao} e remuneração de "
        f"{salario}. {rotulo_do_cpf}: {cpf}. Data de nascimento: {nascimento}.",
        f"{pessoa['nome']} — {cargo} — entrada {admissao}, nascimento {nascimento}, {rotulo_do_cpf} {cpf}, "
        f"salário {salario}.",
    ]
    primeiro = sorteio.choice(modelos)
    return [primeiro, f"{endereco} {contato}"]


def documento_n4(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict], dict]:
    """Frases, rótulos e formatos variados; cada pessoa em dois parágrafos."""
    documento = Document()
    documento.add_paragraph("Olá, pessoal do banco. Abaixo, os dados dos novos colaboradores para abrir as contas.")
    gabaritos = []
    for pessoa in pessoas:
        for paragrafo in paragrafos_n4(sorteio, pessoa):
            documento.add_paragraph(paragrafo)
        gabaritos.append(gabarito_da_pessoa(pessoa, campos_com_o_cbo(pessoa, CAMPOS_DO_N4)))
    documento.add_paragraph("Qualquer dúvida, estou à disposição. Atenciosamente, Administrativo.")
    return documento, gabaritos, {}


# ============================== N5 · Texto misturado ==============================

TITULOS_DO_N5 = ["FICHA {numero:03d}", "ANOTAÇÃO RECEBIDA DO GESTOR", "DADOS ENVIADOS PELA FILIAL",
                 "INFORMAÇÕES DO E-MAIL DE CONTRATAÇÃO", "FICHA DE ADMISSÃO {numero}", "RESUMO DA INTEGRAÇÃO",
                 "MENSAGEM COPIADA DO RECRUTAMENTO", "PENDÊNCIAS MISTURADAS COM CADASTRO"]
ARMADILHAS = ["admissao_em_conflito", "salario_com_extras", "cpf_depois", "pis_ilegivel"]
CAMPOS_DO_N5 = CAMPOS_DO_N4 + ["nome_mae", "estado_civil", "numero_documento", "nis_pis"]


def terceiro(sorteio: random.Random, pessoa: dict) -> tuple[str, str, dict]:
    """Uma pessoa de fora do cadastro (dependente ou contato de emergência). Devolve (frase, nome, valores)."""
    nome = f"{sorteio.choice(NOMES_FEMININOS + NOMES_MASCULINOS)} {pessoa['nome'].split()[1]}"
    if sorteio.random() < 0.5:
        nascimento = sortear_data(sorteio, date(2008, 1, 1), date(2022, 12, 31))
        frase = f"Dependente informado: {nome}, filho(a), nasc. {data_curta(nascimento)}; falta a certidão."
        return frase, nome, {"nome": nome, "data": data_curta(nascimento)}
    telefone = f"{pessoa['ddd']} 3{sorteio.randint(100, 999)}-{sorteio.randint(1000, 9999)}"
    frase = f"Contato de emergência: {nome} (irmão/irmã), telefone {telefone}."
    return frase, nome, {"nome": nome, "telefone": telefone}


def paragrafos_n5(sorteio: random.Random, pessoa: dict, armadilha: str) -> tuple[list[str], dict]:
    """Os parágrafos de uma pessoa no nível misturado, com a armadilha dela. Devolve (parágrafos, ajuste do gabarito)."""
    ajuste = {"armadilhas": [], "espera_duvida": False, "terceiros": [], "valor_renda": None, "sem": []}
    titulo = sorteio.choice(TITULOS_DO_N5).format(numero=pessoa["numero"])
    nome_do_titulo = pessoa["nome"].upper() if sorteio.random() < 0.4 else pessoa["nome"]
    rotulo_do_cpf = sorteio.choice(["CPF", "cpf_formatado =", "Documento fiscal", "Cadastro pessoa:", "reg_cliente:"])
    cpf = cpf_variado(sorteio, pessoa["cpf"])
    admissao = data_variada(sorteio, pessoa["admissao"])
    salario = salario_variado(sorteio, pessoa["salario"])
    # A linha dos documentos (CPF, RG, nascimento); a armadilha do CPF tira o CPF daqui
    linha_dos_documentos = (f"{rotulo_do_cpf} {cpf} | identidade {pessoa['rg']} SSP-{pessoa['uf']} | "
                            f"nascimento {data_variada(sorteio, pessoa['nascimento'])}")
    if armadilha == "cpf_depois":
        linha_dos_documentos = (f"identidade {pessoa['rg']} SSP-{pessoa['uf']} | nascimento "
                                f"{data_variada(sorteio, pessoa['nascimento'])}. O CPF ainda não veio: "
                                f"{ele_ou_ela(pessoa, 'ela', 'ele')} vai mandar depois.")
        ajuste["sem"].append("cpf")
        ajuste["espera_duvida"] = True
    # A linha da contratação (cargo, admissão, salário); duas armadilhas mexem aqui. O cargo leva o código da
    # profissão logo depois quando a pessoa tem um (sem ele, só o cargo, como antes)
    cargo = f"{pessoa['cargo'].capitalize()}{trecho_do_cbo(pessoa)}"
    linha_da_contratacao = f"{cargo}. Admissão: {admissao}. Salário mensal {salario}."
    if armadilha == "admissao_em_conflito":
        outra_data = pessoa["admissao"] + timedelta(days=sorteio.choice([7, 14, 21]))
        linha_da_contratacao = (f"{cargo}. A ficha diz admissão em {data_curta(pessoa['admissao'])}, "
                                f"mas o gestor escreveu {data_curta(outra_data)}; confirmar. Salário mensal {salario}.")
        ajuste["armadilhas"].append("data_admissao")
        ajuste["espera_duvida"] = True
    if armadilha == "salario_com_extras":
        comissao = sorteio.choice(["1,5%", "2%", "3%"])
        garantia = dinheiro(Decimal(sorteio.choice([600, 800, 900, 1200])))
        linha_da_contratacao = (f"{cargo}. Início {admissao}. Base salarial {salario} fixo, mais "
                                f"comissão de {comissao} sobre vendas e garantia mínima de R$ {garantia} nos três "
                                "primeiros meses.")
    # A linha do endereço e do contato
    linha_do_endereco = (f"End.: {pessoa['rua']} {pessoa['numero_da_casa']}, {pessoa['bairro']}, {pessoa['cidade']} "
                         f"{pessoa['uf']}, {pessoa['cep']}. Celular {celular_variado(sorteio, pessoa)}. "
                         f"E-mail pessoal: {pessoa['email']}.")
    # A linha pessoal (mãe, estado civil, PIS) e o ruído (benefícios e uma pessoa de fora)
    linha_do_pis = f"PIS {pessoa['pis']}."
    if armadilha == "pis_ilegivel":
        linha_do_pis = f"PIS ilegível na ficha, parece {pessoa['pis']}."
        ajuste["armadilhas"].append("nis_pis")
        ajuste["espera_duvida"] = True
    frase_do_terceiro, _, valores_do_terceiro = terceiro(sorteio, pessoa)
    ajuste["terceiros"].append(valores_do_terceiro)
    linha_pessoal = (f"Nome da mãe: {pessoa['mae']}. Estado civil: {pessoa['estado_civil']}. {linha_do_pis} "
                     f"Vale-refeição de R$ {sorteio.choice(['32,00', '38,50', '42,00'])} por dia. {frase_do_terceiro}")
    paragrafos = [titulo, nome_do_titulo, linha_dos_documentos, linha_da_contratacao, linha_do_endereco, linha_pessoal]
    return paragrafos, ajuste


def documento_n5(sorteio: random.Random, pessoas: list[dict]) -> tuple[Document, list[dict], dict]:
    """Blocos com títulos variados, ruído, dados de terceiros e uma armadilha por pessoa (as 4 se repetem)."""
    documento = Document()
    documento.add_paragraph("Relação de colaboradores para cadastro")
    documento.add_paragraph("Material reunido pelo administrativo a partir de fichas, e-mails e anotações. Alguns "
                            "itens ainda precisam de confirmação.")
    gabaritos = []
    for posicao, pessoa in enumerate(pessoas):
        armadilha = ARMADILHAS[posicao % len(ARMADILHAS)]
        paragrafos, ajuste = paragrafos_n5(sorteio, pessoa, armadilha)
        for paragrafo in paragrafos:
            documento.add_paragraph(paragrafo)
        campos = []
        # Os campos do N5 (com o do CBO quando a pessoa tem o código), menos os que a armadilha tira
        for campo in campos_com_o_cbo(pessoa, CAMPOS_DO_N5):
            if campo not in ajuste["sem"] and campo not in ajuste["armadilhas"]:
                campos.append(campo)
        gabarito = gabarito_da_pessoa(pessoa, campos)
        gabarito["armadilhas"] = ajuste["armadilhas"]
        gabarito["espera_duvida"] = ajuste["espera_duvida"]
        gabarito["terceiros"] = ajuste["terceiros"]
        gabarito["armadilha"] = armadilha
        gabaritos.append(gabarito)
    return documento, gabaritos, {}


# ============================== Gravação ==============================

def gerar_conjunto(conjunto: str, pasta: Path) -> None:
    """Gera os 7 documentos de um conjunto (N1, N2, N3, dois N4 e dois N5) e o gabarito, na pasta indicada."""
    semente = SEMENTE_POR_CONJUNTO[conjunto]
    sorteio = random.Random(semente)
    pasta.mkdir(parents=True, exist_ok=True)
    planos = [("N1_tabela", "N1", documento_n1), ("N2_fichas", "N2", documento_n2),
              ("N3_texto_simples", "N3", documento_n3), ("N4_texto_variado_a", "N4", documento_n4),
              ("N4_texto_variado_b", "N4", documento_n4), ("N5_texto_misturado_a", "N5", documento_n5),
              ("N5_texto_misturado_b", "N5", documento_n5)]
    gabarito_geral = {"conjunto": conjunto, "semente": semente, "documentos": []}
    numero_da_pessoa = 0
    for nome_do_documento, nivel, montar in planos:
        pessoas = []
        for _ in range(PESSOAS_POR_DOCUMENTO):
            numero_da_pessoa += 1
            pessoas.append(sortear_pessoa(sorteio, numero_da_pessoa))
        # N4 e N5 sorteiam formatos: recebem o sorteio; N1 a N3 são fixos
        if nivel in ("N4", "N5"):
            documento, gabaritos, cabecalho_para_campo = montar(sorteio, pessoas)
        else:
            documento, gabaritos, cabecalho_para_campo = montar(pessoas)
        arquivo = f"{nome_do_documento}.docx"
        documento.save(pasta / arquivo)
        gabarito_geral["documentos"].append({"arquivo": arquivo, "nivel": nivel, "pessoas": gabaritos,
                                             "cabecalho_para_campo": cabecalho_para_campo})
        print(f"{conjunto}/{arquivo}: {len(gabaritos)} pessoas")
    (pasta / "gabarito.json").write_text(json.dumps(gabarito_geral, ensure_ascii=False, indent=1), encoding="utf-8")


def main() -> None:
    """Gera os dois conjuntos, cada um na sua pasta."""
    for conjunto in SEMENTE_POR_CONJUNTO:
        gerar_conjunto(conjunto, pasta_do_conjunto(conjunto))


if __name__ == "__main__":
    main()
