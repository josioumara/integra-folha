"""Leitura de documento Word (.docx) enviado pela empresa (ADR-72).

A empresa manda o documento do jeito que ele estiver. Este arquivo tenta, em ordem, do mais barato ao mais caro:
1. TABELA: o documento tem uma tabela de funcionários. Ela é lida direto, sem IA e sem custo.
   (Antes das fichas, a leitura em services/ingestao.py confere se a lista foi colada como TEXTO EM COLUNAS,
   "Nome;CPF;Cargo" em cada parágrafo: se foi, ela é lida como um CSV, também sem IA; ver texto_sem_tabela.)
2. FICHAS: o documento é uma sequência de "Rótulo: valor" ("Nome: Maria Souza", "CPF: 529.982.247-25"...).
   Cada ficha vira uma linha, também sem IA.
3. TEXTO CORRIDO: nada disso. O Agente Leitor de Documentos (IA) preenche os campos do layout, pessoa por pessoa,
   lendo o texto real (ADR-101; versão 2 do Leitor no ADR-73).

O resultado tem o mesmo jeito da leitura de um CSV (uma lista de linhas de texto, a primeira é o cabeçalho), então
tudo o que vem depois (achar o cabeçalho, guardrail, retrato das colunas, Interpretador) é igual ao da planilha.
"""
import io
import re
from dataclasses import dataclass, field

from docx import Document

from agents import leitor_de_documentos
from models.contratos import carregar_layout
from services import detector_de_dados, guardrail_injecao, teto_de_gasto

# Maior texto corrido aceito (cerca de 150 funcionários): a divisão em blocos manda o documento inteiro, numerado,
# para a IA pequena. Acima disso, a empresa divide o documento.
MAXIMO_DE_CARACTERES_PARA_A_IA = 80_000
# Uma linha de ficha: um rótulo curto e sem números, dois-pontos e o valor. Ex.: "Data de admissão: 05/03/2026".
# Sem números no rótulo porque "Funcionário 1: A Maria..." é texto corrido numerado, não uma ficha
PADRAO_DA_LINHA_DE_FICHA = re.compile(r"^\s*([^:\d]{2,40}?)\s*:\s*(.+?)\s*$")
# Pelo menos esta fração das linhas precisa ser "Rótulo: valor" para o documento ser lido como fichas
FRACAO_MINIMA_DE_LINHAS_DE_FICHA = 0.8


class DocumentoRecusado(ValueError):
    """O documento não pode ser lido. A mensagem vai para a empresa, em linguagem simples.

    execucoes_dos_agentes: quando a recusa veio depois de a IA trabalhar (a IA caiu no meio da leitura, ou leu e não
    achou ninguém), as execuções medidas do Leitor e do Conferidor. O envio não chega a existir, mas o trabalho da IA
    aconteceu: services/processamentos.py grava essas execuções mesmo assim. Nada pessoal: só horários e situação.
    """

    def __init__(self, mensagem: str, execucoes_dos_agentes: list[dict] | None = None):
        # A mensagem continua sendo o texto do erro (str(erro)), como antes
        super().__init__(mensagem)
        # Sem execuções informadas, uma lista vazia nova (nunca dividida com outro erro)
        if execucoes_dos_agentes is None:
            execucoes_dos_agentes = []
        self.execucoes_dos_agentes = execucoes_dos_agentes


@dataclass
class LeituraDoWord:
    """As linhas que saíram do Word (a primeira é o cabeçalho) e o que a empresa precisa saber."""

    linhas: list[list[str]]                                  # cabeçalho + uma linha por funcionário
    como_foi_lido: str                                       # "tabela", "fichas" ou "texto corrido"
    avisos: list[str] = field(default_factory=list)          # recados para a empresa
    duvidas: list[str] = field(default_factory=list)         # todas as perguntas da IA, em texto (texto corrido)
    perguntas: list[dict] = field(default_factory=list)      # as presas a uma pessoa: {registro, campo, pergunta}
    duvidas_gerais: list[str] = field(default_factory=list)  # as que não são de ninguém (trecho não lido)
    origem_das_colunas: dict = field(default_factory=dict)   # {coluna: {campo, rotulo, pessoas}} (texto corrido)
    alertas_guardrail: list[dict] = field(default_factory=list)  # parágrafos removidos pelo guardrail
    uso_da_ia: dict | None = None                            # modo, modelo e custo, quando a IA foi chamada


# ============================== Abrir o documento ==============================

def _abrir(conteudo: bytes):
    """Abre o .docx. Documento corrompido (ou que não é Word de verdade) é recusado."""
    try:
        return Document(io.BytesIO(conteudo))
    except Exception as erro:
        raise DocumentoRecusado("Não foi possível abrir o documento Word. Ele pode estar corrompido; salve de novo "
                                "como .docx e envie outra vez.") from erro


def texto_sem_tabela(conteudo: bytes) -> str:
    """O texto do Word, um parágrafo por linha, quando o documento NÃO tem uma tabela de funcionários; senão, "".

    Serve para a leitura (services/ingestao.py) conferir se a empresa colou a lista como texto em colunas
    ("Ana Souza;529.982.247-25;Analista" em cada parágrafo). Se colou, a lista é lida como um CSV, por regra e sem
    custo, antes das fichas e do texto corrido.
    Os parágrafos NÃO perdem os espaços e tabulações das pontas: a tabulação no começo é uma célula vazia.
    Recusa (DocumentoRecusado) o documento que não abre.
    """
    documento = _abrir(conteudo)
    # Com uma tabela de funcionários, quem vale é a tabela (o caminho de sempre)
    linhas_da_tabela, _ = _linhas_das_tabelas(_tabelas(documento))
    if linhas_da_tabela:
        return ""
    # Os parágrafos com algum texto, como estão
    paragrafos = []
    for paragrafo in documento.paragraphs:
        if paragrafo.text.strip():
            paragrafos.append(paragrafo.text)
    return "\n".join(paragrafos)


def _paragrafos(documento) -> list[str]:
    """O texto de cada parágrafo do documento, sem os espaços das pontas (parágrafos vazios ficam de fora)."""
    paragrafos = []
    for paragrafo in documento.paragraphs:
        texto = paragrafo.text.strip()
        if texto:
            paragrafos.append(texto)
    return paragrafos


def _tabelas(documento) -> list[list[list[str]]]:
    """Cada tabela do documento, como uma lista de linhas; cada linha, uma lista com o texto das células.

    Célula mesclada na HORIZONTAL (ex.: um título ocupando a linha inteira): o texto fica só na primeira coluna do
    grupo, e as outras ficam vazias. Por quê: repetido em cada coluna, o título teria cara de cabeçalho. A biblioteca
    do Word entrega a mesma célula uma vez por coluna que ela ocupa; é assim que o grupo é reconhecido.
    Célula mesclada na VERTICAL (ex.: a unidade valendo para duas pessoas): o texto vale para cada linha do grupo.
    Exemplo: título mesclado em 3 colunas → ["Lista de funcionários", "", ""].
    """
    tabelas = []
    for tabela in documento.tables:
        linhas = []
        for linha in tabela.rows:
            celulas = []
            celula_anterior = None
            for celula in linha.cells:
                # A mesma célula da coluna anterior: é a continuação de uma mesclagem na horizontal
                if celula is celula_anterior:
                    celulas.append("")
                else:
                    celulas.append(celula.text.strip())
                celula_anterior = celula
            linhas.append(celulas)
        tabelas.append(linhas)
    return tabelas


# ============================== 1. Tabela ==============================

def _linhas_das_tabelas(tabelas: list[list[list[str]]]) -> tuple[list[list[str]], list[str]]:
    """As linhas da maior tabela do documento. Devolve (linhas, avisos); sem tabela útil, devolve ([], []).

    Uma tabela que o Word partiu em duas (por exemplo, na virada da página) volta a ser uma só: as tabelas com o
    mesmo cabeçalho da maior são somadas a ela.
    """
    # Tabela útil tem cabeçalho e pelo menos um funcionário (duas linhas)
    tabelas_uteis = []
    for tabela in tabelas:
        if len(tabela) >= 2:
            tabelas_uteis.append(tabela)
    if not tabelas_uteis:
        return [], []
    # A maior tabela (a que tem mais linhas) é a dos funcionários
    maior = tabelas_uteis[0]
    for tabela in tabelas_uteis:
        if len(tabela) > len(maior):
            maior = tabela
    linhas = list(maior)
    avisos = []
    # Soma as outras tabelas com o mesmo cabeçalho (sem repetir o cabeçalho)
    for tabela in tabelas_uteis:
        if tabela is not maior and tabela[0] == maior[0]:
            linhas.extend(tabela[1:])
    # Avisa se alguma tabela ficou de fora
    tabelas_de_fora = 0
    for tabela in tabelas_uteis:
        if tabela is not maior and tabela[0] != maior[0]:
            tabelas_de_fora += 1
    if tabelas_de_fora:
        avisos.append(f"O documento tem {tabelas_de_fora + 1} tabelas diferentes; li a maior, com "
                      f"{len(maior) - 1} linhas. Se os funcionários estão em outra, mande só ela.")
    return linhas, avisos


# ============================== 2. Fichas "Rótulo: valor" ==============================

def _documento_e_de_fichas(paragrafos: list[str]) -> bool:
    """True se quase todas as linhas são "Rótulo: valor" e há pelo menos dois rótulos diferentes."""
    if not paragrafos:
        return False
    linhas_de_ficha = 0
    rotulos = set()
    for paragrafo in paragrafos:
        encontrado = PADRAO_DA_LINHA_DE_FICHA.match(paragrafo)
        if encontrado:
            linhas_de_ficha += 1
            rotulos.add(encontrado.group(1).lower())
    return linhas_de_ficha >= FRACAO_MINIMA_DE_LINHAS_DE_FICHA * len(paragrafos) and len(rotulos) >= 2


def _linhas_das_fichas(paragrafos: list[str]) -> list[list[str]]:
    """Transforma as fichas em tabela. Uma ficha nova começa quando um rótulo se repete.

    Exemplo:
        "Nome: Ana", "CPF: 111", "Nome: Bia", "CPF: 222"
        -> [["Nome", "CPF"], ["Ana", "111"], ["Bia", "222"]]
    Linha que não é "Rótulo: valor" (ex.: "Segue a lista") é ignorada.
    """
    # Os rótulos, na ordem em que aparecem pela primeira vez (viram o cabeçalho)
    rotulos_em_ordem = []
    fichas = []
    ficha_atual = {}
    for paragrafo in paragrafos:
        encontrado = PADRAO_DA_LINHA_DE_FICHA.match(paragrafo)
        if encontrado is None:
            continue
        rotulo = encontrado.group(1).strip()
        valor = encontrado.group(2).strip()
        # Rótulo repetido: a ficha anterior terminou
        if rotulo in ficha_atual:
            fichas.append(ficha_atual)
            ficha_atual = {}
        ficha_atual[rotulo] = valor
        if rotulo not in rotulos_em_ordem:
            rotulos_em_ordem.append(rotulo)
    # A última ficha também entra
    if ficha_atual:
        fichas.append(ficha_atual)
    # Monta a tabela: cabeçalho e, para cada ficha, o valor de cada rótulo (vazio se a ficha não tem)
    linhas = [rotulos_em_ordem]
    for ficha in fichas:
        linha = []
        for rotulo in rotulos_em_ordem:
            linha.append(ficha.get(rotulo, ""))
        linhas.append(linha)
    return linhas


# ============================== 3. Texto corrido (IA) ==============================

def _tirar_paragrafos_suspeitos(paragrafos: list[str]) -> tuple[list[str], list[dict]]:
    """Troca o parágrafo com cara de ordem para a IA pelo aviso do guardrail, ANTES de a IA ler (ADR-38).

    Devolve (parágrafos limpos, alertas com o número do parágrafo).
    """
    limpos = []
    alertas = []
    for numero, paragrafo in enumerate(paragrafos, start=1):
        if guardrail_injecao.e_suspeito(paragrafo):
            alertas.append({"linha": numero, "coluna": 0, "onde": "parágrafo",
                            "padroes": guardrail_injecao.padroes_encontrados(paragrafo)})
            limpos.append(guardrail_injecao.SUBSTITUTO)
        else:
            limpos.append(paragrafo)
    return limpos, alertas


def _rotulo_sem_ligacoes(rotulo: str) -> str:
    """O rótulo sem as palavras de ligação das pontas, para o mesmo jeito de chamar o dado virar uma coluna só.

    Ex.: "com salário de" → "salário"; "salário de" → "salário"; "Também entrou a" → "entrou"; "CPF" → "CPF".
    """
    palavras = rotulo.split()
    # Tira as de ligação do começo ("com", "Também", "a")
    while palavras and detector_de_dados.so_palavras_de_ligacao(palavras[0]):
        palavras.pop(0)
    # E as do fim ("de", "em")
    while palavras and detector_de_dados.so_palavras_de_ligacao(palavras[-1]):
        palavras.pop()
    return " ".join(palavras)


def _rotulo_da_pessoa(tabela, numero_da_pessoa: int, campo: str) -> str:
    """O rótulo (já sem as palavras de ligação) que esta pessoa tem neste campo; vazio se a IA achou pelo lugar."""
    return _rotulo_sem_ligacoes(tabela.rotulos_das_pessoas[numero_da_pessoa].get(campo, ""))


def _chaves_dos_valores(tabela) -> list[tuple[str, str]]:
    """Cada jeito de a empresa chamar um dado, na ordem do layout: [(campo, rótulo em minúsculas)].

    O rótulo vazio quer dizer que a IA achou o dado pelo lugar no texto (ex.: o nome logo abaixo do título).
    """
    chaves = []
    for posicao_do_campo, campo in enumerate(tabela.colunas):
        for numero_da_pessoa, valores_da_pessoa in enumerate(tabela.funcionarios):
            # Só conta o valor que a pessoa tem
            if not valores_da_pessoa[posicao_do_campo]:
                continue
            rotulo = _rotulo_da_pessoa(tabela, numero_da_pessoa, campo)
            chave = (campo, rotulo.casefold())
            if chave not in chaves:
                chaves.append(chave)
    return chaves


def _rotulo_como_escrito(tabela, chave: tuple[str, str]) -> str:
    """O rótulo como a empresa mais escreveu (maiúsculas e acentos): "admissão" → "Admissão"."""
    campo, rotulo_em_minusculas = chave
    contagem = {}
    for numero_da_pessoa in range(len(tabela.rotulos_das_pessoas)):
        rotulo = _rotulo_da_pessoa(tabela, numero_da_pessoa, campo)
        if rotulo.casefold() == rotulo_em_minusculas:
            contagem[rotulo] = contagem.get(rotulo, 0) + 1
    mais_usado = ""
    for rotulo, vezes in contagem.items():
        if not mais_usado or vezes > contagem[mais_usado]:
            mais_usado = rotulo
    return mais_usado


def _nomes_das_colunas(tabela, chaves: list[tuple[str, str]]) -> list[str]:
    """O nome de cada coluna da tabela montada: o rótulo do documento, numerado quando se repete.

    O mesmo rótulo em dois campos acontece quando a IA dividiu um dado em pedaços (ex.: "Endereço" virou rua, número e
    bairro) ou quando a mesma palavra puxa dois dados ("entrou" antes do nome e antes da data): as colunas ficam
    "Endereço (1)", "Endereço (2)"... e a tela mostra um exemplo de cada. Sem rótulo: "Sem rótulo 1", "Sem rótulo 2".
    Ex.: [("cpf", "cpf"), ("logradouro", "endereço"), ("numero", "endereço")] → ["CPF", "Endereço (1)", "Endereço (2)"].
    """
    # Quantas vezes cada rótulo aparece (em campos diferentes)
    vezes_do_rotulo = {}
    for _, rotulo_em_minusculas in chaves:
        vezes_do_rotulo[rotulo_em_minusculas] = vezes_do_rotulo.get(rotulo_em_minusculas, 0) + 1
    ja_usadas = {}
    nomes = []
    for chave in chaves:
        rotulo_em_minusculas = chave[1]
        ja_usadas[rotulo_em_minusculas] = ja_usadas.get(rotulo_em_minusculas, 0) + 1
        parte = ja_usadas[rotulo_em_minusculas]
        if not rotulo_em_minusculas:
            # Sem rótulo: numerado só se houver mais de um
            nome = "Sem rótulo" if vezes_do_rotulo[""] == 1 else f"Sem rótulo {parte}"
        elif vezes_do_rotulo[rotulo_em_minusculas] == 1:
            nome = _rotulo_como_escrito(tabela, chave)
        else:
            nome = f"{_rotulo_como_escrito(tabela, chave)} ({parte})"
        nomes.append(nome)
    return nomes


def tabela_pelos_rotulos(tabela) -> tuple[list[list[str]], dict]:
    """A tabela do texto corrido com UMA COLUNA POR RÓTULO do documento, e não por campo do banco.

    Por quê: a empresa confere o jeito dela ("Documento fiscal", "Registro do cliente") virando o padrão do banco, e
    cada jeito numa linha própria. Se a IA juntou dois dados diferentes no mesmo campo (ex.: "Admissão" e "Efetivação"
    como data de admissão), a empresa troca o campo só daquela linha.
    Recebe: a TabelaDoDocumento do Leitor (colunas = campos do layout; rotulos_das_pessoas).
    Devolve: (linhas — cabeçalho + uma linha por pessoa; origem — {coluna: {"campo", "rotulo", "pessoas"}}).
    Cada pessoa tem o valor de um campo numa coluna só; as outras colunas do mesmo campo ficam vazias nela.
    """
    chaves = _chaves_dos_valores(tabela)
    nomes = _nomes_das_colunas(tabela, chaves)
    origem = {}
    for chave, nome in zip(chaves, nomes):
        origem[nome] = {"campo": chave[0], "rotulo": _rotulo_como_escrito(tabela, chave), "pessoas": 0}
    linhas = [nomes]
    for numero_da_pessoa, valores_da_pessoa in enumerate(tabela.funcionarios):
        linha = [""] * len(chaves)
        for posicao_do_campo, campo in enumerate(tabela.colunas):
            valor = valores_da_pessoa[posicao_do_campo]
            if not valor:
                continue
            # A coluna do jeito que ESTA pessoa chamou o dado
            rotulo = _rotulo_da_pessoa(tabela, numero_da_pessoa, campo)
            posicao_da_coluna = chaves.index((campo, rotulo.casefold()))
            linha[posicao_da_coluna] = valor
            origem[nomes[posicao_da_coluna]]["pessoas"] += 1
        linhas.append(linha)
    return linhas, origem


def linhas_pelos_campos(leitura: LeituraDoWord) -> list[list[str]]:
    """O caminho de volta: a tabela com uma coluna por CAMPO do layout (a régua do EXP-010 compara por campo).

    Tabela e fichas não mudam (não têm origem). No texto corrido, as colunas do mesmo campo se juntam: cada pessoa
    tem o valor numa só delas. Ex.: colunas "CPF" e "Documento fiscal" (as duas do campo cpf) → uma coluna "cpf".
    """
    if not leitura.origem_das_colunas:
        return leitura.linhas
    cabecalho = leitura.linhas[0]
    # Os campos, na ordem em que aparecem nas colunas
    campos = []
    for nome_da_coluna in cabecalho:
        campo = leitura.origem_das_colunas[nome_da_coluna]["campo"]
        if campo not in campos:
            campos.append(campo)
    linhas = [campos]
    for linha in leitura.linhas[1:]:
        valores_por_campo = [""] * len(campos)
        for nome_da_coluna, valor in zip(cabecalho, linha):
            if valor:
                valores_por_campo[campos.index(leitura.origem_das_colunas[nome_da_coluna]["campo"])] = valor
        linhas.append(valores_por_campo)
    return linhas


def _marcar_o_guardrail(execucoes_medidas: list[dict], alertas: list[dict]) -> list[dict]:
    """Marca "guardrail disparado" na execução do Leitor quando o guardrail de entrada tirou algum parágrafo.

    Recebe: as execuções medidas na leitura (a do Leitor e, se rodou, a do Conferidor); os alertas do guardrail de
    entrada (um por parágrafo tirado). Devolve a mesma lista. O guardrail age antes da IA, aqui neste arquivo; por
    isso é aqui que a execução do Leitor fica sabendo. Ex.: 1 parágrafo "ignore as instruções..." tirado → True.
    """
    for execucao in execucoes_medidas:
        if execucao["agente"] == leitor_de_documentos.NOME_DO_AGENTE and alertas:
            execucao["guardrail_disparado"] = True
    return execucoes_medidas


def _ler_texto_corrido(paragrafos: list[str], cliente, campos) -> LeituraDoWord:
    """Pede ao Agente Leitor de Documentos os campos de cada pessoa do texto corrido. Recusa se não achar ninguém.

    campos: os campos do layout vigente (a IA preenche direto esses campos).
    """
    # 1. Guardrail de entrada antes de qualquer IA
    paragrafos_limpos, alertas = _tirar_paragrafos_suspeitos(paragrafos)
    texto = "\n".join(paragrafos_limpos)
    # Texto grande demais para uma leitura só
    if len(texto) > MAXIMO_DE_CARACTERES_PARA_A_IA:
        raise DocumentoRecusado("O documento é grande demais para a leitura de texto corrido (cerca de 150 "
                                "funcionários por envio). Divida em partes ou mande uma planilha.")
    # 2. A IA preenche os campos do layout, pessoa por pessoa, lendo o texto real (ADR-101)
    if cliente is None:
        cliente = leitor_de_documentos.cliente_padrao()
    try:
        tabela = leitor_de_documentos.ler(texto, campos, cliente)
    except leitor_de_documentos.IAIndisponivel as erro:
        # As execuções medidas até a falha seguem no erro (quem cria o envio as grava, com o guardrail marcado)
        execucoes_medidas = _marcar_o_guardrail(list(erro.execucoes_dos_agentes), alertas)
        # O motivo técnico (ex.: sem crédito no provedor) fica no registro do servidor; a empresa recebe o recado simples
        raise DocumentoRecusado("A leitura de texto corrido pelo Agente Leitor está indisponível agora. Tente de novo "
                                "em alguns minutos; se continuar, mande os dados numa tabela ou planilha.",
                                execucoes_medidas) from erro
    except teto_de_gasto.TetoDeGastoAtingido as erro:
        # A IA foi pausada pelo teto de gasto (ADR-131): o documento não entra (a empresa continua com ele) e as
        # execuções medidas, com o ERRO da pausa, vão para a Telemetria
        execucoes_medidas = _marcar_o_guardrail(list(getattr(erro, "execucoes_dos_agentes", ())), alertas)
        raise DocumentoRecusado(teto_de_gasto.RECADO_DO_ARQUIVO_RECUSADO, execucoes_medidas) from erro
    # As execuções medidas nesta leitura (o Leitor e, se rodou, o Conferidor), com o guardrail de entrada marcado
    execucoes_medidas = _marcar_o_guardrail(tabela.uso.get(leitor_de_documentos.CHAVE_DAS_EXECUCOES, []), alertas)
    # 3. Ninguém encontrado: recusa, com as dúvidas da IA na explicação (a IA trabalhou: as execuções seguem no erro)
    if not tabela.funcionarios:
        explicacao = "Não encontrei funcionários neste documento."
        if tabela.duvidas:
            explicacao += " " + " ".join(tabela.duvidas)
        raise DocumentoRecusado(explicacao, execucoes_medidas)
    # Uma coluna por jeito de a empresa chamar o dado (a tela mostra cada um numa linha, com o campo do banco)
    linhas, origem = tabela_pelos_rotulos(tabela)
    leitura = LeituraDoWord(linhas=linhas, como_foi_lido="texto corrido", duvidas=list(tabela.duvidas),
                            perguntas=list(tabela.perguntas), duvidas_gerais=list(tabela.duvidas_gerais),
                            origem_das_colunas=origem, alertas_guardrail=alertas)
    # A empresa sabe que foi a IA que montou a tabela e que deve conferir
    leitura.avisos.insert(0, f"Texto corrido: o Agente Leitor montou a lista com {len(tabela.funcionarios)} "
                             "funcionário(s). Você confere tudo antes de enviar ao banco.")
    if alertas:
        leitura.avisos.append(f"{len(alertas)} trecho(s) com cara de ordem para o Agente Leitor foram tirados pelo "
                              "Guardrail antes da leitura.")
    # O uso da IA vai para a auditoria: chamadas, modelos, tokens, custo e contagens (nada pessoal). Vão junto as
    # execuções medidas (chave "execucoes_dos_agentes"), que services/processamentos.py tira dali e grava nas execuções
    leitura.uso_da_ia = dict(tabela.uso)
    leitura.uso_da_ia["funcionarios"] = len(tabela.funcionarios)
    leitura.uso_da_ia["duvidas"] = len(tabela.duvidas)
    return leitura


# ============================== A leitura ==============================

def ler_word(conteudo: bytes, cliente=None, campos=None) -> LeituraDoWord:
    """Lê o Word da empresa: tabela, fichas ou texto corrido (nessa ordem). Devolve a LeituraDoWord.

    cliente: o LLM, usado só no texto corrido (sem informar, usa o do modo configurado; os testes passam um falso).
    campos: os campos do layout vigente; sem informar, a versão 1 do layout (data/contratos/layout_v1.csv).
    Levanta DocumentoRecusado com a mensagem para a empresa, se o documento não servir.
    """
    documento = _abrir(conteudo)
    paragrafos = _paragrafos(documento)
    # 1. Tabela
    linhas, avisos = _linhas_das_tabelas(_tabelas(documento))
    if linhas:
        avisos.insert(0, "Li a tabela do documento Word.")
        return LeituraDoWord(linhas=linhas, como_foi_lido="tabela", avisos=avisos)
    # Documento sem tabela e sem texto
    if not paragrafos:
        raise DocumentoRecusado("O documento Word está vazio.")
    # 2. Fichas "Rótulo: valor"
    if _documento_e_de_fichas(paragrafos):
        return LeituraDoWord(linhas=_linhas_das_fichas(paragrafos), como_foi_lido="fichas",
                             avisos=["Li o documento Word como fichas (\"Rótulo: valor\"), uma por funcionário."])
    # 3. Texto corrido, com a IA
    if campos is None:
        campos = carregar_layout()
    return _ler_texto_corrido(paragrafos, cliente, campos)
