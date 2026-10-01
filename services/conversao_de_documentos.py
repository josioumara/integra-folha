"""Conversão de documentos que não são .docx para o caminho que a aplicação já sabe ler (ADR-89). Sem IA.

Para que serve: o Word do LibreOffice (.odt) e o texto formatado (.rtf) chegam em formatos diferentes, mas o que
importa é o conteúdo: parágrafos e tabelas. Em vez de um leitor novo para cada formato, este arquivo converte:
    - .odt → um Word (.docx) montado em memória, com os mesmos parágrafos e tabelas, na mesma ordem; ele segue o
      caminho do Word (tabela e fichas por regra, texto corrido pelo Leitor de Documentos);
    - .rtf → texto simples, que segue o caminho do .txt (com colunas, como CSV; sem colunas, como um Word).
Assim, as regras de leitura e as conferências continuam num lugar só.
"""
import io

from docx import Document
from odf import table as tabela_odf
from odf import text as texto_odf
from odf.opendocument import load
from odf.teletype import extractText
from striprtf.striprtf import rtf_to_text

# O separador de célula que o striprtf escreve nas tabelas do RTF (ex.: "Maria|529|")
SEPARADOR_DAS_TABELAS_DO_RTF = "|"
# O atributo do ODF que diz quantas vezes uma célula se repete (célula vazia repetida é gravada uma vez só)
ATRIBUTO_DE_REPETICAO = "numbercolumnsrepeated"
# O tipo que um documento de texto do LibreOffice declara (uma planilha .ods declara outro)
TIPO_DO_DOCUMENTO_DE_TEXTO = "application/vnd.oasis.opendocument.text"
# Uma linha de tabela do .odt com mais células vazias repetidas que isto é o "resto" da página, não dado
MAXIMO_DE_REPETICOES_LIDAS = 50


class DocumentoIlegivel(ValueError):
    """O documento não pôde ser aberto (corrompido ou com a extensão trocada). A mensagem é para a empresa."""


def _celulas_da_linha(linha_odf) -> list[str]:
    """O texto de cada célula de uma linha de tabela do .odt, expandindo as células repetidas.

    Recebe: a linha (table:table-row). Devolve: os textos, na ordem. Exemplo: 3 células vazias gravadas como uma
    célula "repetida 3 vezes" viram ["", "", ""].
    """
    celulas = []
    for celula in linha_odf.getElementsByType(tabela_odf.TableCell):
        texto = extractText(celula).strip()
        # Quantas vezes a célula se repete (1 quando o atributo não existe)
        repeticoes = int(celula.getAttribute(ATRIBUTO_DE_REPETICAO) or 1)
        # Repetição enorme de célula vazia é o fim da linha na planilha, não dado: para aqui
        if repeticoes > MAXIMO_DE_REPETICOES_LIDAS and not texto:
            break
        for _ in range(repeticoes):
            celulas.append(texto)
    # Tira as células vazias do fim da linha
    while celulas and not celulas[-1]:
        celulas.pop()
    return celulas


def _acrescentar_tabela(documento, tabela) -> None:
    """Copia uma tabela do .odt para o Word em memória, linha por linha (linha sem nenhum texto fica de fora)."""
    linhas = []
    for linha_odf in tabela.getElementsByType(tabela_odf.TableRow):
        celulas = _celulas_da_linha(linha_odf)
        if any(celulas):
            linhas.append(celulas)
    if not linhas:
        return
    # A tabela do Word tem o número de colunas da linha mais comprida
    colunas = 0
    for linha in linhas:
        colunas = max(colunas, len(linha))
    tabela_do_word = documento.add_table(rows=len(linhas), cols=colunas)
    for numero_da_linha, celulas in enumerate(linhas):
        for numero_da_coluna, texto in enumerate(celulas):
            tabela_do_word.cell(numero_da_linha, numero_da_coluna).text = texto


def word_do_odt(conteudo: bytes) -> bytes:
    """Converte um .odt num .docx em memória, com os parágrafos e as tabelas na mesma ordem.

    Recebe: os bytes do .odt. Devolve: os bytes do .docx. Levanta DocumentoIlegivel se o arquivo não abre.
    Exemplo: um .odt com o parágrafo "Seguem os funcionários" e uma tabela Nome/CPF vira um Word com o mesmo parágrafo
    e a mesma tabela, que o leitor do Word lê por regra.
    """
    try:
        documento_odf = load(io.BytesIO(conteudo))
    except Exception as erro:
        raise DocumentoIlegivel("Não foi possível abrir o documento .odt: pode estar corrompido ou com a extensão "
                                "trocada.") from erro
    # O pacote abre, mas é outro tipo do LibreOffice (ex.: uma planilha .ods com o nome .odt)
    if documento_odf.mimetype != TIPO_DO_DOCUMENTO_DE_TEXTO:
        raise DocumentoIlegivel("O arquivo não é um documento de texto do LibreOffice (.odt). Se for uma planilha, "
                                "salve como .ods.")
    documento = Document()
    # O corpo do texto, na ordem em que aparece: parágrafos, títulos, listas e tabelas
    for elemento in documento_odf.text.childNodes:
        nome = getattr(elemento, "qname", (None, None))[1]
        if nome == "table":
            _acrescentar_tabela(documento, elemento)
        elif nome in ("p", "h", "list"):
            # Uma lista vira um parágrafo por item; parágrafo e título, um parágrafo
            itens = [elemento]
            if nome == "list":
                itens = elemento.getElementsByType(texto_odf.P)
            for item in itens:
                texto = extractText(item).strip()
                if texto:
                    documento.add_paragraph(texto)
    em_memoria = io.BytesIO()
    documento.save(em_memoria)
    return em_memoria.getvalue()


def texto_do_rtf(conteudo: bytes) -> str:
    """Converte um .rtf em texto simples. Uma linha de tabela vira as células separadas por "|".

    Recebe: os bytes do .rtf. Devolve: o texto. Levanta DocumentoIlegivel se o conteúdo não é RTF.
    Exemplo: a tabela "Maria | 529..." vira a linha "Maria|529..." (o "|" que sobra no fim da linha sai).
    Por que "cp1252": o RTF é texto ASCII com os acentos escritos como códigos (ex.: \\'e9 = é), na tabela do Windows.
    """
    texto_bruto = conteudo.decode("cp1252", errors="replace")
    if not texto_bruto.lstrip().startswith("{\\rtf"):
        raise DocumentoIlegivel("O arquivo não é um .rtf válido: pode estar com a extensão trocada.")
    texto = rtf_to_text(texto_bruto, encoding="cp1252", errors="replace")
    # A linha de tabela termina com um "|" a mais: sai, para a regra da tabela contar as colunas certas
    linhas = []
    for linha in texto.splitlines():
        if linha.endswith(SEPARADOR_DAS_TABELAS_DO_RTF):
            linha = linha[:-1]
        linhas.append(linha)
    return "\n".join(linhas)
