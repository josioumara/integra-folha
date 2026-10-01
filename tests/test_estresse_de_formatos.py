"""Teste de estresse dos formatos de arquivo aceitos no envio (leitura por regra, sem IA).

Para que serve este arquivo: a empresa manda a lista de funcionários do jeito que o sistema dela exporta. Aqui
geramos na hora (numa pasta temporária do teste) arquivos "difíceis", do tipo que aparece no mundo real, e
conferimos que a ingestão (services/ingestao.py) faz a coisa CERTA com cada um:
- aceita e lê direito (cabeçalho certo, células certas, sem juntar nem perder linhas); ou
- recusa com uma mensagem clara, que diz à empresa o que fazer.

O que é estressado: CSV com cada separador (; , | e tabulação); codificações (UTF-8, UTF-8 com BOM, Windows-1252,
Latin-1 e UTF-16); aspas, separador e quebra de linha dentro da célula; linhas e colunas vazias; cabeçalho fora da
primeira linha; duas linhas de cabeçalho; células mescladas e várias abas no Excel; fórmulas; datas como número de
série; números com vírgula decimal e ponto de milhar; CPF com e sem máscara e com zeros à esquerda perdidos; arquivo
vazio; só cabeçalho; extensão trocada; arquivo corrompido ou cortado; arquivo no limite de tamanho, de linhas e de
colunas; Word com tabela e Word em texto (só os caminhos sem IA).

Nenhum teste chama IA: planilha e CSV nunca usam IA, e no Word usamos um cliente que FALHA se for chamado (prova que
a leitura foi por regra). Todos os dados são fictícios; os CPFs são válidos só na conta do dígito verificador.

Os 18 defeitos que estes testes encontraram (antes marcados com xfail) já foram corrigidos. Se aparecer um defeito
novo, o teste dele pode ser marcado com xfail(strict=True) até a correção: "strict" faz a bateria avisar quando o
defeito for corrigido (aí é só tirar a marca).
"""
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pandas
import pytest
from docx import Document
from openpyxl import Workbook

from services import ingestao, normalizador
from services.documentos import cpf_valido
from tests.test_leitura_de_word import ClienteQueNaoPodeSerChamado

# CPFs fictícios (válidos só na conta do dígito verificador), com máscara
CPF_DA_ANA = "529.982.247-25"
CPF_DA_BIA = "111.444.777-35"
CPF_DO_CAIO = "123.456.789-09"
# CPFs fictícios que começam com zero: são estes que o Excel estraga quando guarda o CPF como número
CPF_COM_UM_ZERO = "012.345.678-90"
CPF_COM_DOIS_ZEROS = "001.234.567-97"

# O Excel antigo (.xls) de teste que já existe no projeto (usado aqui só para gerar a versão cortada)
ARQUIVO_XLS_DE_TESTE = Path(__file__).resolve().parent / "dados" / "lista_excel_antigo.xls"

# Um megabyte, em bytes (o limite de tamanho é dado em bytes)
UM_MEGABYTE = 1024 * 1024


# ============================== Ajudantes: gerar e ler os arquivos ==============================

def ler_do_disco(pasta: Path, nome_do_arquivo: str, conteudo: bytes, limite_em_bytes: int = 5 * UM_MEGABYTE):
    """Grava o arquivo na pasta do teste e lê de volta, como o envio faz. Devolve a Leitura da ingestão.

    Recebe: a pasta temporária, o nome (com a extensão que a empresa usou), os bytes e o limite de tamanho.
    Exemplo: ler_do_disco(tmp_path, "lista.csv", b"Nome;CPF\\nAna;52998224725\\n") → Leitura com 1 funcionário.
    """
    # Grava os bytes num arquivo de verdade (o nome importa: a extensão decide como ler)
    caminho = pasta / nome_do_arquivo
    caminho.write_bytes(conteudo)
    # Lê do disco e entrega para a ingestão, com um cliente de IA que falha se for chamado
    return ingestao.ler_arquivo(caminho.read_bytes(), nome_do_arquivo, limite_em_bytes,
                                cliente=ClienteQueNaoPodeSerChamado(modo="mock"))


def texto_do_csv(linhas_do_arquivo: list[str], fim_de_linha: str = "\n") -> str:
    """Junta as linhas de um CSV num texto só, com o fim de linha pedido (e um no final, como os sistemas gravam).

    Exemplo: texto_do_csv(["Nome;CPF", "Ana;1"]) → "Nome;CPF\\nAna;1\\n".
    """
    # Cada linha ganha o fim de linha no final
    return fim_de_linha.join(linhas_do_arquivo) + fim_de_linha


def bytes_da_planilha(livro: Workbook) -> bytes:
    """Os bytes de uma planilha .xlsx montada na memória com o openpyxl."""
    # Salva a planilha numa "folha em branco" na memória, em vez de num arquivo
    memoria = io.BytesIO()
    livro.save(memoria)
    # Devolve o que foi gravado
    return memoria.getvalue()


def trocar_valor_gravado_na_planilha(conteudo: bytes, valor_gravado: str, valor_novo: str) -> bytes:
    """Troca um número exatamente como está gravado dentro do .xlsx (no XML da primeira aba). Devolve os novos bytes.

    Por que existe: o openpyxl arredonda o número ao gravar, mas o Excel grava o valor com todos os dígitos (ex.:
    "1234.5600000000002"). Assim o teste recebe o mesmo arquivo que o Excel gravaria.
    Exemplo: trocar_valor_gravado_na_planilha(conteudo, "1234.56", "1234.5600000000002").
    """
    # O .xlsx é um pacote zip: abre o pacote original e prepara um pacote novo na memória
    pacote_original = zipfile.ZipFile(io.BytesIO(conteudo))
    memoria = io.BytesIO()
    pacote_novo = zipfile.ZipFile(memoria, "w", zipfile.ZIP_DEFLATED)
    # Copia cada arquivo do pacote; no XML da primeira aba, troca o valor
    for nome_interno in pacote_original.namelist():
        dados = pacote_original.read(nome_interno)
        if nome_interno == "xl/worksheets/sheet1.xml":
            dados = dados.replace(f"<v>{valor_gravado}</v>".encode("utf-8"), f"<v>{valor_novo}</v>".encode("utf-8"))
        pacote_novo.writestr(nome_interno, dados)
    # Fecha o pacote novo (grava o índice do zip) e devolve os bytes
    pacote_novo.close()
    return memoria.getvalue()


def bytes_do_word(paragrafos: list[str]) -> bytes:
    """Os bytes de um Word (.docx) só com parágrafos de texto, um por item da lista."""
    # Documento novo e vazio
    documento = Document()
    # Um parágrafo por texto pedido
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    # Salva na memória e devolve os bytes
    memoria = io.BytesIO()
    documento.save(memoria)
    return memoria.getvalue()


def perfil_pelo_nome(leitura) -> dict:
    """O retrato de cada coluna (tipo provável, amostras, aviso), encontrado pelo nome da coluna."""
    perfis_por_nome = {}
    # Um retrato por coluna, guardado pelo nome dela
    for perfil in ingestao.perfil_das_colunas(leitura):
        perfis_por_nome[perfil["nome"]] = perfil
    return perfis_por_nome


def primeira_coluna(leitura) -> list[str]:
    """O valor da primeira coluna de cada funcionário (em geral, o nome), para conferir quem foi lido."""
    valores = []
    # Pega a primeira célula de cada linha de funcionário
    for linha in leitura.linhas:
        valores.append(linha[0])
    return valores


def tem_aviso_com(leitura, trecho: str) -> bool:
    """True se algum aviso da leitura contém o trecho (ex.: "abas")."""
    # Procura o trecho em cada aviso
    for aviso in leitura.avisos:
        if trecho in aviso:
            return True
    return False


# ============================== 1. CSV: separadores ==============================

@pytest.mark.parametrize("separador", [";", ",", "|", "\t"], ids=["ponto_e_virgula", "virgula", "barra", "tabulacao"])
def test_csv_com_cada_separador_e_lido_certo(tmp_path, separador):
    """O separador é descoberto sozinho: ; , | ou tabulação dão a mesma leitura."""
    # As mesmas linhas, montadas com o separador do caso
    linhas_da_tabela = [["Nome", "CPF", "Cargo", "Salário"],
                        ["Ana Souza", CPF_DA_ANA, "Analista", "3150.00"],
                        ["Bia Lima", CPF_DA_BIA, "Vendedora", "2800.50"]]
    linhas_do_arquivo = []
    # Cada linha da tabela vira uma linha de texto, com as células unidas pelo separador
    for linha in linhas_da_tabela:
        linhas_do_arquivo.append(separador.join(linha))
    # Grava o texto em UTF-8
    conteudo = texto_do_csv(linhas_do_arquivo).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # O separador certo foi descoberto, e as células chegaram inteiras
    assert leitura.separador == separador
    assert leitura.cabecalhos == linhas_da_tabela[0]
    assert leitura.linhas == linhas_da_tabela[1:]


def test_csv_com_ponto_e_virgula_e_virgula_decimal_nas_celulas(tmp_path):
    """Separador ";" com muitas vírgulas nos valores (dinheiro "3.150,00"): a vírgula não pode virar separador."""
    # Um CSV com ";" e duas vírgulas por linha de pessoa (a vírgula é o decimal)
    conteudo = texto_do_csv(["Nome;Salário;Adicional",
                             "Ana Souza;3.150,00;120,50",
                             "Bia Lima;2.800,00;80,00",
                             "Caio Reis;12.500,75;0,00"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # O ";" é o separador; as vírgulas ficam dentro das células
    assert leitura.separador == ";"
    assert leitura.linhas[0] == ["Ana Souza", "3.150,00", "120,50"]


@pytest.mark.parametrize("fim_de_linha", [
    pytest.param("\r\n", id="windows"),
    pytest.param("\n", id="linux"),
    # O fim de linha só com "\r" (Mac antigo e alguns sistemas legados)
    pytest.param("\r", id="mac_antigo"),
])
def test_csv_com_cada_tipo_de_fim_de_linha(tmp_path, fim_de_linha):
    """Windows (\\r\\n), Mac antigo (\\r) e Linux (\\n): as linhas são separadas do mesmo jeito."""
    # O mesmo CSV, com o fim de linha do caso
    conteudo = texto_do_csv(["Nome;CPF", f"Ana Souza;{CPF_DA_ANA}", f"Bia Lima;{CPF_DA_BIA}"],
                            fim_de_linha).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # Duas pessoas, sem sobrar "\r" no fim da última célula
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA], ["Bia Lima", CPF_DA_BIA]]


def test_csv_com_a_dica_sep_do_excel_na_primeira_linha(tmp_path):
    """O Excel às vezes grava "sep=;" na primeira linha: ela não é cabeçalho nem funcionário."""
    # A dica "sep=;" na linha 1, o cabeçalho na linha 2 e uma pessoa
    conteudo = texto_do_csv(["sep=;", "Nome;CPF;Cargo", f"Ana Souza;{CPF_DA_ANA};Analista"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # O cabeçalho é a linha 2, e a dica não virou funcionário
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo"]
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA, "Analista"]]


def test_extensao_em_maiusculas_e_aceita(tmp_path):
    """"LISTA.CSV" é o mesmo formato que "lista.csv"."""
    # Um CSV comum, com o nome todo em maiúsculas
    conteudo = texto_do_csv(["Nome;CPF", f"Ana Souza;{CPF_DA_ANA}"]).encode("utf-8")
    leitura = ler_do_disco(tmp_path, "LISTA.CSV", conteudo)
    # Foi lido como CSV, com a pessoa
    assert leitura.formato == "csv"
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA]]


# ============================== 2. CSV: codificações ==============================

@pytest.mark.parametrize("codificacao_do_arquivo, codificacao_esperada", [
    ("utf-8", "utf-8"),
    ("utf-8-sig", "utf-8"),      # UTF-8 com a marca BOM no começo (o "CSV UTF-8" do Excel)
    ("cp1252", "cp1252"),        # Windows-1252, o padrão do Excel em português
    ("latin-1", "cp1252"),       # ISO-8859-1: para os acentos do português, os mesmos bytes do Windows-1252
], ids=["utf8", "utf8_com_bom", "windows_1252", "latin1"])
def test_csv_com_acentos_em_cada_codificacao(tmp_path, codificacao_do_arquivo, codificacao_esperada):
    """Os acentos chegam certos em qualquer codificação comum, e o BOM não gruda no nome da primeira coluna."""
    linhas_da_tabela = [["Nome", "Função", "Endereço"],
                        ["Maria Conceição", "Técnica de manutenção", "Praça da Sé, 100"],
                        ["João Araújo", "Açougueiro", "Rua São João, 7"]]
    linhas_do_arquivo = []
    # Cada linha da tabela vira uma linha de texto separada por ";"
    for linha in linhas_da_tabela:
        linhas_do_arquivo.append(";".join(linha))
    # Grava o texto na codificação do caso
    conteudo = texto_do_csv(linhas_do_arquivo).encode(codificacao_do_arquivo)
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # A primeira coluna é "Nome", sem a marca BOM (o caractere invisível "\ufeff") colada na frente
    assert leitura.cabecalhos == linhas_da_tabela[0]
    # Os acentos das células chegaram certos
    assert leitura.linhas == linhas_da_tabela[1:]
    # A codificação foi descoberta sozinha
    assert leitura.codificacao == codificacao_esperada


def test_txt_em_utf16_do_excel_e_lido(tmp_path):
    """O Excel salva "Texto Unicode (.txt)" em UTF-16 com tabulação: é uma exportação comum e deve ser lida."""
    # Tabulação, fim de linha do Windows e UTF-16 (cada letra ocupa 2 bytes, e muitos deles são o byte zero)
    conteudo = texto_do_csv(["Nome\tCPF\tFunção", f"Maria Conceição\t{CPF_DA_ANA}\tAnalista"],
                            "\r\n").encode("utf-16")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.txt", conteudo)
    # Certo: as colunas e os acentos chegam como na planilha
    assert leitura.cabecalhos == ["Nome", "CPF", "Função"]
    assert leitura.linhas == [["Maria Conceição", CPF_DA_ANA, "Analista"]]


def test_txt_separado_por_tabulacao_em_windows_1252(tmp_path):
    """O "Texto (separado por tabulações)" do Excel: .txt, tabulação e Windows-1252."""
    # Tabulação, fim de linha do Windows e Windows-1252
    conteudo = texto_do_csv(["Nome\tCPF\tFunção", f"João Araújo\t{CPF_DA_BIA}\tAçougueiro"], "\r\n").encode("cp1252")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.txt", conteudo)
    # Lido como tabela, com os acentos certos
    assert leitura.formato == "txt"
    assert leitura.cabecalhos == ["Nome", "CPF", "Função"]
    assert leitura.linhas == [["João Araújo", CPF_DA_BIA, "Açougueiro"]]


# ============================== 3. CSV: aspas e conteúdo dentro da célula ==============================

def test_aspas_separador_e_quebra_de_linha_dentro_da_celula(tmp_path):
    """Entre aspas, o ";" e a quebra de linha são parte da célula; aspas dobradas ("") viram uma aspa."""
    # Pessoa 1: nome com ";" e endereço com quebra de linha; pessoa 2: apelido entre aspas dobradas
    conteudo = texto_do_csv(["Nome;Endereço;CPF",
                             f"\"Souza; Maria\";\"Rua A, 100\nApto 2\";{CPF_DA_ANA}",
                             f"\"Ana \"\"Aninha\"\" Lima\";Rua B;{CPF_DA_BIA}"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # Duas pessoas (a quebra de linha dentro da célula não cria uma terceira)
    assert leitura.linhas == [["Souza; Maria", "Rua A, 100\nApto 2", CPF_DA_ANA],
                              ["Ana \"Aninha\" Lima", "Rua B", CPF_DA_BIA]]
    # Cada pessoa é um registro: a Ana está na linha 3, como o Excel mostra (a célula com quebra fica numa linha só)
    assert leitura.numeros_linha == [2, 3]


def test_csv_com_virgula_e_valores_americanos_entre_aspas(tmp_path):
    """Separador "," com o salário no jeito americano entre aspas ("3,150.00"): a vírgula entre aspas fica."""
    # Nome e salário com vírgula, protegidos por aspas
    conteudo = texto_do_csv(["Name,CPF,Salary",
                             f"\"Souza, Maria\",{CPF_DA_ANA},\"3,150.00\"",
                             f"Bia Lima,{CPF_DA_BIA},\"12,500.50\""]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # A vírgula é o separador, e as vírgulas entre aspas ficaram dentro das células
    assert leitura.separador == ","
    assert leitura.linhas == [["Souza, Maria", CPF_DA_ANA, "3,150.00"], ["Bia Lima", CPF_DA_BIA, "12,500.50"]]


def test_aspas_abertas_sem_fechar_nao_engolem_as_outras_linhas(tmp_path):
    """Uma aspa esquecida no começo de uma célula não pode fazer as pessoas seguintes sumirem em silêncio.

    Certo: ler as 3 pessoas, ou recusar dizendo que há aspas sem fechar. Errado: 1 "pessoa" com o resto do arquivo.
    """
    # A aspa no começo do nome da Ana nunca é fechada
    conteudo = texto_do_csv(["Nome;CPF;Cargo",
                             f"\"Ana Souza;{CPF_DA_ANA};Analista",
                             f"Bia Lima;{CPF_DA_BIA};Vendedora",
                             f"Caio Reis;{CPF_DO_CAIO};Gerente"]).encode("utf-8")
    # Tenta ler: uma recusa com mensagem também é um resultado certo
    try:
        leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    except ingestao.ArquivoRecusado as recusa:
        # Recusar também é certo, desde que a mensagem aponte as aspas
        assert "aspas" in str(recusa)
        return
    # Se aceitou, as três pessoas precisam estar lá
    assert len(leitura.linhas) == 3


def test_aspas_sem_fechar_em_arquivo_grande_da_uma_mensagem_clara(tmp_path):
    """A mesma aspa esquecida, mas com 5.000 pessoas depois: o resto do arquivo vira uma célula de mais de 128 KB.

    Certo: ler as pessoas ou recusar com uma mensagem para a empresa (ArquivoRecusado), nunca um erro técnico.
    """
    linhas_do_arquivo = ["Nome;CPF;Cargo", f"\"Ana Souza;{CPF_DA_ANA};Analista"]
    # Muitas pessoas depois da aspa esquecida
    for numero_da_pessoa in range(1, 5001):
        linhas_do_arquivo.append(f"Pessoa {numero_da_pessoa};{CPF_DA_BIA};Vendedora")
    conteudo = texto_do_csv(linhas_do_arquivo).encode("utf-8")
    # Tenta ler: uma recusa com mensagem também é um resultado certo (um erro técnico, não)
    try:
        leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    except ingestao.ArquivoRecusado as recusa:
        # Recusar é certo, desde que a mensagem aponte as aspas
        assert "aspas" in str(recusa)
        return
    # Se aceitou, todas as pessoas precisam estar lá
    assert len(leitura.linhas) == 5001


# ============================== 4. Linhas e colunas vazias ==============================

def test_linhas_vazias_e_linhas_so_com_separadores_sao_puladas(tmp_path):
    """Linha em branco ou só com ";;;" no meio e no fim não vira funcionário; o número da linha continua certo."""
    # Duas pessoas, com linhas vazias e linhas só de separadores entre elas e no fim
    conteudo = texto_do_csv(["Nome;CPF;Cargo",
                             f"Ana Souza;{CPF_DA_ANA};Analista",
                             "",
                             ";;",
                             f"Bia Lima;{CPF_DA_BIA};Vendedora",
                             "", "", ";;"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # Só as duas pessoas viraram funcionários
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]
    # A Bia está na linha 5 do arquivo (contando como no Excel)
    assert leitura.numeros_linha == [2, 5]


def test_coluna_com_nome_e_sem_valores_continua_na_leitura(tmp_path):
    """Uma coluna com nome e toda vazia (ex.: "Complemento") é real: fica, com o retrato "VAZIA"."""
    # A coluna "Complemento" tem nome, mas nenhuma pessoa a preencheu
    conteudo = texto_do_csv(["Nome;Complemento;CPF",
                             f"Ana Souza;;{CPF_DA_ANA}",
                             f"Bia Lima;;{CPF_DA_BIA}"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # A coluna continua lá, e o retrato diz que ela está vazia
    assert leitura.cabecalhos == ["Nome", "Complemento", "CPF"]
    assert perfil_pelo_nome(leitura)["Complemento"]["tipo_provavel"] == "VAZIA"


def test_separadores_sobrando_no_fim_nao_criam_colunas(tmp_path):
    """Exportações do Excel costumam gravar "Nome;CPF;;;": colunas sem nome e sem nenhum valor não são colunas."""
    # Três separadores sobrando no fim de cada linha
    conteudo = texto_do_csv(["Nome;CPF;;;",
                             f"Ana Souza;{CPF_DA_ANA};;;",
                             f"Bia Lima;{CPF_DA_BIA};;;"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # Certo: só as duas colunas de verdade
    assert leitura.cabecalhos == ["Nome", "CPF"]


# ============================== 5. Cabeçalho fora do lugar ==============================

def test_csv_com_titulo_data_e_linha_em_branco_antes_do_cabecalho(tmp_path):
    """Relatório com título e data antes da tabela: o cabeçalho é achado na linha 4, e a empresa é avisada."""
    # Título, data de emissão e uma linha em branco antes do cabeçalho
    conteudo = texto_do_csv(["Relatório de funcionários",
                             "Emitido em 01/09/2026",
                             "",
                             "Nome;CPF;Cargo;Salário",
                             f"Ana Souza;{CPF_DA_ANA};Analista;3.150,00",
                             f"Bia Lima;{CPF_DA_BIA};Vendedora;2.800,00"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "relatorio.csv", conteudo)
    # O cabeçalho certo, as pessoas nas linhas 5 e 6 e o aviso de onde o cabeçalho estava
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo", "Salário"]
    assert leitura.numeros_linha == [5, 6]
    assert tem_aviso_com(leitura, "cabeçalho foi encontrado na linha 4")


def test_csv_com_linha_de_dados_da_empresa_antes_do_cabecalho(tmp_path):
    """Relatórios de ERP trazem "Empresa / CNPJ / Emitido em" numa linha acima da tabela: não é o cabeçalho."""
    # Linha 1: três células de texto com os dados da empresa; linha 2: o cabeçalho de verdade
    conteudo = texto_do_csv(["Empresa: Aurora Ltda;CNPJ 11.222.333/0001-81;Emitido em 01/09/2026",
                             "Nome;CPF;Cargo;Salário",
                             f"Ana Souza;{CPF_DA_ANA};Analista;3.150,00",
                             f"Bia Lima;{CPF_DA_BIA};Vendedora;2.800,00"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "relatorio.csv", conteudo)
    # Certo: o cabeçalho é a linha 2
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo", "Salário"]
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


def test_excel_com_titulo_mesclado_acima_do_cabecalho(tmp_path):
    """Planilha com o título mesclado de A1 a D1 e uma linha em branco: o cabeçalho de verdade está na linha 3."""
    # Planilha nova, na primeira aba
    livro = Workbook()
    aba = livro.active
    # Linha 1: o título, mesclado em cima das 4 colunas (o jeito mais comum de "enfeitar" a planilha)
    aba.append(["Funcionários admitidos em setembro"])
    aba.merge_cells("A1:D1")
    # Linha 2 em branco; linha 3 é o cabeçalho; depois, as pessoas
    aba.append([])
    aba.append(["Nome", "CPF", "Cargo", "Salário"])
    aba.append(["Ana Souza", CPF_DA_ANA, "Analista", 3150])
    aba.append(["Bia Lima", CPF_DA_BIA, "Vendedora", 2800])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # Certo: o cabeçalho da linha 3 e as duas pessoas
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo", "Salário"]
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


def test_excel_com_a_tabela_comecando_na_celula_c5(tmp_path):
    """A tabela começa na célula C5 (linhas 1 a 4 e colunas A e B vazias): só as colunas da tabela contam."""
    # Planilha nova, na primeira aba
    livro = Workbook()
    aba = livro.active
    # Cabeçalho e pessoas escritos a partir da coluna C (3), linha 5
    linhas_da_tabela = [["Nome", "CPF", "Cargo"],
                        ["Ana Souza", CPF_DA_ANA, "Analista"],
                        ["Bia Lima", CPF_DA_BIA, "Vendedora"]]
    for numero_da_linha, linha in enumerate(linhas_da_tabela, start=5):
        for numero_da_coluna, valor in enumerate(linha, start=3):
            # Escreve cada valor na sua célula (linha 5 em diante, coluna C em diante)
            aba.cell(row=numero_da_linha, column=numero_da_coluna, value=valor)
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # Certo: só as 3 colunas da tabela
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo"]
    assert leitura.linhas == linhas_da_tabela[1:]


# ============================== 6. Duas linhas de cabeçalho ==============================

def test_csv_com_linha_de_grupos_acima_do_cabecalho(tmp_path):
    """A linha de grupos ("Dados pessoais", "Contrato") está incompleta: o cabeçalho é a linha de baixo."""
    # Linha 1: os grupos, com células vazias entre eles; linha 2: os nomes das colunas
    conteudo = texto_do_csv(["Dados pessoais;;Contrato;",
                             "Nome;CPF;Cargo;Salário",
                             f"Ana Souza;{CPF_DA_ANA};Analista;3.150,00"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # O cabeçalho é a linha 2, e só a Ana é funcionária
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo", "Salário"]
    assert primeira_coluna(leitura) == ["Ana Souza"]


def test_excel_com_duas_linhas_de_cabecalho_e_grupos_mesclados(tmp_path):
    """Linha 1 com grupos mesclados (A1:B1 "Dados pessoais", C1:D1 "Contrato") e linha 2 com os nomes das colunas.

    Certo: cada coluna tem o nome da linha 2 (sozinho ou junto do grupo, ex.: "Dados pessoais - CPF"), e a linha
    "Nome, CPF, Cargo, Salário" nunca vira funcionário.
    """
    # Planilha nova, na primeira aba
    livro = Workbook()
    aba = livro.active
    # Linha 1: os dois grupos, cada um mesclado sobre as suas duas colunas
    aba.append(["Dados pessoais", None, "Contrato", None])
    aba.merge_cells("A1:B1")
    aba.merge_cells("C1:D1")
    # Linha 2: os nomes das colunas; depois, as pessoas
    aba.append(["Nome", "CPF", "Cargo", "Salário"])
    aba.append(["Ana Souza", CPF_DA_ANA, "Analista", 3150])
    aba.append(["Bia Lima", CPF_DA_BIA, "Vendedora", 2800])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # O nome de cada coluna vem da linha 2
    assert "Nome" in leitura.cabecalhos[0]
    assert "CPF" in leitura.cabecalhos[1]
    # Só as duas pessoas são funcionários
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


# ============================== 7. Excel: abas, mesclagem e fórmulas ==============================

def test_excel_com_a_lista_na_primeira_aba_e_outra_aba_depois(tmp_path):
    """A lista está na primeira aba; a aba "Observações" fica de fora, e a empresa é avisada."""
    # A primeira aba, com a lista
    livro = Workbook()
    aba = livro.active
    aba.title = "Funcionarios"
    aba.append(["Nome", "CPF"])
    aba.append(["Ana Souza", CPF_DA_ANA])
    # Uma segunda aba com anotações
    outra_aba = livro.create_sheet("Observacoes")
    outra_aba.append(["Conferir o CPF da Ana com o RH"])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # A lista da primeira aba foi lida, com o aviso da outra aba
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA]]
    assert "A planilha tem 2 abas; só a aba Funcionarios foi lida." in leitura.avisos


def test_excel_com_aba_de_resumo_antes_da_lista(tmp_path):
    """A aba "Resumo" (totais) vem antes da aba com os funcionários: a lida deve ser a da lista (ou recusar e dizer)."""
    # A primeira aba é um resumo com totais
    livro = Workbook()
    resumo = livro.active
    resumo.title = "Resumo"
    resumo.append(["Resumo da folha de setembro"])
    resumo.append(["Funcionários", 2])
    resumo.append(["Emitido em", datetime(2026, 9, 1)])
    # A segunda aba tem a lista de funcionários
    lista = livro.create_sheet("Funcionarios")
    lista.append(["Nome", "CPF", "Cargo"])
    lista.append(["Ana Souza", CPF_DA_ANA, "Analista"])
    lista.append(["Bia Lima", CPF_DA_BIA, "Vendedora"])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # Certo: a aba da lista
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo"]
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


def test_excel_com_celula_mesclada_na_vertical_e_na_horizontal(tmp_path):
    """Unidade mesclada em 2 linhas e endereço mesclado em 2 colunas: o valor vale para cada célula do grupo."""
    # O cabeçalho e as duas pessoas (as células vazias são as que a mesclagem cobre)
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF", "Unidade", "Endereço", "Complemento"])
    aba.append(["Ana Souza", CPF_DA_ANA, "Fábrica Campinas", "Rua A, 100", None])
    aba.append(["Bia Lima", CPF_DA_BIA, None, "Rua B, 200", "Casa"])
    # A unidade vale para as duas pessoas; o endereço da Ana ocupa também a coluna do complemento
    aba.merge_cells("C2:C3")
    aba.merge_cells("D2:E2")
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # O valor de cada grupo aparece em todas as células dele, e a empresa é avisada
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA, "Fábrica Campinas", "Rua A, 100", "Rua A, 100"],
                              ["Bia Lima", CPF_DA_BIA, "Fábrica Campinas", "Rua B, 200", "Casa"]]
    assert tem_aviso_com(leitura, "células mescladas")


def test_excel_com_formula_na_aba_lida_e_recusado_com_o_caminho(tmp_path):
    """Fórmula na lista: o valor guardado pode estar velho; a mensagem diz como salvar só os valores."""
    # O salário com reajuste é uma fórmula na célula D2
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF", "Salário", "Salário com reajuste"])
    aba.append(["Ana Souza", CPF_DA_ANA, 3000, "=C2*1.1"])
    # Recusa dizendo onde está a fórmula e como salvar só os valores
    with pytest.raises(ingestao.ArquivoRecusado, match=r"fórmulas.*D2.*Valores"):
        ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))


def test_excel_com_formula_so_em_outra_aba_e_aceito(tmp_path):
    """A fórmula está numa aba que não é lida (cálculos): a lista é aceita."""
    # A lista, na primeira aba, sem fórmulas
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF", "Salário"])
    aba.append(["Ana Souza", CPF_DA_ANA, 3000])
    # A fórmula fica numa segunda aba, de cálculos
    calculos = livro.create_sheet("Calculos")
    calculos.append(["Total", "=SUM(Funcionarios!C2:C2)"])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # A lista foi aceita e lida
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA, "3000"]]


def test_excel_com_formatacao_esquecida_la_embaixo_nao_conta_como_linhas(tmp_path):
    """Alguém formatou a coluna inteira: o Excel guarda células vazias até a linha 30.000. Só 2 pessoas contam."""
    # A lista com duas pessoas
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF"])
    aba.append(["Ana Souza", CPF_DA_ANA])
    aba.append(["Bia Lima", CPF_DA_BIA])
    # Uma célula vazia, só com formato de texto, lá na linha 30.000
    aba.cell(row=30_000, column=1).number_format = "@"
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # Certo: as duas pessoas
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


# ============================== 8. Excel: datas e números ==============================

def test_excel_com_datas_de_verdade_e_data_como_numero_de_serie(tmp_path):
    """Data formatada vira AAAA-MM-DD (com aviso); data como número de série sem formato (45352) chega como está e
    o Normalizador a converte (45352 = 01/03/2024)."""
    # Uma data, uma data com hora e um número de série de data sem o formato de data
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "Admissão", "Último ponto", "Nascimento (sem formato)"])
    aba.append(["Ana Souza", date(2024, 3, 1), datetime(2024, 3, 1, 8, 30), 45352])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # As datas viram AAAA-MM-DD (com o aviso); o número de série chega como número
    assert leitura.linhas == [["Ana Souza", "2024-03-01", "2024-03-01 08:30:00", "45352"]]
    assert tem_aviso_com(leitura, "AAAA-MM-DD")
    # O número de série vira a data certa no Normalizador
    assert normalizador.converter_data("45352", "DMY") == date(2024, 3, 1)


def test_excel_com_numeros_inteiros_decimais_e_cpf_como_numero(tmp_path):
    """Número vira o texto do número; CPF gravado como número sem perder zeros chega inteiro."""
    # CPF, salário e dependentes gravados como número
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF", "Salário", "Dependentes"])
    aba.append(["Ana Souza", 52998224725, 3150.5, 2])
    aba.append(["Bia Lima", 11144477735, 2800.0, 0])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # Cada número vira o seu texto ("2800.0" vira "2800", sem casas inventadas)
    assert leitura.linhas == [["Ana Souza", "52998224725", "3150.5", "2"], ["Bia Lima", "11144477735", "2800", "0"]]


def test_excel_com_ruido_de_ponto_flutuante_no_salario(tmp_path):
    """Valor colado de uma conta do Excel: o arquivo guarda 1234.5600000000002, o Excel mostra 1234,56.

    Certo: ler como o Excel mostra (até 15 dígitos significativos), senão o salário vira pendência de "mais de duas
    casas decimais" sem a empresa ter errado nada.
    """
    # A lista com o salário de 1234,56
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "Salário"])
    aba.append(["Ana Souza", 1234.56])
    # Grava o salário com o ruído, do jeito que o Excel gravaria
    conteudo = trocar_valor_gravado_na_planilha(bytes_da_planilha(livro), "1234.56", "1234.5600000000002")
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", conteudo)
    # Certo: o salário que a empresa vê no Excel
    assert leitura.linhas == [["Ana Souza", "1234.56"]]


def test_excel_com_cpf_como_numero_que_perdeu_os_zeros(tmp_path):
    """CPF que começa com zero, gravado como número: chega sem os zeros, o retrato avisa e o Normalizador recompõe
    (só porque o dígito verificador confirma)."""
    # Planilha nova, na primeira aba
    livro = Workbook()
    aba = livro.active
    aba.append(["Nome", "CPF"])
    # 012.345.678-90 e 001.234.567-97 gravados como número: o Excel guarda 1234567890 e 123456797
    aba.append(["Ana Souza", 1234567890])
    aba.append(["Bia Lima", 123456797])
    # Lê a planilha
    leitura = ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))
    # A leitura não inventa zeros: o texto é o que o Excel guardou
    assert leitura.linhas == [["Ana Souza", "1234567890"], ["Bia Lima", "123456797"]]
    # O retrato da coluna avisa que zeros podem ter sumido
    assert "zeros à esquerda" in perfil_pelo_nome(leitura)["CPF"]["aviso"]
    # O Normalizador recompõe os dois, porque o dígito verificador confere
    assert normalizador.converter_documento("1234567890", 11, cpf_valido, True) == ("01234567890", True)
    assert normalizador.converter_documento("123456797", 11, cpf_valido, True) == ("00123456797", True)


# ============================== 9. CSV: números e CPF ==============================

def test_csv_com_virgula_decimal_e_ponto_de_milhar(tmp_path):
    """Dinheiro no jeito brasileiro chega como está; o retrato diz "dinheiro" e o Normalizador converte certo."""
    # Ponto de milhar, "R$", milhões e uma casa decimal só
    valores_do_arquivo = ["3.150,00", "R$ 12.500,50", "1.234.567,89", "980,5"]
    linhas_do_arquivo = ["Nome;Salário"]
    # Uma pessoa por valor
    for posicao, valor in enumerate(valores_do_arquivo, start=1):
        linhas_do_arquivo.append(f"Pessoa {posicao};{valor}")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", texto_do_csv(linhas_do_arquivo).encode("utf-8"))
    # A leitura não mexe no número: cada valor chega exatamente como estava no arquivo
    valores_lidos = []
    for linha in leitura.linhas:
        # O salário é a segunda coluna
        valores_lidos.append(linha[1])
    assert valores_lidos == valores_do_arquivo
    # O retrato reconhece a coluna como dinheiro
    assert perfil_pelo_nome(leitura)["Salário"]["tipo_provavel"] == "DECIMAL_MONETARIO"
    # O Normalizador descobre que a vírgula é o decimal e converte cada valor
    convencao = normalizador.convencao_decimal(valores_lidos)
    assert convencao == "virgula"
    valores_convertidos = []
    for valor in valores_lidos:
        # Converte um valor de cada vez, com a convenção da coluna
        valores_convertidos.append(normalizador.converter_decimal(valor, convencao))
    # Os valores certos, com duas casas
    assert valores_convertidos == [Decimal("3150.00"), Decimal("12500.50"), Decimal("1234567.89"), Decimal("980.50")]


def test_csv_com_cpf_com_e_sem_mascara(tmp_path):
    """CPF com máscara, sem máscara e com espaços sobrando: chega como está (sem os espaços das pontas) e a coluna é
    reconhecida como CPF."""
    # Com máscara, sem máscara, com espaços nas pontas e com zero na frente (preservado, porque CSV é texto)
    conteudo = texto_do_csv(["Nome;CPF",
                             f"Ana Souza;{CPF_DA_ANA}",
                             "Bia Lima;11144477735",
                             f"Caio Reis;  {CPF_DO_CAIO}  ",
                             f"Duda Alves;{CPF_COM_UM_ZERO}"]).encode("utf-8")
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    cpfs_lidos = []
    for linha in leitura.linhas:
        # O CPF é a segunda coluna
        cpfs_lidos.append(linha[1])
    # Cada CPF chega como estava (só sem os espaços das pontas)
    assert cpfs_lidos == [CPF_DA_ANA, "11144477735", CPF_DO_CAIO, CPF_COM_UM_ZERO]
    # A coluna é reconhecida como CPF
    assert perfil_pelo_nome(leitura)["CPF"]["tipo_provavel"] == "CPF"


def test_csv_com_cpf_que_perdeu_os_zeros_e_sinalizado(tmp_path):
    """Muito comum: o CSV foi aberto e salvo no Excel, e os CPFs que começam com zero perderam os zeros.

    Certo: a coluna é reconhecida como CPF (os zeros recolocados passam no dígito verificador) ou o retrato avisa
    que zeros à esquerda podem ter sumido, como já acontece com a planilha .xlsx.
    """
    # A Bia (012.345.678-90) e o Caio (001.234.567-97) perderam os zeros da frente
    conteudo = texto_do_csv(["Nome;CPF",
                             "Ana Souza;52998224725",
                             "Bia Lima;1234567890",
                             "Caio Reis;123456797",
                             "Duda Alves;11144477735"]).encode("utf-8")
    # Lê o arquivo e pega o retrato da coluna do CPF
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    perfil_do_cpf = perfil_pelo_nome(leitura)["CPF"]
    # Certo: CPF reconhecido ou aviso de zeros
    assert perfil_do_cpf["tipo_provavel"] == "CPF" or "zeros" in perfil_do_cpf.get("aviso", "")


# ============================== 10. Vazio e só cabeçalho ==============================

@pytest.mark.parametrize("nome_do_arquivo, conteudo, trecho_da_mensagem", [
    ("lista.csv", b"", "vazio"),
    ("lista.csv", b"\n\n   \n\t\n", "nenhuma linha preenchida"),
    ("lista.csv", "\ufeff".encode("utf-8"), "nenhuma linha preenchida"),
    ("lista.csv", b";;;\n;;;\n", "nenhuma linha preenchida"),
    ("lista.csv", b"Nome;CPF;Cargo\n\n;;\n\n", "nenhuma linha de funcion"),
], ids=["zero_bytes", "so_linhas_em_branco", "so_o_bom", "so_separadores", "cabecalho_e_linhas_vazias"])
def test_csv_vazio_ou_so_com_cabecalho_e_recusado_com_o_motivo(tmp_path, nome_do_arquivo, conteudo,
                                                              trecho_da_mensagem):
    """Arquivo sem funcionário é recusado, e a mensagem diz o porquê."""
    # Recusa com a mensagem do caso
    with pytest.raises(ingestao.ArquivoRecusado, match=trecho_da_mensagem):
        ler_do_disco(tmp_path, nome_do_arquivo, conteudo)


def test_excel_vazio_e_recusado(tmp_path):
    """Planilha sem nenhuma célula preenchida (só a aba em branco)."""
    # Uma planilha nova, sem nada escrito, é recusada
    with pytest.raises(ingestao.ArquivoRecusado, match="nenhuma linha preenchida"):
        ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(Workbook()))


def test_excel_so_com_cabecalho_e_recusado(tmp_path):
    """Planilha só com a linha dos nomes das colunas: não tem funcionário."""
    # Só o cabeçalho
    livro = Workbook()
    livro.active.append(["Nome", "CPF", "Cargo"])
    # Recusa dizendo que falta funcionário
    with pytest.raises(ingestao.ArquivoRecusado, match="nenhuma linha de funcionário"):
        ler_do_disco(tmp_path, "lista.xlsx", bytes_da_planilha(livro))


# ============================== 11. Extensão trocada e arquivo corrompido ==============================

def test_csv_salvo_com_extensao_xlsx_e_recusado_com_mensagem_clara(tmp_path):
    """Um CSV renomeado para .xlsx (o próprio Excel não abre): recusa dizendo que a extensão pode estar trocada."""
    # Um CSV comum...
    conteudo = texto_do_csv(["Nome;CPF", f"Ana Souza;{CPF_DA_ANA}"]).encode("utf-8")
    # ...enviado com o nome .xlsx
    with pytest.raises(ingestao.ArquivoRecusado, match="extensão trocada"):
        ler_do_disco(tmp_path, "lista.xlsx", conteudo)


def test_xlsx_salvo_com_extensao_csv_e_recusado_com_mensagem_clara(tmp_path):
    """Uma planilha .xlsx renomeada para .csv: o conteúdo é binário, não texto."""
    # Uma planilha .xlsx comum...
    livro = Workbook()
    livro.active.append(["Nome", "CPF"])
    livro.active.append(["Ana Souza", CPF_DA_ANA])
    # ...enviada com o nome .csv
    with pytest.raises(ingestao.ArquivoRecusado, match="extensão trocada"):
        ler_do_disco(tmp_path, "lista.csv", bytes_da_planilha(livro))


def test_word_salvo_com_extensao_xlsx_e_recusado(tmp_path):
    """Um Word renomeado para .xlsx: os dois são zip, mas a planilha não abre; recusa em vez de quebrar."""
    # Um Word enviado com o nome .xlsx
    with pytest.raises(ingestao.ArquivoRecusado, match="Não foi possível abrir a planilha"):
        ler_do_disco(tmp_path, "lista.xlsx", bytes_do_word(["Nome: Ana Souza", f"CPF: {CPF_DA_ANA}"]))


def test_xls_que_e_uma_tabela_html_exportada_por_sistema(tmp_path):
    """Muitos ERPs e bancos exportam "Excel" como uma página HTML com extensão .xls; o Excel abre normalmente.

    Certo: ler a tabela (o conteúdo é uma tabela de verdade).
    """
    # Uma página HTML com uma tabela: cabeçalho e uma pessoa
    conteudo = ("<html><head><meta charset=\"utf-8\"></head><body><table>"
                "<tr><th>Nome</th><th>CPF</th><th>Cargo</th></tr>"
                f"<tr><td>Ana Souza</td><td>{CPF_DA_ANA}</td><td>Analista</td></tr>"
                "</table></body></html>").encode("utf-8")
    # Lê o arquivo, com o nome .xls que o sistema deu
    leitura = ler_do_disco(tmp_path, "exportacao.xls", conteudo)
    # Certo: a tabela lida
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo"]
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA, "Analista"]]


def planilha_ods_de_teste() -> bytes:
    """Uma planilha do LibreOffice (.ods) pequena, feita na hora com o pandas."""
    # O cabeçalho e uma pessoa, como uma tabela do pandas
    tabela = pandas.DataFrame([["Nome", "CPF"], ["Ana Souza", CPF_DA_ANA]])
    # Grava no formato do LibreOffice, sem o cabeçalho e o índice do próprio pandas
    memoria = io.BytesIO()
    tabela.to_excel(memoria, engine="odf", header=False, index=False)
    return memoria.getvalue()


def planilha_xlsx_de_teste() -> bytes:
    """Uma planilha .xlsx pequena (cabeçalho e uma pessoa)."""
    # O cabeçalho e uma pessoa
    livro = Workbook()
    livro.active.append(["Nome", "CPF"])
    livro.active.append(["Ana Souza", CPF_DA_ANA])
    return bytes_da_planilha(livro)


def word_de_teste() -> bytes:
    """Um Word (.docx) pequeno, com uma ficha "Rótulo: valor"."""
    return bytes_do_word(["Nome: Ana Souza", f"CPF: {CPF_DA_ANA}"])


@pytest.mark.parametrize("nome_do_arquivo, gerar_conteudo", [
    ("lista.xlsx", planilha_xlsx_de_teste),
    ("lista.ods", planilha_ods_de_teste),
    ("lista.xls", ARQUIVO_XLS_DE_TESTE.read_bytes),
    ("lista.docx", word_de_teste),
], ids=["xlsx", "ods", "xls", "docx"])
def test_arquivo_cortado_no_meio_e_recusado_como_corrompido(tmp_path, nome_do_arquivo, gerar_conteudo):
    """Download interrompido: só a primeira metade dos bytes chegou. A assinatura confere, mas o arquivo não abre."""
    # O arquivo inteiro e válido do formato do caso
    conteudo_inteiro = gerar_conteudo()
    # Fica só com a primeira metade dos bytes
    metade = conteudo_inteiro[:len(conteudo_inteiro) // 2]
    # Recusa dizendo que o arquivo pode estar corrompido (e não um erro técnico)
    with pytest.raises(ingestao.ArquivoRecusado, match="corrompido"):
        ler_do_disco(tmp_path, nome_do_arquivo, metade)


# ============================== 12. Limites: tamanho, linhas e colunas ==============================

def csv_com_tamanho_exato(tamanho_em_bytes: int) -> bytes:
    """Um CSV válido com exatamente o tamanho pedido: cabeçalho, pessoas com uma observação de 100 letras e a última
    observação completada até o byte exato.

    Exemplo: csv_com_tamanho_exato(1_048_576) → cerca de 8.000 pessoas, 1 MB certinho.
    """
    # O cabeçalho
    comeco = "Nome;CPF;Observação\n".encode("utf-8")
    # Uma linha de pessoa típica, com uma observação comprida (para chegar a 1 MB sem passar do limite de linhas)
    linha_de_pessoa = f"Pessoa;{CPF_DA_ANA};{'x' * 100}\n".encode("utf-8")
    # Quantas pessoas cabem inteiras, deixando espaço para a última linha
    quantidade_de_pessoas = (tamanho_em_bytes - len(comeco)) // len(linha_de_pessoa) - 1
    # O cabeçalho seguido das pessoas inteiras
    conteudo = comeco + linha_de_pessoa * quantidade_de_pessoas
    # A última pessoa tem a observação do tamanho que falta para o byte exato (menos o fim de linha)
    ultima_linha_sem_observacao = f"Ultima pessoa;{CPF_DA_ANA};".encode("utf-8")
    letras_que_faltam = tamanho_em_bytes - len(conteudo) - len(ultima_linha_sem_observacao) - 1
    # Junta tudo: as pessoas, a última linha completada e o fim de linha
    return conteudo + ultima_linha_sem_observacao + b"x" * letras_que_faltam + b"\n"


def test_arquivo_exatamente_no_limite_de_tamanho_e_aceito(tmp_path):
    """No limite exato (1 MB aqui), o arquivo passa."""
    # Um CSV de 1 MB certinho (confere o tamanho antes, para o teste não mentir)
    conteudo = csv_com_tamanho_exato(UM_MEGABYTE)
    assert len(conteudo) == UM_MEGABYTE
    # Lê com o limite de 1 MB
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo, limite_em_bytes=UM_MEGABYTE)
    # A última pessoa chegou inteira
    assert leitura.linhas[-1][0] == "Ultima pessoa"


def test_arquivo_um_byte_acima_do_limite_e_recusado_com_o_limite(tmp_path):
    """Um byte acima do limite: recusado, e a mensagem diz o limite em MB."""
    # Um CSV com 1 MB e 1 byte
    conteudo = csv_com_tamanho_exato(UM_MEGABYTE + 1)
    # Recusa dizendo o limite
    with pytest.raises(ingestao.ArquivoRecusado, match="limite de 1 MB"):
        ler_do_disco(tmp_path, "lista.csv", conteudo, limite_em_bytes=UM_MEGABYTE)


def csv_com_pessoas(quantidade_de_pessoas: int, linhas_em_branco_no_fim: int = 0) -> bytes:
    """Um CSV com o cabeçalho e a quantidade pedida de pessoas fictícias (e, se pedido, linhas em branco no fim)."""
    linhas_do_arquivo = ["Nome;CPF;Cargo"]
    # Uma linha por pessoa, com um nome diferente em cada uma
    for numero_da_pessoa in range(1, quantidade_de_pessoas + 1):
        linhas_do_arquivo.append(f"Pessoa {numero_da_pessoa};{CPF_DA_ANA};Analista")
    # As linhas em branco que alguns sistemas deixam no fim do arquivo
    for _linha_vazia in range(linhas_em_branco_no_fim):
        linhas_do_arquivo.append("")
    # Junta as linhas num texto só e grava em UTF-8
    return texto_do_csv(linhas_do_arquivo).encode("utf-8")


def test_arquivo_com_o_maximo_de_linhas_e_aceito(tmp_path):
    """Exatamente 20.000 funcionários: é o máximo, e passa."""
    # Lê um CSV com 20.000 pessoas
    leitura = ler_do_disco(tmp_path, "lista.csv", csv_com_pessoas(ingestao.MAXIMO_DE_LINHAS))
    # Todas chegaram
    assert len(leitura.linhas) == ingestao.MAXIMO_DE_LINHAS


def test_arquivo_com_uma_linha_acima_do_maximo_e_recusado(tmp_path):
    """20.001 funcionários: recusado, com o pedido para dividir o arquivo."""
    # Um CSV com 20.001 pessoas é recusado, com o pedido para dividir
    with pytest.raises(ingestao.ArquivoRecusado, match="mais de 20000 linhas: divida"):
        ler_do_disco(tmp_path, "lista.csv", csv_com_pessoas(ingestao.MAXIMO_DE_LINHAS + 1))


def test_linhas_em_branco_no_fim_nao_contam_para_o_limite(tmp_path):
    """20.000 funcionários e 30 linhas em branco no fim: continuam sendo 20.000 funcionários."""
    # 20.000 pessoas e 30 linhas vazias no fim do arquivo
    conteudo = csv_com_pessoas(ingestao.MAXIMO_DE_LINHAS, linhas_em_branco_no_fim=30)
    # Lê o arquivo
    leitura = ler_do_disco(tmp_path, "lista.csv", conteudo)
    # Certo: as 20.000 pessoas
    assert len(leitura.linhas) == ingestao.MAXIMO_DE_LINHAS


def csv_com_colunas(quantidade_de_colunas: int) -> bytes:
    """Um CSV com a quantidade pedida de colunas ("Campo 1", "Campo 2"...) e uma pessoa."""
    nomes_das_colunas = []
    valores_da_pessoa = []
    # Uma coluna de cada vez: o nome no cabeçalho e um valor na linha da pessoa
    for numero_da_coluna in range(1, quantidade_de_colunas + 1):
        nomes_das_colunas.append(f"Campo {numero_da_coluna}")
        valores_da_pessoa.append(f"valor {numero_da_coluna}")
    # O cabeçalho e a pessoa, separados por ";", em UTF-8
    return texto_do_csv([";".join(nomes_das_colunas), ";".join(valores_da_pessoa)]).encode("utf-8")


def test_arquivo_com_o_maximo_de_colunas_e_aceito(tmp_path):
    """Exatamente 200 colunas: passa."""
    # Lê um CSV com 200 colunas
    leitura = ler_do_disco(tmp_path, "lista.csv", csv_com_colunas(ingestao.MAXIMO_DE_COLUNAS))
    # Todas as colunas chegaram
    assert len(leitura.cabecalhos) == ingestao.MAXIMO_DE_COLUNAS


def test_arquivo_com_uma_coluna_acima_do_maximo_e_recusado(tmp_path):
    """201 colunas: recusado, com o pedido para conferir se é a lista de funcionários."""
    # Um CSV com 201 colunas é recusado
    with pytest.raises(ingestao.ArquivoRecusado, match="mais de 200 colunas"):
        ler_do_disco(tmp_path, "lista.csv", csv_com_colunas(ingestao.MAXIMO_DE_COLUNAS + 1))


# ============================== 13. Word (só os caminhos sem IA) ==============================

def test_word_com_tabela_acentos_e_uma_tabela_de_assinatura(tmp_path):
    """Word com um parágrafo, a tabela de funcionários (com uma linha vazia no fim, como o Word deixa) e uma
    tabelinha de assinatura: a tabela maior é a lida, sem IA."""
    # Documento novo, com um parágrafo de apresentação
    documento = Document()
    documento.add_paragraph("Segue a lista de admissões de setembro.")
    # A tabela de funcionários, com a última linha vazia
    linhas_da_tabela = [["Nome", "CPF", "Função"],
                        ["Maria Conceição", CPF_COM_UM_ZERO, "Técnica de manutenção"],
                        ["João Araújo", CPF_DA_BIA, "Açougueiro"],
                        ["", "", ""]]
    tabela = documento.add_table(rows=len(linhas_da_tabela), cols=3)
    for numero_da_linha, linha in enumerate(linhas_da_tabela):
        for numero_da_coluna, valor in enumerate(linha):
            # Escreve cada valor na sua célula da tabela
            tabela.cell(numero_da_linha, numero_da_coluna).text = valor
    # A tabelinha de assinatura, de uma linha só
    assinatura = documento.add_table(rows=1, cols=2)
    assinatura.cell(0, 0).text = "Assinatura do RH"
    assinatura.cell(0, 1).text = "Data"
    # Salva o documento na memória e lê
    memoria = io.BytesIO()
    documento.save(memoria)
    leitura = ler_do_disco(tmp_path, "lista.docx", memoria.getvalue())
    # A tabela dos funcionários, sem a linha vazia e sem chamar a IA
    assert leitura.cabecalhos == ["Nome", "CPF", "Função"]
    assert leitura.linhas == linhas_da_tabela[1:3]
    assert leitura.uso_da_ia is None


def test_word_com_tabela_e_celula_mesclada_na_vertical(tmp_path):
    """A unidade mesclada em duas linhas da tabela do Word vale para as duas pessoas."""
    # Documento novo com uma tabela de 3 linhas e 3 colunas
    documento = Document()
    tabela = documento.add_table(rows=3, cols=3)
    valores = [["Nome", "CPF", "Unidade"], ["Ana Souza", CPF_DA_ANA, ""], ["Bia Lima", CPF_DA_BIA, ""]]
    for numero_da_linha, linha in enumerate(valores):
        for numero_da_coluna, valor in enumerate(linha):
            # Escreve cada valor na sua célula da tabela
            tabela.cell(numero_da_linha, numero_da_coluna).text = valor
    # Mescla a unidade das duas pessoas e escreve o valor depois de mesclar
    unidade = tabela.cell(1, 2).merge(tabela.cell(2, 2))
    unidade.text = "Fábrica Campinas"
    # Salva o documento na memória e lê
    memoria = io.BytesIO()
    documento.save(memoria)
    leitura = ler_do_disco(tmp_path, "lista.docx", memoria.getvalue())
    # A unidade aparece nas duas pessoas
    assert leitura.linhas == [["Ana Souza", CPF_DA_ANA, "Fábrica Campinas"],
                              ["Bia Lima", CPF_DA_BIA, "Fábrica Campinas"]]


def test_word_com_titulo_mesclado_dentro_da_tabela(tmp_path):
    """A primeira linha da tabela é um título mesclado nas 3 colunas; o cabeçalho de verdade é a segunda linha."""
    # Documento novo com uma tabela de 4 linhas (título, cabeçalho e duas pessoas)
    documento = Document()
    tabela = documento.add_table(rows=4, cols=3)
    valores = [["", "", ""], ["Nome", "CPF", "Cargo"], ["Ana Souza", CPF_DA_ANA, "Analista"],
               ["Bia Lima", CPF_DA_BIA, "Vendedora"]]
    for numero_da_linha, linha in enumerate(valores):
        for numero_da_coluna, valor in enumerate(linha):
            # Escreve cada valor na sua célula da tabela
            tabela.cell(numero_da_linha, numero_da_coluna).text = valor
    # O título ocupa a primeira linha inteira
    titulo = tabela.cell(0, 0).merge(tabela.cell(0, 2))
    titulo.text = "Lista de funcionários"
    # Salva o documento na memória e lê
    memoria = io.BytesIO()
    documento.save(memoria)
    leitura = ler_do_disco(tmp_path, "lista.docx", memoria.getvalue())
    # Certo: o cabeçalho da segunda linha
    assert leitura.cabecalhos == ["Nome", "CPF", "Cargo"]
    assert primeira_coluna(leitura) == ["Ana Souza", "Bia Lima"]


def test_word_em_texto_com_fichas_titulo_e_acentos_sem_ia(tmp_path):
    """Word em texto, no jeito "Rótulo: valor", com título e linhas em branco: lido por regra, sem IA."""
    # Um título, uma linha em branco e duas fichas com acentos nos rótulos e nos valores
    paragrafos = ["Admissões de setembro", "",
                  "Nome: Maria Conceição", f"CPF: {CPF_COM_UM_ZERO}", "Função: Técnica de manutenção", "",
                  "Nome: João Araújo", f"CPF: {CPF_DA_BIA}", "Função: Açougueiro"]
    # Lê o Word (o cliente de IA falha se for chamado)
    leitura = ler_do_disco(tmp_path, "fichas.docx", bytes_do_word(paragrafos))
    # Uma coluna por rótulo, uma linha por ficha, e nenhuma IA usada
    assert leitura.cabecalhos == ["Nome", "CPF", "Função"]
    assert leitura.linhas == [["Maria Conceição", CPF_COM_UM_ZERO, "Técnica de manutenção"],
                              ["João Araújo", CPF_DA_BIA, "Açougueiro"]]
    assert leitura.uso_da_ia is None


@pytest.mark.parametrize("separador", [";", "\t"], ids=["ponto_e_virgula", "tabulacao"])
def test_word_com_texto_em_colunas_e_lido_por_regra(tmp_path, separador):
    """A empresa colou a planilha no Word como texto: cada parágrafo é uma linha com colunas.

    Certo: ler como tabela, por regra e sem custo (o mesmo que o .txt com colunas já faz).
    """
    linhas_da_tabela = [["Nome", "CPF", "Cargo"],
                        ["Ana Souza", CPF_DA_ANA, "Analista"],
                        ["Bia Lima", CPF_DA_BIA, "Vendedora"]]
    paragrafos = []
    # Cada linha da tabela vira um parágrafo, com as células unidas pelo separador
    for linha in linhas_da_tabela:
        paragrafos.append(separador.join(linha))
    # Lê o Word (o cliente de IA falha se for chamado)
    leitura = ler_do_disco(tmp_path, "lista.docx", bytes_do_word(paragrafos))
    # Certo: a tabela lida por regra
    assert leitura.cabecalhos == linhas_da_tabela[0]
    assert leitura.linhas == linhas_da_tabela[1:]
