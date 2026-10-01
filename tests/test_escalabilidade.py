"""Escalabilidade sem mudar código: o banco cadastra campos novos e o arquivo com eles é homologado.

É o ensaio da demo, automatizado:
1. o especialista do banco salva uma versão nova do layout com dois campos (vale-refeição e PLR), do tipo
   "dinheiro" da lista fechada, como faria na tela Parâmetros;
2. a Aurora envia a sua planilha com duas colunas a mais ("Vale Refeição" e "PLR");
3. o arquivo passa pelo fluxo inteiro (as decisões da empresa pelo gabarito) e é homologado;
4. o arquivo final traz os dois campos novos, com os valores convertidos, na versão 2 do layout.
Nenhuma linha de código muda: só o parâmetro.
"""
import csv
import io

import openpyxl
import pytest

from eval import avaliacao_do_fluxo
from services import banco, homologacao, parametros, processamentos
from tests.test_fluxo_empresa import busca_falsa

# Os campos novos, como o banco cadastraria na tela
CAMPOS_NOVOS = [
    {"campo": "valor_vale_refeicao", "grupo": "Benefícios", "tipo": "DECIMAL_MONETARIO", "obrigatorio": False,
     "sensivel": False, "uso_comercial_permitido": False, "descricao": "Valor mensal do vale-refeição",
     "regra": "Valor em reais, com duas casas", "nao_confundir_com": "valor_renda", "exemplo": "650,00"},
    {"campo": "valor_plr", "grupo": "Remuneração variável", "tipo": "DECIMAL_MONETARIO", "obrigatorio": False,
     "sensivel": False, "uso_comercial_permitido": False, "descricao": "Participação nos lucros e resultados",
     "regra": "Valor em reais, com duas casas", "nao_confundir_com": "valor_renda", "exemplo": "1.200,50"},
]
# As colunas novas na planilha da Aurora e o campo de cada uma
COLUNAS_NOVAS = {"Vale Refeição": "valor_vale_refeicao", "PLR": "valor_plr"}
NOME_DO_ARQUIVO = "aurora_com_beneficios.xlsx"


@pytest.fixture(scope="module", autouse=True)
def gerar_envios():
    """Os arquivos de envio não vão para o Git: gera antes dos testes."""
    from scripts.gerar_dados import main
    main()


def _planilha_com_colunas_novas(pasta) -> None:
    """A planilha da Aurora com "Vale Refeição" e "PLR" no fim, gravada na pasta informada."""
    planilha = openpyxl.load_workbook(avaliacao_do_fluxo.PASTA_DOS_ENVIOS / "aurora_carga_inicial.xlsx")
    aba = planilha.active
    coluna_do_vale = aba.max_column + 1
    coluna_da_plr = aba.max_column + 2
    # Cabeçalho na linha 1; valores no formato brasileiro, como o RH digitaria
    aba.cell(row=1, column=coluna_do_vale, value="Vale Refeição")
    aba.cell(row=1, column=coluna_da_plr, value="PLR")
    for linha in range(2, aba.max_row + 1):
        aba.cell(row=linha, column=coluna_do_vale, value="R$ 650,00")
        aba.cell(row=linha, column=coluna_da_plr, value="1.200,50")
    planilha.save(pasta / NOME_DO_ARQUIVO)


def test_campo_novo_cadastrado_chega_ao_arquivo_homologado(tmp_path, monkeypatch):
    """Layout v2 com dois campos novos: o arquivo com eles é homologado e o arquivo final os traz convertidos."""
    conexao = banco.conectar(tmp_path / "teste.db")
    # 1. O banco salva a versão 2 do layout: os 44 campos de sempre + os 2 novos
    _versao_atual, campos_atuais = parametros.layout_ativo(conexao)
    linhas_do_layout = []
    for campo in campos_atuais:
        linhas_do_layout.append(campo.model_dump(mode="json"))
    linhas_do_layout.extend(CAMPOS_NOVOS)
    assert parametros.salvar_layout(conexao, linhas_do_layout, "especialista.banco") == 2

    # 2. A Aurora envia a planilha com as duas colunas novas
    _planilha_com_colunas_novas(tmp_path)
    monkeypatch.setattr(avaliacao_do_fluxo, "PASTA_DOS_ENVIOS", tmp_path)
    gabarito = dict(avaliacao_do_fluxo.carregar_gabaritos()["aurora_carga_inicial"])
    gabarito["arquivo"] = NOME_DO_ARQUIVO
    mapeamento = dict(gabarito["mapeamento"])
    mapeamento.update(COLUNAS_NOVAS)
    gabarito["mapeamento"] = mapeamento

    # 3. O fluxo inteiro, com a pessoa simulada decidindo pelo gabarito
    medido = avaliacao_do_fluxo.processar_arquivo(conexao, "aurora_com_beneficios", gabarito,
                                                  avaliacao_do_fluxo.carregar_verdade(), busca=busca_falsa)
    assert medido["status_final"] == "HOMOLOGADO"

    # 4. O arquivo final: versão 2 do layout, com os campos novos convertidos para o padrão do banco
    processamento_id = None
    for perfil in processamentos.listar(conexao, "EMP001"):
        if perfil.nome_arquivo == NOME_DO_ARQUIVO:
            processamento_id = perfil.processamento_id
    homologado = homologacao.obter(conexao, processamento_id)
    assert homologado["relatorio"]["versao_layout"] == 2
    linhas = list(csv.DictReader(io.StringIO(homologado["arquivo"].decode("utf-8")), delimiter=";"))
    assert len(linhas) == homologado["relatorio"]["registros_homologados"]
    for linha in linhas:
        assert linha["valor_vale_refeicao"] == "650.00"
        assert linha["valor_plr"] == "1200.50"
    conexao.close()
