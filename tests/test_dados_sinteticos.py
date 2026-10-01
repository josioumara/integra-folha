"""Dados sintéticos: reprodutíveis, coerentes com o gabarito e sem vazar o vocabulário de teste (ADR-21)."""
import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.documentos import cnpj_valido, cpf_valido
from tests.test_vocabulario_segregado import normalizar

# Pastas dos dados gerados e dos gabaritos
RAIZ = Path(__file__).resolve().parent.parent
SINTETICO = RAIZ / "data" / "synthetic"
GOLDEN = RAIZ / "data" / "golden"


@pytest.fixture(scope="module", autouse=True)
def gerar_uma_vez():
    """Os arquivos de envio não vão para o Git: gera antes dos testes."""
    from scripts.gerar_dados import main
    main()


def ler_csv(caminho: Path, separador: str = ",", codificacao: str = "utf-8") -> list[dict]:
    """As linhas de um CSV, como dicionários."""
    with open(caminho, encoding=codificacao, newline="") as arquivo:
        return list(csv.DictReader(arquivo, delimiter=separador))


def ler_envio(gabarito: dict) -> tuple[list[str], list[list]]:
    """Lê um arquivo de envio do jeito que o gabarito descreve (a ingestão faz isso sem gabarito)."""
    caminho = SINTETICO / "envios" / gabarito["arquivo"]
    formato = gabarito["formato"]
    if formato["tipo"] == "csv":
        with open(caminho, encoding=formato["codificacao"], newline="") as arquivo:
            linhas = list(csv.reader(arquivo, delimiter=formato["separador"]))
    else:
        linhas = []
        for linha in load_workbook(caminho).active.iter_rows(values_only=True):
            linhas.append(list(linha))
    cabecalho = linhas[formato["linha_do_cabecalho"] - 1]
    return list(cabecalho), linhas[formato["linha_do_cabecalho"]:]


def gabaritos() -> list[dict]:
    """Todos os gabaritos, em ordem alfabética."""
    todos = []
    for caminho in sorted(GOLDEN.glob("*.json")):
        todos.append(json.loads(caminho.read_text(encoding="utf-8")))
    return todos


def _ler_gabarito(nome: str) -> dict:
    """O gabarito de um arquivo, pelo nome sem extensão."""
    return json.loads((GOLDEN / f"{nome}.json").read_text(encoding="utf-8"))


def _nome_do_arquivo(gabarito: dict) -> str:
    """O nome do caso de teste: o arquivo do gabarito."""
    return gabarito["arquivo"]


def test_gabarito_tem_256_funcionarios_com_cpf_e_cnpj_validos():
    """256 pessoas, todas com CPF e CNPJ válidos e nenhum CPF repetido."""
    verdade = ler_csv(SINTETICO / "funcionarios_truth.csv")
    assert len(verdade) == 256
    cpfs = set()
    for funcionario in verdade:
        assert cpf_valido(funcionario["cpf"])
        assert cnpj_valido(funcionario["cnpj_empregador"])
        cpfs.add(funcionario["cpf"])
    # Ninguém repetido no gabarito
    assert len(cpfs) == 256


def _coluna_do_campo(gabarito: dict, campo: str) -> str:
    """A coluna do arquivo que alimenta o campo (a ambígua vale pelo campo após a confirmação)."""
    for coluna, campo_mapeado in gabarito["mapeamento"].items():
        if campo_mapeado is None:
            campo_mapeado = gabarito["colunas_ambiguas"].get(coluna, {}).get("campo_apos_confirmacao")
        if campo_mapeado == campo:
            return coluna
    raise KeyError(campo)


@pytest.mark.parametrize("gabarito", gabaritos(), ids=_nome_do_arquivo)
def test_arquivo_bate_com_o_gabarito(gabarito):
    """Cabeçalho, número de linhas e cada erro injetado estão onde o gabarito diz."""
    cabecalho, linhas = ler_envio(gabarito)
    assert cabecalho == list(gabarito["mapeamento"]) + gabarito["colunas_extras"]
    assert len(linhas) == len(gabarito["funcionario_ids"])
    for erro in gabarito["erros"]:
        # Estes são verificados nos testes específicos abaixo
        if erro["tipo"] in ("PESSOA_DUPLICADA_NO_ARQUIVO", "JA_HOMOLOGADO_NA_EMPRESA"):
            continue
        coluna = _coluna_do_campo(gabarito, erro["campo"])
        valor = linhas[erro["linha"] - 1][cabecalho.index(coluna)]
        esperado = erro["valor_no_arquivo"]
        if isinstance(valor, (int, float)):
            # O Excel devolve 1100.0 como 1100
            assert float(valor) == float(esperado)
        elif erro["tipo"] == "CAMPO_OBRIGATORIO_VAZIO":
            assert valor in ("", None)
        else:
            assert str(valor) == esperado


def test_nenhum_cabecalho_dos_envios_vem_do_vocabulario_de_teste():
    """Os arquivos das empresas usam só nomes do treino: a prova continua inédita."""
    termos_do_teste = set()
    for linha in ler_csv(RAIZ / "data" / "vocabulario" / "teste.csv"):
        termos_do_teste.add(normalizar(linha["termo"]))
    for gabarito in gabaritos():
        for coluna in gabarito["mapeamento"]:
            assert normalizar(coluna) not in termos_do_teste, f"{coluna!r} ({gabarito['arquivo']}) vazou do teste"


def test_vale_verde_tem_cabecalho_na_linha_4_e_colunas_extras():
    """Vale Verde: títulos em cima da tabela e duas colunas que não são do layout."""
    gabarito = _ler_gabarito("vale_verde_carga_inicial")
    assert gabarito["formato"]["linha_do_cabecalho"] == 4
    assert gabarito["colunas_extras"] == ["Obs. RH", "Cód. Interno"]


def test_brisa_grava_cpf_como_numero_e_perde_zeros_a_esquerda():
    """Brisa: o CPF está gravado como número no Excel."""
    gabarito = _ler_gabarito("brisa_carga_inicial")
    cabecalho, linhas = ler_envio(gabarito)
    posicao_do_cpf = cabecalho.index(_coluna_do_campo(gabarito, "cpf"))
    for linha in linhas:
        assert isinstance(linha[posicao_do_cpf], int)


def test_atlantico_tem_vencimentos_ambiguo_e_renda_dez_vezes_acima():
    """Atlântico: "Vencimentos" é ambígua e as rendas injetadas passam do teto do cargo."""
    gabarito = _ler_gabarito("atlantico_carga_inicial")
    assert gabarito["mapeamento"]["Vencimentos"] is None
    assert gabarito["colunas_ambiguas"]["Vencimentos"]["campo_apos_confirmacao"] == "valor_renda"
    faixas = {}
    for faixa in ler_csv(SINTETICO / "faixas_referencia_cargo.csv"):
        faixas[faixa["cargo"]] = faixa
    verdade = {}
    for funcionario in ler_csv(SINTETICO / "funcionarios_truth.csv"):
        verdade[funcionario["funcionario_id"]] = funcionario
    for erro in gabarito["erros"]:
        if erro["tipo"] != "RENDA_FORA_DO_CARGO":
            continue
        cargo = verdade[gabarito["funcionario_ids"][erro["linha"] - 1]]["cargo"]
        assert float(erro["valor_no_arquivo"]) > float(faixas[cargo]["faixa_maxima"])


def test_prisma_tem_uma_pessoa_duplicada_no_arquivo():
    """Prisma: uma pessoa aparece duas vezes."""
    funcionario_ids = _ler_gabarito("prisma_carga_inicial")["funcionario_ids"]
    assert len(funcionario_ids) == len(set(funcionario_ids)) + 1


def test_inclusao_da_brisa_traz_alguem_da_carga_inicial():
    """A inclusão da Brisa traz alguém que já estava na carga inicial."""
    inicial = _ler_gabarito("brisa_carga_inicial")
    inclusao = _ler_gabarito("brisa_inclusao")
    assert set(inicial["funcionario_ids"]) & set(inclusao["funcionario_ids"])


def test_inclusao_da_aurora_usa_as_mesmas_colunas_e_a_da_horizonte_nao():
    """Aurora repete as colunas (reuso do mapeamento); Horizonte muda os nomes (a IA interpreta de novo)."""
    def colunas(nome):
        """As colunas do layout no arquivo."""
        return list(_ler_gabarito(nome)["mapeamento"])
    assert colunas("aurora_inclusao") == colunas("aurora_carga_inicial")
    assert colunas("horizonte_inclusao") != colunas("horizonte_carga_inicial")


def _conteudo_dos_arquivos_gerados() -> dict:
    """Nome -> bytes dos CSVs gerados e dos gabaritos."""
    conteudos = {}
    for caminho in list(SINTETICO.glob("*.csv")) + list(GOLDEN.glob("*.json")):
        conteudos[caminho.name] = caminho.read_bytes()
    return conteudos


def _conteudo_das_planilhas() -> dict:
    """Nome -> células das planilhas enviadas (os bytes do .xlsx mudam com a data de gravação)."""
    planilhas = {}
    for caminho in (SINTETICO / "envios").glob("*.xlsx"):
        planilhas[caminho.name] = ler_envio({"arquivo": caminho.name,
                                             "formato": {"tipo": "xlsx", "linha_do_cabecalho": 1}})
    return planilhas


def test_mesma_seed_gera_o_mesmo_conteudo():
    """Gerar de novo dá os mesmos arquivos e as mesmas planilhas."""
    antes = _conteudo_dos_arquivos_gerados()
    planilhas_antes = _conteudo_das_planilhas()
    from scripts.gerar_dados import main
    main()
    assert antes == _conteudo_dos_arquivos_gerados()
    planilhas_depois = _conteudo_das_planilhas()
    for nome, conteudo in planilhas_antes.items():
        assert planilhas_depois[nome] == conteudo
