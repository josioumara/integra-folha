"""Testes do gerador da prova por tipo de arquivo (scripts/gerar_prova_por_tipo.py).

O que é conferido:
- as MESMAS pessoas aparecem no arquivo de mesmo número de todos os 7 tipos (a diferença entre os tipos é só o
  formato), e o nome e o CPF de cada uma estão mesmo dentro do arquivo;
- os nomes das colunas do T3 e do T4 vêm só do vocabulário da PROVA, nunca do de treino;
- todo arquivo traz os 4 obrigatórios de cada pessoa (menos onde a armadilha do texto misturado tira um de propósito);
- o código CBO de cada cargo existe na tabela oficial, e a renda que o gerador sorteia fica dentro da faixa pública;
- gerar de novo dá o mesmo conjunto que está gravado, e cada arquivo gravado bate com a impressão digital do manifesto.
"""
import csv
import io
import json
from decimal import Decimal

from docx import Document
from openpyxl import load_workbook

from baselines.baseline_mapper import ler_termos, normalizar
from eval.congelamento import impressao_digital
from scripts import gerar_cabecalhos, gerar_dados
from scripts import gerar_prova_por_tipo as gerador
from services import tabela_cbo

# A prova gravada no repositório
PASTA_GRAVADA = gerador.PASTA_DA_PROVA


def _manifesto_gravado() -> dict:
    """O manifesto da prova gravada."""
    return json.loads((PASTA_GRAVADA / "manifesto.json").read_text(encoding="utf-8"))


def _celulas_da_planilha(conteudo: bytes, nome: str) -> list[list[str]]:
    """As linhas da planilha (Excel ou CSV), com cada célula como texto."""
    if nome.endswith(".csv"):
        return list(csv.reader(io.StringIO(conteudo.decode("utf-8")), delimiter=";"))
    linhas = []
    for linha in load_workbook(io.BytesIO(conteudo)).active.iter_rows(values_only=True):
        celulas = []
        for valor in linha:
            celulas.append("" if valor is None else str(valor))
        linhas.append(celulas)
    return linhas


def _texto_do_arquivo(caminho) -> str:
    """Todo o texto do arquivo numa string só (as células da planilha ou os parágrafos e as tabelas do Word)."""
    conteudo = caminho.read_bytes()
    if caminho.suffix == ".docx":
        documento = Document(io.BytesIO(conteudo))
        pedacos = []
        for paragrafo in documento.paragraphs:
            pedacos.append(paragrafo.text)
        for tabela in documento.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    pedacos.append(celula.text)
        return "\n".join(pedacos)
    pedacos = []
    for linha in _celulas_da_planilha(conteudo, caminho.name):
        pedacos.append(";".join(linha))
    return "\n".join(pedacos)


def _so_digitos(texto: str) -> str:
    """Só os dígitos do texto."""
    digitos = ""
    for caractere in texto:
        if caractere.isdigit():
            digitos += caractere
    return digitos


def test_as_mesmas_pessoas_em_todos_os_tipos():
    manifesto = _manifesto_gravado()
    valores_da_pessoa = {}
    for pessoa in manifesto["pessoas"]:
        valores_da_pessoa[pessoa["pessoa_id"]] = pessoa["valores"]
    # O arquivo de número N de todos os tipos tem as mesmas pessoas, na mesma ordem
    pessoas_por_numero = {}
    for entrada in manifesto["arquivos"]:
        pessoas_por_numero.setdefault(entrada["numero"], set()).add(tuple(entrada["pessoas"]))
    assert len(pessoas_por_numero) == gerador.ARQUIVOS_POR_TIPO
    for conjuntos in pessoas_por_numero.values():
        assert len(conjuntos) == 1
    # O nome e o CPF de cada pessoa estão mesmo no arquivo (menos o CPF que a armadilha do texto misturado tira)
    for entrada in manifesto["arquivos"]:
        texto = _texto_do_arquivo(PASTA_GRAVADA / entrada["arquivo"])
        digitos_do_texto = _so_digitos(texto)
        for gabarito in entrada["gabarito"]:
            valores = valores_da_pessoa[gabarito["pessoa_id"]]
            assert valores["nome_completo"].lower() in texto.lower(), (entrada["arquivo"], gabarito["pessoa_id"])
            if gabarito.get("armadilha") != "cpf_depois":
                assert valores["cpf"] in digitos_do_texto, (entrada["arquivo"], gabarito["pessoa_id"])


def test_nomes_das_colunas_so_do_vocabulario_da_prova():
    manifesto = _manifesto_gravado()
    vocabulario_da_prova = gerar_cabecalhos.ler_vocabulario("teste")
    # Todos os termos de treino (e as ambíguas e a mais do treino), já no jeito de comparar
    termos_de_treino = set()
    for termo, _campo in ler_termos(gerador.RAIZ / "data" / "vocabulario" / "treino.csv"):
        termos_de_treino.add(normalizar(termo))
    for termo in gerar_cabecalhos.AMBIGUAS["treino"] + gerar_cabecalhos.EXTRAS["treino"]:
        termos_de_treino.add(normalizar(termo))
    colunas_conferidas = 0
    for entrada in manifesto["arquivos"]:
        if entrada["tipo"] not in ("T3", "T4"):
            continue
        for coluna in entrada["colunas"]:
            origem = coluna["origem"]
            # O termo de origem é da lista da prova certa
            if origem in ("vocabulario_da_prova", "sem_nome"):
                assert coluna["termo"] in vocabulario_da_prova[coluna["campos"][0]], coluna
            elif origem == "cbo_da_prova":
                assert coluna["termo"] in gerador.NOMES_DO_CBO_NA_PROVA
            elif origem == "ambigua_da_prova":
                assert coluna["termo"] in gerar_cabecalhos.AMBIGUAS["teste"]
            elif origem == "a_mais_da_prova":
                assert coluna["termo"] in gerar_cabecalhos.EXTRAS["teste"]
            else:
                assert origem == "juncao_nome_cpf"
            # Nem o termo nem o nome escrito no arquivo são termos de treino
            assert normalizar(coluna["termo"]) not in termos_de_treino, coluna
            assert normalizar(coluna["cabecalho"]) not in termos_de_treino, coluna
            colunas_conferidas += 1
    assert colunas_conferidas > 100


def test_os_quatro_obrigatorios_em_todo_arquivo():
    manifesto = _manifesto_gravado()
    for entrada in manifesto["arquivos"]:
        for gabarito in entrada["gabarito"]:
            for campo in gerador.CAMPOS_OBRIGATORIOS:
                # A armadilha do texto misturado tira o CPF ("vem depois") ou a admissão (duas datas) de propósito
                tirado_pela_armadilha = campo in gabarito["armadilhas"] or (
                    campo == "cpf" and gabarito.get("armadilha") == "cpf_depois")
                if not tirado_pela_armadilha:
                    assert campo in gabarito["campos"], (entrada["arquivo"], gabarito["pessoa_id"], campo)


def test_cbo_de_cada_cargo_existe_e_a_renda_fica_na_faixa_publica():
    faixas = {}
    with open(gerador.RAIZ / "data" / "cbo" / "faixas_salariais_cbo.csv", encoding="utf-8", newline="") as arquivo:
        for linha in csv.DictReader(arquivo):
            if linha["minimo"] and linha["maximo"]:
                faixas[linha["codigo_cbo"]] = (Decimal(linha["minimo"]), Decimal(linha["maximo"]))
    # Todo cargo do gerador da base sintética tem código
    assert set(gerador.CBO_DO_CARGO) == set(gerar_dados.CARGOS)
    for cargo, codigo in gerador.CBO_DO_CARGO.items():
        assert tabela_cbo.existe(codigo), cargo
        # O gerador sorteia a renda entre 75% e 125% da mediana do cargo
        mediana = gerar_dados.CARGOS[cargo][1]
        minimo, maximo = faixas[codigo]
        assert minimo <= mediana * Decimal("0.75") and mediana * Decimal("1.25") <= maximo, cargo


def test_gerar_de_novo_da_o_conjunto_gravado(tmp_path):
    manifesto_novo = gerador.gerar_prova(tmp_path)
    manifesto_gravado = _manifesto_gravado()
    # O manifesto é o mesmo, fora a impressão digital (o Excel e o Word guardam a hora em que foram salvos)
    for manifesto in (manifesto_novo, manifesto_gravado):
        for entrada in manifesto["arquivos"]:
            entrada.pop("sha256")
    assert json.loads(json.dumps(manifesto_novo, ensure_ascii=False)) == manifesto_gravado
    # E o conteúdo de cada arquivo é o mesmo
    for entrada in manifesto_gravado["arquivos"]:
        assert _texto_do_arquivo(tmp_path / entrada["arquivo"]) == _texto_do_arquivo(PASTA_GRAVADA / entrada["arquivo"])


def test_arquivos_gravados_batem_com_a_impressao_digital():
    manifesto = _manifesto_gravado()
    assert len(manifesto["arquivos"]) == len(gerador.TIPOS) * gerador.ARQUIVOS_POR_TIPO
    for entrada in manifesto["arquivos"]:
        assert impressao_digital(PASTA_GRAVADA / entrada["arquivo"]) == entrada["sha256"], entrada["arquivo"]
