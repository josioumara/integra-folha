"""A régua do experimento EXP-010 (eval/avaliacao_do_leitor.py) e os documentos de teste congelados (ADR-73).

O que se prova aqui:
    - valores iguais depois da padronização contam como acerto ("14 de setembro de 1991" = "14/09/1991");
    - pessoas são casadas pelo CPF, pelo e-mail ou pelo nome;
    - armadilha preenchida, pergunta que faltou e dado de terceiro que vazou aparecem na medida;
    - os documentos gerados batem com os gravados, nos dois conjuntos (a prova não muda sem ninguém ver);
    - os níveis lidos por regra (tabela e fichas) acertam 100%, sem IA.
"""
import json

from docx import Document

from eval import avaliacao_do_leitor
from models.contratos import carregar_layout
from services import leitura_de_word
from scripts import gerar_documentos_de_teste

CAMPOS = carregar_layout()
CONJUNTOS = tuple(gerar_documentos_de_teste.SEMENTE_POR_CONJUNTO)


def _campo(nome: str):
    """O campo do layout pelo nome."""
    for campo in CAMPOS:
        if campo.campo == nome:
            return campo
    raise KeyError(nome)


def test_valores_iguais_depois_da_padronizacao():
    assert avaliacao_do_leitor.valores_iguais("14/09/1991", "14 de setembro de 1991", _campo("data_nascimento"))
    assert avaliacao_do_leitor.valores_iguais("4.350,00", "R$ 4.350,00", _campo("valor_renda"))
    assert avaliacao_do_leitor.valores_iguais("(31) 99770-1182", "+55 31 99770-1182", _campo("telefone_celular"))
    assert avaliacao_do_leitor.valores_iguais("44.827.196-3", "44.827.196-3 SSP-SP", _campo("numero_documento"))
    assert avaliacao_do_leitor.valores_iguais("Maria Souza", "MARIA  SOUZA", _campo("nome_completo"))
    assert not avaliacao_do_leitor.valores_iguais("4.350,00", "4.530,00", _campo("valor_renda"))
    assert not avaliacao_do_leitor.valores_iguais("14/09/1991", "data ruim", _campo("data_nascimento"))


def test_pessoas_sao_casadas_pelo_cpf_pelo_email_ou_pelo_nome():
    esperadas = [{"campos": {"cpf": "529.982.247-25", "nome_completo": "Ana Lima"}},
                 {"campos": {"email_pessoal": "bia@exemplo.com", "nome_completo": "Bia Reis"}},
                 {"campos": {"nome_completo": "Caio Souza Prado"}},
                 {"campos": {"nome_completo": "Dora Alves"}}]
    lidas = [{"nome_completo": "CAIO SOUSA PRADO"}, {"email_pessoal": "bia@exemplo.com"}, {"cpf": "52998224725"}]
    assert avaliacao_do_leitor.casar_pessoas(esperadas, lidas) == [(0, 2), (1, 1), (2, 0), (3, None)]


def test_medida_conta_armadilha_pergunta_e_terceiro():
    documento = {"pessoas": [{"campos": {"cpf": "529.982.247-25", "nome_completo": "Ana Lima"},
                              "armadilhas": ["data_admissao"], "espera_duvida": True,
                              "terceiros": [{"nome": "Pedro Lima", "telefone": "31 3322-1100"}]}]}
    linhas = [["nome_completo", "cpf", "data_admissao", "nome_mae"],
              ["Ana Lima", "529.982.247-25", "10/03/2024", "Pedro Lima"]]
    medida = avaliacao_do_leitor.medir_documento(documento, linhas, [], CAMPOS)
    assert medida["campos_certos"] == 2 and medida["campos_esperados"] == 2
    # A armadilha foi preenchida, a pergunta não veio e o nome do dependente entrou num campo
    assert medida["armadilhas"] == 1 and medida["armadilhas_respeitadas"] == 0
    assert medida["perguntas_esperadas"] == 1 and medida["perguntas_feitas"] == 0
    assert medida["terceiros_vazados"] == 1


def test_resumo_calcula_acerto_e_intervalo():
    medidas = [{"campos_certos": 9, "campos_esperados": 10, "pessoas_esperadas": 2, "pessoas_encontradas": 2,
                "custo_usd": 0.1, "por_pessoa": [{"certos": 5, "esperados": 5}, {"certos": 4, "esperados": 5}]}]
    resumo = avaliacao_do_leitor.resumir(medidas)
    assert resumo["acerto_por_campo"] == 0.9
    assert 0.8 <= resumo["intervalo_95"][0] <= 0.9 <= resumo["intervalo_95"][1] <= 1.0
    assert resumo["custo_por_funcionario_usd"] == 0.05


def test_documentos_gerados_batem_com_os_gravados(tmp_path):
    """Gerar de novo (numa pasta temporária) dá o mesmo gabarito e o mesmo texto: a prova não muda escondida."""
    for conjunto in CONJUNTOS:
        pasta_gravada = gerar_documentos_de_teste.pasta_do_conjunto(conjunto)
        gerar_documentos_de_teste.gerar_conjunto(conjunto, tmp_path / conjunto)
        gabarito_novo = json.loads((tmp_path / conjunto / "gabarito.json").read_text(encoding="utf-8"))
        gabarito_gravado = json.loads((pasta_gravada / "gabarito.json").read_text(encoding="utf-8"))
        assert gabarito_novo == gabarito_gravado
        for documento in gabarito_gravado["documentos"]:
            texto_novo = []
            for paragrafo in Document(tmp_path / conjunto / documento["arquivo"]).paragraphs:
                texto_novo.append(paragrafo.text)
            texto_gravado = []
            for paragrafo in Document(pasta_gravada / documento["arquivo"]).paragraphs:
                texto_gravado.append(paragrafo.text)
            assert texto_novo == texto_gravado


def test_tabela_e_fichas_acertam_tudo_sem_ia():
    for conjunto in CONJUNTOS:
        pasta = gerar_documentos_de_teste.pasta_do_conjunto(conjunto)
        gabarito = json.loads((pasta / "gabarito.json").read_text(encoding="utf-8"))
        for documento in gabarito["documentos"]:
            if documento["nivel"] not in ("N1", "N2"):
                continue
            leitura = leitura_de_word.ler_word((pasta / documento["arquivo"]).read_bytes())
            medida = avaliacao_do_leitor.medir_documento(documento, leitura.linhas, leitura.duvidas, CAMPOS)
            assert medida["campos_certos"] == medida["campos_esperados"], documento["arquivo"]
