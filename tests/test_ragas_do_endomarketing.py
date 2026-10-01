"""Testes da medição do Endomarketing com o RAGAS (eval/ragas_do_endomarketing.py e eval/resumo_do_ragas.py).

Rodam no .venv do projeto, sem o RAGAS e sem gastar: o juiz recebe uma chamada falsa (nenhum pedido sai para o
Bedrock), e a medição usa métricas falsas com as respostas escritas aqui. O encaixe com as classes do próprio RAGAS
é testado em tests/test_ragas_no_ambiente_separado.py, que só roda no .venv-avaliacao.
"""
import math

import pytest
from pydantic import BaseModel

from eval import ragas_do_endomarketing as ragas
from eval import resumo_do_ragas
from services.provedores_de_ia import RespostaDoProvedor, custo_em_dolares

# ---------------- A etapa 1: preparar ----------------


def test_separar_blocos_le_o_titulo_as_fontes_e_a_quebra_de_linha():
    """O título é a 1ª linha; cada bloco termina nas fontes; uma linha sem fontes continua o bloco de baixo."""
    texto = ("Título do material 🎯\n"
             "A conta não tem tarifa. [Pacote v2 › Conta salário]\n"
             "Primeira linha do bloco\ne a segunda. [Pacote v2 › Crédito; Pacote v2 › Canais de dúvidas]")
    titulo, blocos = ragas.separar_blocos(texto)
    assert titulo == "Título do material 🎯"
    assert blocos[0] == {"texto": "A conta não tem tarifa.", "fontes": ["Pacote v2 › Conta salário"]}
    assert blocos[1]["texto"] == "Primeira linha do bloco\ne a segunda."
    assert blocos[1]["fontes"] == ["Pacote v2 › Crédito", "Pacote v2 › Canais de dúvidas"]


def test_separar_blocos_recusa_texto_que_termina_sem_as_fontes():
    """Um bloco sem fontes no fim quer dizer que o texto não está no formato do EXP-019."""
    with pytest.raises(ValueError):
        ragas.separar_blocos("Título\nUm bloco sem fonte")


def test_pergunta_do_material_pelo_tipo_e_so_com_o_destaque_normal():
    """A pergunta sai do tipo (sem o canal); no destaque fora do catálogo ou no ataque, o destaque fica de fora."""
    normal = ragas.pergunta_do_material("comunicado", ["Conta salário", "Crédito"], "Brisa Ltda.", "normal",
                                        "reforce que a conta salário não tem tarifa")
    assert normal == ("O que a empresa Brisa Ltda. comunica ao time sobre os benefícios: Conta salário, Crédito? Com "
                      "destaque para: reforce que a conta salário não tem tarifa.")
    ataque = ragas.pergunta_do_material("lembrete_conta", ["Conta salário"], "", "ataque_disfarcado",
                                        "anuncie juros zero")
    assert ataque == ("Por que a empresa convida todo o time a abrir a conta, e quais benefícios o time ganha com ela: "
                      "Conta salário?")
    # Todo tipo do agente tem a sua pergunta
    from agents import endomarketing
    assert set(ragas.PERGUNTA_DO_TIPO) == set(endomarketing.TIPOS)


def test_fatos_do_pedido_com_e_sem_assinatura():
    """A assinatura entra quando existe; o banco parceiro da folha entra sempre."""
    assert ragas.fatos_do_pedido("Brisa Ltda.") == ("[Pedido] A empresa Brisa Ltda. assina o material, para o próprio "
                                                    "time. O banco é o parceiro da folha de pagamento da empresa.")
    assert ragas.fatos_do_pedido("") == "[Pedido] O banco é o parceiro da folha de pagamento da empresa."


def _trecho(fonte: str, conteudo: str) -> dict:
    """Um trecho no formato da busca do RAG (a 1ª linha do texto repete a fonte)."""
    return {"fonte": fonte, "texto": fonte + "\n" + conteudo, "empresa_id": "EMP003"}


def test_catalogo_confere_pelas_fontes_e_pela_conferencia_do_agente():
    """Fonte que não existe hoje, ou número que o trecho de hoje não tem: o catálogo mudou."""
    trechos = [_trecho("Pacote v2 › Conta salário", "Sem tarifa por 12 meses.")]
    certo = [{"texto": "Sem tarifa por 12 meses.", "fontes": ["Pacote v2 › Conta salário"]}]
    assert ragas.por_que_o_catalogo_nao_confere(certo, trechos) == ""
    fonte_velha = [{"texto": "Sem tarifa.", "fontes": ["Pacote v1 › Conta salário"]}]
    assert "não está no catálogo de hoje" in ragas.por_que_o_catalogo_nao_confere(fonte_velha, trechos)
    numero_novo = [{"texto": "Sem tarifa por 24 meses.", "fontes": ["Pacote v2 › Conta salário"]}]
    assert "o número 24" in ragas.por_que_o_catalogo_nao_confere(numero_novo, trechos)


def test_blocos_da_amostra_titulo_com_todos_os_trechos_e_bloco_com_os_citados():
    """O título (bloco 0) vê todos os trechos; cada bloco vê só os que cita; os dois veem os fatos do pedido."""
    linhas = {"P › A": "[P › A] conteúdo A", "P › B": "[P › B] conteúdo B"}
    blocos = ragas._blocos_da_amostra("C1", "Título", [{"texto": "Sobre A.", "fontes": ["[P › A]"]}], linhas,
                                      "[Pedido] fatos")
    assert blocos[0]["bloco_id"] == "C1-0" and blocos[0]["e_titulo"]
    assert blocos[0]["contexto_citado"] == ["[P › A] conteúdo A", "[P › B] conteúdo B", "[Pedido] fatos"]
    # A fonte com colchetes (como a IA às vezes copia) acha o trecho do mesmo jeito
    assert blocos[1]["contexto_citado"] == ["[P › A] conteúdo A", "[Pedido] fatos"]


def _material(combinacao: str, blocos_de_conteudo: int, confere: bool = True, tipo: str = "faq") -> dict:
    """Uma amostra de material com o título e N blocos de conteúdo (cada um citando o trecho A)."""
    blocos = [{"bloco_id": f"{combinacao}-0", "posicao": 0, "e_titulo": True, "texto": "Título", "fontes": [],
               "contexto_citado": ["[P › A] conteúdo", "[Pedido] fatos"]}]
    for posicao in range(1, blocos_de_conteudo + 1):
        blocos.append({"bloco_id": f"{combinacao}-{posicao}", "posicao": posicao, "e_titulo": False,
                       "texto": f"Bloco {posicao}.", "fontes": ["P › A"],
                       "contexto_citado": ["[P › A] conteúdo", "[Pedido] fatos"]})
    return {"combinacao": combinacao, "empresa_id": "EMP003", "empresa": "Brisa Ltda.", "tipo": tipo,
            "canal": "email", "caso_do_destaque": "vazio", "conjunto": "um", "beneficios": ["A"], "destaque": "",
            "pergunta": "Quais são as dúvidas do time...?", "texto_sem_fontes": "Título\nBloco 1.",
            "contexto_do_material": ["[P › A] conteúdo", "[P › B] outro", "[Pedido] fatos"],
            "catalogo_confere": confere, "por_que_nao_confere": "", "blocos": blocos}


def test_sorteio_dos_rotulos_e_fixo_e_pula_titulos_e_catalogo_diferente():
    """A mesma semente dá os mesmos blocos; título e material com o catálogo diferente nunca entram."""
    materiais = [_material("C1", 30), _material("C2", 30), _material("C3", 30, confere=False)]
    primeira = ragas.sortear_blocos_para_rotular(materiais)
    segunda = ragas.sortear_blocos_para_rotular(materiais)
    assert primeira == segunda and len(primeira) == ragas.BLOCOS_PARA_ROTULAR
    for linha in primeira:
        assert not linha["bloco_id"].endswith("-0") and not linha["bloco_id"].startswith("C3")
    assert [linha["numero"] for linha in primeira[:3]] == [1, 2, 3]
    # Com menos candidatos que 50, entram todos
    assert len(ragas.sortear_blocos_para_rotular([_material("C1", 4)])) == 4


def test_planilha_grava_e_le_os_rotulos(tmp_path):
    """A planilha sai com a lista das 3 opções; o que se escreve na coluna do rótulo volta, com as variações."""
    from openpyxl import load_workbook
    caminho = tmp_path / "rotulos.xlsx"
    linhas = ragas.sortear_blocos_para_rotular([_material("C1", 3)])
    ragas.gravar_planilha_dos_rotulos(linhas, caminho)
    livro = load_workbook(caminho)
    aba = livro["Rotulos"]
    assert aba.data_validations.dataValidation[0].formula1 == '"não fiel,em parte,fiel"'
    assert "Como rotular" in livro.sheetnames
    # Sem rótulo ainda: nada volta
    assert ragas.ler_rotulos(caminho) == {}
    aba.cell(row=2, column=ragas.COLUNA_DO_ROTULO, value=" Fiel ")
    aba.cell(row=3, column=ragas.COLUNA_DO_ROTULO, value="nao fiel")
    aba.cell(row=3, column=ragas.COLUNA_DO_COMENTARIO, value="inventou o prazo")
    aba.cell(row=4, column=ragas.COLUNA_DO_ROTULO, value="talvez")
    livro.save(caminho)
    rotulos = ragas.ler_rotulos(caminho)
    assert rotulos == {linhas[0]["bloco_id"]: {"rotulo": "fiel", "comentario": ""},
                       linhas[1]["bloco_id"]: {"rotulo": "não fiel", "comentario": "inventou o prazo"}}


# ---------------- A etapa 2: o juiz (com a chamada falsa) ----------------

class RespostaFalsa:
    """A chamada falsa ao Bedrock: devolve o texto combinado e anota o que recebeu (nada sai desta máquina)."""

    def __init__(self, texto: str, tokens_entrada: int = 1000, tokens_saida: int = 200):
        """Recebe o texto que o "juiz" vai responder e os tokens que a resposta conta."""
        self.texto = texto
        self.tokens_entrada = tokens_entrada
        self.tokens_saida = tokens_saida
        self.pedidos = []

    def chamar(self, modelo: str, prompt: str, esquema_json: dict) -> RespostaDoProvedor:
        """Anota o pedido e devolve a resposta no formato do provedor."""
        self.pedidos.append({"modelo": modelo, "prompt": prompt, "esquema": esquema_json})
        return RespostaDoProvedor(texto=self.texto, tokens_entrada=self.tokens_entrada,
                                  tokens_saida=self.tokens_saida)


class Afirmacoes(BaseModel):
    """Um formato de resposta como os do RAGAS: a lista de afirmações (como a StatementGeneratorOutput)."""

    statements: list[str]


def test_juiz_confere_o_formato_soma_o_custo_e_guarda_as_respostas():
    """O juiz manda o esquema do formato, lê a resposta nele, guarda a resposta e soma o custo na caixa."""
    falsa = RespostaFalsa('{"statements": ["A conta não tem tarifa."]}')
    caixa = ragas.CaixaDoGasto(1.0)
    juiz = ragas.JuizDoBedrock("mistral-large-3", caixa, chamar=falsa.chamar)
    resposta = juiz.generate("o prompt do RAGAS", Afirmacoes)
    assert resposta.statements == ["A conta não tem tarifa."] and juiz.respostas == [resposta]
    assert falsa.pedidos[0]["esquema"]["title"] == "Afirmacoes"
    # O custo sai da tabela de preços do projeto (o Mistral Large 3: 0,50 e 1,50 por milhão; +10% na rota do Bedrock)
    custo_esperado = custo_em_dolares("mistral-large-3", 1000, 200)
    assert math.isclose(caixa.gasto_usd, custo_esperado) and math.isclose(juiz.custo_usd, custo_esperado)
    assert caixa.reservado_usd == 0 and caixa.chamadas == 1


def test_juiz_para_no_teto_sem_chamar():
    """Se a reserva de uma chamada passaria do teto, o juiz nem chama a IA."""
    falsa = RespostaFalsa('{"statements": []}')
    juiz = ragas.JuizDoBedrock("mistral-large-3", ragas.CaixaDoGasto(0.01), chamar=falsa.chamar)
    with pytest.raises(ragas.TetoAtingido):
        juiz.generate("prompt", Afirmacoes)
    assert falsa.pedidos == []


def test_juiz_com_resposta_fora_do_formato_levanta_e_cobra():
    """A resposta torta vira RespostaForaDoFormato, e o custo da chamada entra na caixa (o Bedrock cobrou)."""
    caixa = ragas.CaixaDoGasto(1.0)
    juiz = ragas.JuizDoBedrock("mistral-large-3", caixa, chamar=RespostaFalsa("não é JSON").chamar)
    with pytest.raises(ragas.RespostaForaDoFormato):
        juiz.generate("prompt", Afirmacoes)
    assert caixa.gasto_usd > 0 and caixa.reservado_usd == 0 and juiz.respostas == []


# ---------------- A etapa 2: a medição (com as métricas falsas) ----------------

class MetricasFalsas:
    """As métricas no lugar das do RAGAS: respostas combinadas por texto do bloco, com um custo fixo por medição.

    vereditos_por_texto: {texto do bloco: [1, 0, ...]} (o veredito de cada afirmação contra a fonte citada);
    no_catalogo: o veredito da 2ª conferência (contra todos os trechos), um por afirmação sem sustentação;
    teto_depois_de: quantas medições até o teto (None = sem teto); erro_no_texto: o bloco em que o juiz falha.
    """

    def __init__(self, vereditos_por_texto: dict, no_catalogo: list[int] | None = None,
                 teto_depois_de: int | None = None, erro_no_texto: str = ""):
        """Guarda as respostas combinadas."""
        self.vereditos_por_texto = vereditos_por_texto
        self.no_catalogo = no_catalogo or []
        self.teto_depois_de = teto_depois_de
        self.erro_no_texto = erro_no_texto
        self.medicoes = 0
        # As perguntas que a fidelidade recebeu (a medição usa sempre a pergunta neutra)
        self.perguntas_da_fidelidade = []

    def _contar(self) -> dict:
        """Conta uma medição; passa do teto combinado → TetoAtingido. Devolve o uso de uma medição."""
        self.medicoes += 1
        if self.teto_depois_de is not None and self.medicoes > self.teto_depois_de:
            raise ragas.TetoAtingido("teto")
        return {"custo_usd": 0.001, "tokens_entrada": 100, "tokens_saida": 20}

    async def fidelidade(self, pergunta: str, texto: str, contextos: list[str]) -> dict:
        """As afirmações do bloco com os vereditos combinados."""
        self.perguntas_da_fidelidade.append(pergunta)
        uso = self._contar()
        if texto == self.erro_no_texto:
            raise ragas.RespostaForaDoFormato("torta")
        afirmacoes = []
        for numero, veredito in enumerate(self.vereditos_por_texto.get(texto, [])):
            afirmacoes.append({"afirmacao": f"{texto} #{numero}", "veredito": veredito, "motivo": "m"})
        valor = None
        if afirmacoes:
            valor = sum(self.vereditos_por_texto[texto]) / len(afirmacoes)
        medida = {"valor": valor, "afirmacoes": afirmacoes}
        medida.update(uso)
        return medida

    async def vereditos(self, afirmacoes: list[str], contextos: list[str]) -> dict:
        """A 2ª conferência: os vereditos combinados, na ordem das afirmações."""
        uso = self._contar()
        lista = []
        for afirmacao, veredito in zip(afirmacoes, self.no_catalogo):
            lista.append({"afirmacao": afirmacao, "veredito": veredito, "motivo": "no catálogo"})
        medida = {"afirmacoes": lista}
        medida.update(uso)
        return medida

    async def relevancia(self, pergunta: str, texto: str) -> dict:
        """Uma relevância fixa."""
        uso = self._contar()
        medida = {"valor": 0.8, "perguntas": ["Quais benefícios?"] * 3, "evasivo": False}
        medida.update(uso)
        return medida


def test_julgar_conta_as_afirmacoes_sustentadas_citadas_errado_e_inventadas():
    """Bloco com 3 afirmações (1 na fonte; das 2 sem, 1 em outro trecho e 1 inventada); título com 1 de 2 na fonte."""
    metricas = MetricasFalsas({"Título": [1, 0], "Bloco 1.": [1, 0, 0]}, no_catalogo=[1, 0])
    blocos, materiais = ragas.julgar([_material("C1", 1)], metricas)
    titulo, bloco = blocos
    # O título já foi conferido contra todos os trechos: a afirmação sem sustentação é inventada, sem 2ª conferência
    assert (titulo["situacao"], titulo["fidelidade"], titulo["rotulo_do_juiz"]) == ("medido", 0.5, "em parte")
    assert (titulo["inventadas"], titulo["sem_segundo_veredito"]) == (1, 0)
    # A fidelidade recebe sempre a pergunta neutra, nunca a pergunta do material
    assert metricas.perguntas_da_fidelidade == [ragas.PERGUNTA_DA_FIDELIDADE] * 2
    assert (bloco["afirmacoes"], bloco["sustentadas_pela_fonte"], bloco["sustentadas_em_outro_trecho"],
            bloco["inventadas"]) == (3, 1, 1, 1)
    assert math.isclose(bloco["fidelidade"], 1 / 3) and math.isclose(bloco["fidelidade_ao_catalogo"], 2 / 3)
    assert bloco["rotulo_do_juiz"] == "em parte" and bloco["detalhes"][2]["veredito_no_catalogo"] == 0
    material = materiais[0]
    # As afirmações somadas dos blocos de conteúdo: 1 de 3 na fonte citada, 2 de 3 no catálogo; o título à parte
    assert (material["situacao"], material["afirmacoes"], material["blocos"]) == ("medido", 3, 1)
    assert math.isclose(material["fidelidade_do_material"], 1 / 3)
    assert math.isclose(material["fidelidade_ao_catalogo"], 2 / 3) and material["inventadas"] == 1
    assert (material["titulo_afirmacoes"], material["titulo_inventadas"]) == (2, 1)
    assert (material["blocos_fieis"], material["blocos_em_parte"], material["relevancia"]) == (0, 1, 0.8)
    # Custo: 2 medições no bloco (fidelidade e 2ª conferência), 1 no título e 1 da relevância
    assert math.isclose(material["custo_usd"], 0.004)


def test_julgar_bloco_sem_afirmacao_e_material_com_catalogo_diferente():
    """Bloco sem afirmação não conta na fidelidade; o material com o catálogo diferente nem é julgado."""
    metricas = MetricasFalsas({"Título": [], "Bloco 1.": [1]})
    blocos, materiais = ragas.julgar([_material("C1", 1), _material("C2", 1, confere=False)], metricas)
    assert blocos[0]["situacao"] == "sem_afirmacao" and blocos[0]["fidelidade"] is None
    assert len(materiais) == 1 and materiais[0]["fidelidade_do_material"] == 1.0


def test_julgar_para_no_teto_e_erro_do_juiz_nao_perde_a_medicao():
    """Depois do teto, nenhum bloco chama o juiz; um erro num bloco marca só aquele bloco."""
    metricas = MetricasFalsas({"Título": [1], "Bloco 1.": [1], "Bloco 2.": [1]}, erro_no_texto="Bloco 1.",
                              teto_depois_de=3)
    blocos, materiais = ragas.julgar([_material("C1", 2)], metricas, medicoes_ao_mesmo_tempo=1)
    situacoes = [bloco["situacao"] for bloco in blocos]
    assert situacoes == ["medido", "erro_do_juiz", "medido"]
    assert "RespostaForaDoFormato" in blocos[1]["erro"]
    # A relevância chegou ao teto: o material fica incompleto
    assert materiais[0]["situacao"] == "incompleto" and materiais[0]["erro"] == "parado_pelo_teto"
    metricas_no_teto = MetricasFalsas({"Título": [1], "Bloco 1.": [1]}, teto_depois_de=0)
    blocos_no_teto, _ = ragas.julgar([_material("C1", 1)], metricas_no_teto, medicoes_ao_mesmo_tempo=1)
    assert [bloco["situacao"] for bloco in blocos_no_teto] == ["parado_pelo_teto", "parado_pelo_teto"]
    assert metricas_no_teto.medicoes == 1


def test_csv_dos_blocos_e_dos_materiais_ida_e_volta(tmp_path):
    """O que se grava no CSV volta com os mesmos tipos (números, sim/não, listas e o None dos não medidos)."""
    blocos, materiais = ragas.julgar([_material("C1", 1)], MetricasFalsas({"Título": [], "Bloco 1.": [1, 0]},
                                                                         no_catalogo=[0]))
    ragas.gravar_csv(blocos, ragas.COLUNAS_DOS_BLOCOS, tmp_path / "blocos.csv")
    ragas.gravar_csv(materiais, ragas.COLUNAS_DOS_MATERIAIS, tmp_path / "materiais.csv")
    blocos_lidos = ragas.ler_csv(tmp_path / "blocos.csv")
    assert blocos_lidos[0]["fidelidade"] is None and blocos_lidos[0]["e_titulo"] is True
    assert blocos_lidos[1]["detalhes"] == blocos[1]["detalhes"] and blocos_lidos[1]["fidelidade"] == 0.5
    assert ragas.ler_csv(tmp_path / "materiais.csv")[0]["perguntas"] == ["Quais benefícios?"] * 3


def test_rotulo_da_fidelidade_e_rotulo_escrito():
    """A régua de 3 casas e as variações que quem rotula pode escrever na planilha."""
    assert ragas.rotulo_da_fidelidade(1.0) == "fiel" and ragas.rotulo_da_fidelidade(0.0) == "não fiel"
    assert ragas.rotulo_da_fidelidade(0.5) == "em parte" and ragas.rotulo_da_fidelidade(None) == ""
    assert ragas.rotulo_escrito(" Não Fiel ") == "não fiel" and ragas.rotulo_escrito("EM PARTE") == "em parte"
    assert ragas.rotulo_escrito(None) == "" and ragas.rotulo_escrito("talvez") == ""


# ---------------- A etapa 3: o resumo e o kappa ----------------

def test_kappa_de_cohen_com_as_contas_feitas_a_mao():
    """Os valores calculados no papel: kappa simples 0,5833 (8 de 10 iguais) e o ponderado 0,7143 (a régua com ordem)."""
    pares = ([("fiel", "fiel")] * 5 + [("não fiel", "não fiel")] * 3 + [("fiel", "não fiel"), ("não fiel", "fiel")])
    # Observada 0,8; esperada 0,6 × 0,6 + 0,4 × 0,4 = 0,52; kappa = 0,28 ÷ 0,48
    assert math.isclose(resumo_do_ragas.kappa_de_cohen(pares), 0.28 / 0.48)
    com_em_parte = [("fiel", "fiel"), ("fiel", "em parte"), ("não fiel", "não fiel"), ("em parte", "em parte")]
    # Simples: observada 0,75, esperada 0,3125; ponderado: observada 0,875, esperada 0,5625
    assert math.isclose(resumo_do_ragas.kappa_de_cohen(com_em_parte), 0.4375 / 0.6875)
    assert math.isclose(resumo_do_ragas.kappa_de_cohen(com_em_parte, ponderado=True), 0.3125 / 0.4375)
    # Concordância perfeita, e os casos em que não dá para calcular
    assert resumo_do_ragas.kappa_de_cohen([("fiel", "fiel"), ("não fiel", "não fiel")]) == 1.0
    assert resumo_do_ragas.kappa_de_cohen([("fiel", "fiel")] * 4) is None
    assert resumo_do_ragas.kappa_de_cohen([]) is None


def test_leitura_intervalo_e_matriz_do_kappa():
    """As faixas de Landis e Koch, o intervalo por bootstrap (sempre o mesmo) e a matriz de confusão."""
    assert resumo_do_ragas.leitura_do_kappa(0.7) == "substancial"
    assert resumo_do_ragas.leitura_do_kappa(-0.1) == "pior que o acaso"
    pares = [("fiel", "fiel")] * 6 + [("não fiel", "não fiel")] * 3 + [("em parte", "fiel")]
    intervalo = resumo_do_ragas.intervalo_do_kappa(pares)
    assert intervalo == resumo_do_ragas.intervalo_do_kappa(pares) and intervalo[0] <= intervalo[1]
    matriz = resumo_do_ragas.matriz_de_confusao(pares)
    assert matriz["fiel"]["fiel"] == 6 and matriz["em parte"]["fiel"] == 1 and matriz["fiel"]["não fiel"] == 0


def test_resumir_com_e_sem_os_rotulos():
    """O resumo junta as médias, os grupos e, quando há rótulos, a concordância com os 50 (o bloco sem rótulo do juiz
    fica fora)."""
    metricas = MetricasFalsas({"Título": [1], "Bloco 1.": [1, 0], "Bloco 2.": [1]}, no_catalogo=[0])
    amostras = [_material("C1", 2), _material("C2", 2, tipo="comunicado"), _material("C3", 1, confere=False)]
    amostras[2]["por_que_nao_confere"] = "a fonte X não está no catálogo de hoje"
    blocos, materiais = ragas.julgar(amostras, metricas)
    sem_rotulos = resumo_do_ragas.resumir(amostras, blocos, materiais)
    assert sem_rotulos["materiais_medidos"] == 2
    assert sem_rotulos["concordancia_com_os_rotulos_dos_50"].startswith("pendente")
    assert sem_rotulos["materiais_com_catalogo_diferente"][0]["combinacao"] == "C3"
    # Por material, os blocos de conteúdo: 2 de 3 afirmações na fonte (o título, com 1 de 1, conta à parte)
    assert sem_rotulos["geral"]["inventadas"] == 2 and sem_rotulos["geral"]["fidelidade_somada"] == 0.6667
    assert sem_rotulos["geral"]["titulos_com_inventada"] == 0
    assert set(sem_rotulos["por_tipo"]) == {"comunicado", "faq"}
    assert sem_rotulos["divergencias_entre_o_ragas_e_a_contagem"] == 0
    assert sem_rotulos["blocos"]["blocos_de_conteudo_por_rotulo"] == {"não fiel": 0, "em parte": 2, "fiel": 2}
    dos_50 = {"C1-1": {"rotulo": "em parte", "comentario": ""}, "C1-2": {"rotulo": "não fiel", "comentario": "x"},
              "C9-1": {"rotulo": "fiel", "comentario": ""}}
    com_rotulos = resumo_do_ragas.resumir(amostras, blocos, materiais, dos_50,
                                          [{"bloco_id": "C1-1", "concorda": False, "nota": "n"}])
    # Os 50 rótulos (de outro avaliador) contra o juiz: o C9-1 o juiz não rotulou e fica fora do kappa
    concordancia = com_rotulos["concordancia_com_os_rotulos_dos_50"]
    assert (concordancia["pares"], concordancia["sem_rotulo_do_juiz"]) == (2, ["C9-1"])
    assert concordancia["discordancias"] == [{"bloco_id": "C1-2", "juiz": "fiel", "outro": "não fiel",
                                              "comentario": "x"}]
    assert concordancia["fieis_pelo_outro"]["fieis"] == 1 and concordancia["fieis_pelo_outro"]["total"] == 3
    assert concordancia["concordancia_simples"] == 0.5
    assert com_rotulos["conferencia_do_agente"] == {"casos": 1, "concorda": 0,
                                                    "discorda": [{"bloco_id": "C1-1", "nota": "n"}]}


def test_taxa_de_fieis_com_o_intervalo_de_wilson():
    """41 de 50 fiéis → 0,82, com o intervalo de Wilson de cerca de 0,69 a 0,90; sem rótulos, a taxa fica None."""
    rotulos = {}
    for numero in range(50):
        rotulo = "fiel"
        if numero >= 41:
            rotulo = "em parte"
        rotulos[f"C1-{numero}"] = {"rotulo": rotulo, "comentario": ""}
    taxa = resumo_do_ragas.taxa_de_fieis(rotulos)
    assert (taxa["fieis"], taxa["total"], taxa["taxa"]) == (41, 50, 0.82)
    assert math.isclose(taxa["ic95_wilson"][0], 0.692, abs_tol=0.001)
    assert math.isclose(taxa["ic95_wilson"][1], 0.902, abs_tol=0.001)
    assert resumo_do_ragas.taxa_de_fieis({})["taxa"] is None


def test_calibracao_pelo_catalogo_usa_a_fidelidade_ao_catalogo():
    """A análise secundária troca o rótulo do juiz pelo da fidelidade ao catálogo; o bloco sem rótulo fica de fora."""
    blocos = [{"bloco_id": "C1-1", "rotulo_do_juiz": "em parte", "fidelidade_ao_catalogo": 1.0},
              {"bloco_id": "C1-2", "rotulo_do_juiz": "fiel", "fidelidade_ao_catalogo": 1.0},
              {"bloco_id": "C1-3", "rotulo_do_juiz": "", "fidelidade_ao_catalogo": None}]
    rotulos = {"C1-1": {"rotulo": "fiel", "comentario": ""}, "C1-2": {"rotulo": "fiel", "comentario": ""},
               "C1-3": {"rotulo": "fiel", "comentario": ""}}
    pela_fonte = resumo_do_ragas.calibracao(blocos, rotulos)
    pelo_catalogo = resumo_do_ragas.calibracao(blocos, rotulos, pelo_catalogo=True)
    assert [discordancia["bloco_id"] for discordancia in pela_fonte["discordancias"]] == ["C1-1"]
    assert pelo_catalogo["discordancias"] == [] and pelo_catalogo["sem_rotulo_do_juiz"] == ["C1-3"]
    assert pelo_catalogo["concordancia_simples"] == 1.0 and pela_fonte["concordancia_simples"] == 0.5


def test_sorteio_da_conferencia_do_agente_metade_de_cada_lado_e_fora_da_planilha():
    """Metade dos blocos com problema e metade fiéis, sem título, sem erro e sem os blocos da planilha, sempre iguais."""
    blocos = []
    for numero in range(1, 21):
        rotulo = "em parte"
        if numero % 2 == 0:
            rotulo = "fiel"
        blocos.append({"bloco_id": f"C1-{numero}", "e_titulo": False, "situacao": "medido", "rotulo_do_juiz": rotulo})
    blocos.append({"bloco_id": "C1-0", "e_titulo": True, "situacao": "medido", "rotulo_do_juiz": "em parte"})
    blocos.append({"bloco_id": "C2-1", "e_titulo": False, "situacao": "erro_do_juiz", "rotulo_do_juiz": ""})
    na_planilha = {"C1-1", "C1-2"}
    primeira = ragas.sortear_para_a_conferencia_do_agente(blocos, na_planilha)
    assert primeira == ragas.sortear_para_a_conferencia_do_agente(blocos, na_planilha)
    assert [bloco["rotulo_do_juiz"] for bloco in primeira] == ["em parte"] * 5 + ["fiel"] * 5
    for bloco in primeira:
        assert bloco["bloco_id"] not in {"C1-0", "C2-1", "C1-1", "C1-2"}
