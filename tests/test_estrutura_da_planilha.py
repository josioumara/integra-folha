"""A estrutura da planilha, tratada por regra (ADR-90): linha de total, sem cabeçalho, células mescladas e abas vazias.

O que estes testes provam:
- a linha de total no fim ("Total" e a soma) não vira funcionário, e a empresa é avisada;
- a planilha sem cabeçalho (a primeira linha já é uma pessoa) não perde essa pessoa: as colunas viram "Coluna N";
- o valor de uma célula mesclada (ex.: a unidade em 3 linhas) vale para todas as linhas do grupo;
- com a primeira aba vazia (uma capa), a primeira aba com conteúdo é a lida.
"""
import io

from openpyxl import Workbook

from services import ingestao


def planilha(linhas: list[list], mesclar: list[str] | None = None, aba_vazia_antes: bool = False) -> bytes:
    """Uma planilha .xlsx feita na hora, com as linhas dadas (e, se pedido, células mescladas e uma capa vazia)."""
    livro = Workbook()
    aba = livro.active
    if aba_vazia_antes:
        aba.title = "Capa"
        aba = livro.create_sheet("Funcionarios")
    for linha in linhas:
        aba.append(linha)
    for faixa in mesclar or []:
        aba.merge_cells(faixa)
    memoria = io.BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


def test_linha_de_total_fica_de_fora_com_aviso():
    conteudo = planilha([["Nome", "CPF", "Salário"], ["Ana Souza", "52998224725", 3000],
                         ["Bia Lima", "11144477735", 2500], ["Total", "", 5500]])
    leitura = ingestao.ler_arquivo(conteudo, "lista.xlsx")
    assert [linha[0] for linha in leitura.linhas] == ["Ana Souza", "Bia Lima"]
    assert "A linha 4 parece uma linha de total e ficou de fora." in leitura.avisos


def test_planilha_sem_cabecalho_nao_perde_a_primeira_pessoa():
    conteudo = planilha([["Ana Souza", "52998224725", "Analista"], ["Bia Lima", "11144477735", "Vendedora"]])
    leitura = ingestao.ler_arquivo(conteudo, "lista.xlsx")
    assert leitura.cabecalhos == ["Coluna 1", "Coluna 2", "Coluna 3"]
    assert [linha[0] for linha in leitura.linhas] == ["Ana Souza", "Bia Lima"]
    assert leitura.numeros_linha == [1, 2]
    assert any("não tem linha de cabeçalho" in aviso for aviso in leitura.avisos)


def test_celula_mesclada_vale_para_todas_as_linhas_do_grupo():
    conteudo = planilha([["Nome", "CPF", "Unidade"], ["Ana Souza", "52998224725", "Fábrica Campinas"],
                         ["Bia Lima", "11144477735", None], ["Caio Reis", "12345678909", None]], mesclar=["C2:C4"])
    leitura = ingestao.ler_arquivo(conteudo, "lista.xlsx")
    assert [linha[2] for linha in leitura.linhas] == ["Fábrica Campinas"] * 3
    assert any("células mescladas" in aviso for aviso in leitura.avisos)


def test_aba_vazia_antes_da_lista_e_pulada():
    conteudo = planilha([["Nome", "CPF"], ["Ana Souza", "52998224725"]], aba_vazia_antes=True)
    leitura = ingestao.ler_arquivo(conteudo, "lista.xlsx")
    assert leitura.cabecalhos == ["Nome", "CPF"]
    assert "A planilha tem 2 abas; só a aba Funcionarios foi lida." in leitura.avisos
    assert "As abas antes dela estavam vazias." in leitura.avisos


def test_regras_da_linha_de_total_e_da_linha_de_dado():
    assert ingestao.linha_parece_de_total(["Total geral:", "", "35.200,00"]) is True
    assert ingestao.linha_parece_de_total(["TOTAIS", "12"]) is True
    # Uma pessoa com CPF nunca é linha de total, mesmo com a palavra na linha
    assert ingestao.linha_parece_de_total(["Total", "52998224725"]) is False
    assert ingestao.linha_parece_de_total(["Maria Total", "Analista"]) is False
    assert ingestao.linha_parece_dado(["Ana Souza", "52998224725"]) is True
    assert ingestao.linha_parece_dado(["Admissão", "02/09/2026"]) is True
    assert ingestao.linha_parece_dado(["Nome", "CPF", "Admissão"]) is False
