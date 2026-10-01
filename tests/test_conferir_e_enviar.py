"""Conferir e enviar: a lista final, o número único das pendências e o envio parcial (ADR-126).

O que estes testes provam (com a IA simulada, sem custo; as planilhas são montadas aqui, com pessoas inventadas):
- a lista que vai para o banco mostra o valor de cada campo, venha ele de uma coluna do arquivo ou dos cartões
  ("para todos" ou pessoa a pessoa); o campo que ninguém informou continua fora (ADR-126);
- o número das pendências é um só: o resumo da leitura (painel "Linhas do arquivo"), os contadores da conferência e
  os cartões de "Acompanhar cadastros" dão o mesmo total, contando as informações que faltam no arquivo inteiro
  (ADR-126);
- envio parcial: quem já foi ao banco, noutro envio ainda em análise, ou já está cadastrado fica de
  fora do envio novo, com o motivo; o envio segue com as outras pessoas; o arquivo em que todos já foram mandados antes
  é recusado com o porquê; o mesmo CPF duas vezes no arquivo continua pendência; quem o banco devolveu e a empresa tirou
  da devolução pode voltar; o envio que já está com o banco não é travado pelo arquivo novo.
Cada caso tem variações (outra posição, outra escrita do CPF, outro dado mudado, outras quantidades), para a regra
valer em situações novas e não só no caso que achou o defeito.
"""
import io

import openpyxl
import pytest

from models.contratos import EstadoProcessamento
from services import (acompanhamento, assistente_na_tela, avaliacao_do_banco, cadastro, correcoes, empresas, ingestao,
                      parametros, processamentos, validador)
from tests.test_avaliacao_do_banco import ESPECIALISTA, enviar_a_aurora_ao_banco
from tests.test_correcao import busca_falsa
from tests.test_fluxo_empresa import (ENVIOS, _gabarito, aprovar_no_banco, conexao, escolhas_do_gabarito,  # noqa: F401
                                      gerar_envios, verdade)
from tests.test_pendencias_em_grupo import cpf_valido

# A empresa dos testes, quem clica e o arquivo de onde as planilhas novas são montadas
EMPRESA = "EMP001"
LOGIN = "rh.aurora"
ARQUIVO_BASE = "aurora_carga_inicial"
# Pessoas inventadas para as linhas novas (nome, os 9 primeiros dígitos do CPF e a matrícula)
PESSOAS_NOVAS = [("Helena Prado", "902113554", "770001"), ("Otavio Cunha", "903224665", "770002"),
                 ("Livia Barros", "904335776", "770003"), ("Renato Faria", "905446887", "770004"),
                 ("Marta Quintela", "906557998", "770005")]


@pytest.fixture(autouse=True)
def usar_busca_falsa(monkeypatch):
    """Nenhum uso do RAG depende do índice nem do modelo de embeddings."""
    monkeypatch.setattr("rag.busca.search_rules", busca_falsa)


# ---------------- Planilhas montadas no teste ----------------

def _colunas_do_arquivo_base() -> dict:
    """A coluna (posição, 1 = primeira) de cada campo que a planilha nova mexe: cpf, matricula e nome_completo."""
    gabarito = _gabarito(ARQUIVO_BASE)
    planilha = openpyxl.load_workbook(ENVIOS / gabarito["arquivo"])
    aba = planilha.active
    linha_do_cabecalho = gabarito["formato"]["linha_do_cabecalho"]
    colunas = {}
    for celula in aba[linha_do_cabecalho]:
        campo = gabarito["mapeamento"].get(celula.value)
        if campo in ("cpf", "matricula", "nome_completo"):
            colunas[campo] = celula.column
    return colunas


def _linhas_do_arquivo_base() -> tuple[list, list[list]]:
    """O cabeçalho e as linhas de dados da planilha da Aurora (listas de valores)."""
    gabarito = _gabarito(ARQUIVO_BASE)
    aba = openpyxl.load_workbook(ENVIOS / gabarito["arquivo"]).active
    linha_do_cabecalho = gabarito["formato"]["linha_do_cabecalho"]
    cabecalho = []
    for celula in aba[linha_do_cabecalho]:
        cabecalho.append(celula.value)
    linhas = []
    for linha in aba.iter_rows(min_row=linha_do_cabecalho + 1, values_only=True):
        if any(valor not in (None, "") for valor in linha):
            linhas.append(list(linha))
    return cabecalho, linhas


def _so_digitos(valor) -> str:
    """Só os dígitos de um valor. Ex.: "529.982.247-25" → "52998224725"."""
    digitos = ""
    for caractere in str(valor or ""):
        if caractere.isdigit():
            digitos = digitos + caractere
    return digitos


def _linhas_ja_enviadas(cpfs_enviados: set[str]) -> list[list]:
    """As linhas da planilha base cujo CPF está entre os já enviados (a linha com o CPF errado fica de fora)."""
    posicao_do_cpf = _colunas_do_arquivo_base()["cpf"] - 1
    _, linhas = _linhas_do_arquivo_base()
    escolhidas = []
    for linha in linhas:
        if _so_digitos(linha[posicao_do_cpf]) in cpfs_enviados:
            escolhidas.append(linha)
    return escolhidas


def _pessoa_nova(linha_modelo: list, indice: int, cpf_com_pontos: bool = True) -> list:
    """Uma linha nova a partir de outra: mesmos dados de cargo e endereço, outra pessoa (nome, CPF e matrícula)."""
    colunas = _colunas_do_arquivo_base()
    nome, nove_digitos, matricula = PESSOAS_NOVAS[indice]
    cpf = cpf_valido(nove_digitos)
    if cpf_com_pontos:
        cpf = cpf[:3] + "." + cpf[3:6] + "." + cpf[6:9] + "-" + cpf[9:]
    nova = list(linha_modelo)
    nova[colunas["nome_completo"] - 1] = nome
    nova[colunas["cpf"] - 1] = cpf
    nova[colunas["matricula"] - 1] = matricula
    return nova


def _planilha(linhas: list[list]) -> bytes:
    """Uma planilha .xlsx com o cabeçalho da Aurora e as linhas informadas."""
    cabecalho, _ = _linhas_do_arquivo_base()
    planilha = openpyxl.Workbook()
    aba = planilha.active
    aba.append(cabecalho)
    for linha in linhas:
        aba.append(linha)
    saida = io.BytesIO()
    planilha.save(saida)
    return saida.getvalue()


def _enviar_e_aceitar(conexao, conteudo: bytes, nome_do_arquivo: str) -> str:
    """Envia a planilha, aceita as colunas e confirma os alertas de salário (se houver). Devolve o envio."""
    leitura = cadastro.enviar_arquivo(conexao, EMPRESA, LOGIN, conteudo, nome_do_arquivo, busca=busca_falsa)
    processamento_id = leitura["processamento_id"]
    cadastro.aceitar_mapeamento(conexao, EMPRESA, LOGIN, processamento_id, escolhas_do_gabarito(ARQUIVO_BASE),
                                busca=busca_falsa)
    # Os alertas (ex.: salário longe dos colegas, num arquivo pequeno) são confirmados, como a empresa faria
    for achado in validador.obter(conexao, processamento_id).achados:
        if achado.severidade == validador.ALERTA and not achado.resolvido:
            acompanhamento.confirmar_pendencia(conexao, EMPRESA, LOGIN, processamento_id, achado.regra_id,
                                               achado.linha, "Conferido no teste")
    return processamento_id


def _cpfs_do_envio(conexao, processamento_id: str) -> set[str]:
    """Os CPFs (só dígitos) das pessoas de um envio."""
    cpfs = set()
    for registro in correcoes.dados_atuais(conexao, processamento_id).registros:
        cpfs.add(registro["cpf"])
    return cpfs


def _regras_da_linha(conexao, processamento_id: str) -> dict[int, set[str]]:
    """{linha: as regras que o Validador achou nela}."""
    regras = {}
    for achado in validador.obter(conexao, processamento_id).achados:
        regras.setdefault(achado.linha, set()).add(achado.regra_id)
    return regras


# ---------------- Envio parcial: quem já foi ao banco fica de fora ----------------

@pytest.mark.parametrize("posicoes_das_ja_enviadas, total_de_linhas, cpf_com_pontos", [
    ((0, 1), 5, True),        # as duas primeiras já foram; três pessoas novas depois
    ((2,), 4, False),         # uma no meio; CPF das novas só com dígitos
    ((3, 4), 5, True),        # as duas últimas
])
def test_quem_ja_foi_ao_banco_fica_de_fora_e_o_envio_segue_com_as_outras(conexao, verdade, posicoes_das_ja_enviadas,
                                                                           total_de_linhas, cpf_com_pontos):
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    ja_enviadas = _linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))
    # A planilha nova: as já enviadas nas posições pedidas, pessoas novas no resto
    linhas = []
    indice_da_nova = 0
    indice_da_enviada = 0
    for posicao in range(total_de_linhas):
        if posicao in posicoes_das_ja_enviadas:
            linhas.append(ja_enviadas[indice_da_enviada])
            indice_da_enviada = indice_da_enviada + 1
        else:
            linhas.append(_pessoa_nova(ja_enviadas[0], indice_da_nova, cpf_com_pontos))
            indice_da_nova = indice_da_nova + 1
    processamento_id = _enviar_e_aceitar(conexao, _planilha(linhas), "novos_e_repetidos.xlsx")
    # As linhas já enviadas (cabeçalho na linha 1) ficam de fora, e só com o aviso: nenhuma pendência para a empresa
    linhas_de_fora = set()
    for posicao in posicoes_das_ja_enviadas:
        linhas_de_fora.add(posicao + 2)
    regras = _regras_da_linha(conexao, processamento_id)
    for linha in linhas_de_fora:
        assert regras[linha] == {validador.JA_ENVIADO_AO_BANCO}
    # A lista que vai ao banco não tem essas pessoas; o bloco "fica de fora" diz quem e por quê
    lista = cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)
    linhas_da_lista = set()
    for linha in lista["linhas"]:
        linhas_da_lista.add(linha["linha"])
    assert linhas_da_lista.isdisjoint(linhas_de_fora)
    assert len(lista["linhas"]) == total_de_linhas - len(posicoes_das_ja_enviadas)
    assert len(lista["ficam_de_fora"]) == len(posicoes_das_ja_enviadas)
    for pessoa in lista["ficam_de_fora"]:
        assert "Já foi enviada ao banco" in pessoa["motivo"] and pessoa["nome"] and len(pessoa["cpf"]) == 14
    assert lista["contagem"]["ficam_de_fora"] == len(posicoes_das_ja_enviadas)
    # "Pronto para enviar": só as pessoas novas contam; o envio vai, e ninguém vai duas vezes
    prontos = cadastro.envios_prontos_para_o_banco(conexao, EMPRESA)
    assert [pronto["processamento_id"] for pronto in prontos] == [processamento_id]
    assert prontos[0]["pessoas"] == total_de_linhas - len(posicoes_das_ja_enviadas)
    resultado = cadastro.enviar_prontos_ao_banco(conexao, EMPRESA, LOGIN, conferiu_a_lista=True)
    assert resultado == {"envios": 1, "pessoas": total_de_linhas - len(posicoes_das_ja_enviadas),
                         "ficaram_de_fora": len(posicoes_das_ja_enviadas)}
    # O banco vê só as pessoas novas neste envio
    pessoas_no_banco = avaliacao_do_banco.pessoas_do_envio(conexao, ESPECIALISTA, processamento_id)
    assert len(pessoas_no_banco) == total_de_linhas - len(posicoes_das_ja_enviadas)


def test_quem_ja_foi_ao_banco_fica_de_fora_mesmo_com_outro_dado_mudado(conexao, verdade):
    """A mesma pessoa (mesmo CPF) com o nome em maiúsculas e outra matrícula continua sendo quem já foi ao banco."""
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    ja_enviada = list(_linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))[0])
    colunas = _colunas_do_arquivo_base()
    ja_enviada[colunas["nome_completo"] - 1] = str(ja_enviada[colunas["nome_completo"] - 1]).upper()
    ja_enviada[colunas["matricula"] - 1] = "880099"
    processamento_id = _enviar_e_aceitar(conexao, _planilha([_pessoa_nova(ja_enviada, 0), ja_enviada]),
                                         "mudou_o_nome.xlsx")
    assert _regras_da_linha(conexao, processamento_id)[3] == {validador.JA_ENVIADO_AO_BANCO}


def test_quem_ja_esta_cadastrado_fica_de_fora_com_o_motivo(conexao, verdade):
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    aprovar_no_banco(conexao, enviado, EMPRESA)
    cadastradas = _linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))
    processamento_id = _enviar_e_aceitar(conexao, _planilha([cadastradas[0], _pessoa_nova(cadastradas[0], 1)]),
                                         "inclusao.xlsx")
    lista = cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)
    assert len(lista["linhas"]) == 1 and len(lista["ficam_de_fora"]) == 1
    assert "já homologado" in lista["ficam_de_fora"][0]["motivo"]


def test_o_envio_com_o_banco_nao_e_travado_pelo_arquivo_novo_da_mesma_pessoa(conexao, verdade):
    """O arquivo novo, com a empresa, traz alguém do envio que espera o banco: o banco continua podendo aprovar."""
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    ja_enviadas = _linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))
    novo = _enviar_e_aceitar(conexao, _planilha([ja_enviadas[0], _pessoa_nova(ja_enviadas[0], 2)]), "depois.xlsx")
    assert validador.JA_ENVIADO_AO_BANCO in _regras_da_linha(conexao, novo)[2]
    # O banco aprova o primeiro envio sem pendência nova
    aprovar_no_banco(conexao, enviado, EMPRESA)
    assert processamentos.obter(conexao, enviado).status == EstadoProcessamento.HOMOLOGADO


# ---------------- O arquivo em que todos já foram mandados antes ----------------

@pytest.mark.parametrize("situacao", ["em_analise", "cadastradas", "uma_so"])
def test_arquivo_sem_ninguem_novo_e_recusado_com_o_porque(conexao, verdade, situacao):
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    if situacao == "cadastradas":
        aprovar_no_banco(conexao, enviado, EMPRESA)
    ja_mandadas = _linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))
    if situacao == "uma_so":
        ja_mandadas = ja_mandadas[:1]
    with pytest.raises(ingestao.ArquivoRecusado) as recusa:
        cadastro.enviar_arquivo(conexao, EMPRESA, LOGIN, _planilha(ja_mandadas), "de_novo.xlsx", busca=busca_falsa)
    mensagem = str(recusa.value)
    assert mensagem.startswith("Nada novo para enviar") and "Nada foi enviado" in mensagem
    if situacao == "cadastradas":
        assert f"{len(ja_mandadas)} já cadastrada(s)" in mensagem
    elif situacao == "uma_so":
        assert "a pessoa deste arquivo já foi mandada antes (1 em análise no banco)" in mensagem
    else:
        assert f"as {len(ja_mandadas)} pessoas" in mensagem and f"{len(ja_mandadas)} em análise no banco" in mensagem


def test_mensagem_do_arquivo_sem_ninguem_novo_junta_os_lugares():
    conhecidos = {"1": cadastro.JA_CADASTRADA, "2": cadastro.EM_ANALISE_NO_BANCO, "3": cadastro.EM_ANALISE_NO_BANCO,
                  "4": cadastro.EM_OUTRO_ARQUIVO}
    mensagem = cadastro.mensagem_de_nada_novo({"1", "2", "3", "4"}, conhecidos)
    assert "as 4 pessoas" in mensagem
    assert "(1 já cadastrada(s), 2 em análise no banco, 1 em outro arquivo que ainda não foi ao banco)" in mensagem


# ---------------- O mesmo CPF duas vezes; o devolvido que volta ----------------

def test_mesmo_cpf_duas_vezes_no_arquivo_continua_pendencia(conexao, verdade):
    """A pessoa ainda não foi ao banco: só a empresa sabe qual das duas linhas vale."""
    homologado = enviar_a_aurora_ao_banco(conexao, verdade)
    modelo = _linhas_ja_enviadas(_cpfs_do_envio(conexao, homologado))[0]
    nova = _pessoa_nova(modelo, 3)
    processamento_id = _enviar_e_aceitar(conexao, _planilha([nova, _pessoa_nova(modelo, 4), list(nova)]), "dup.xlsx")
    assert "PESSOA_DUPLICADA" in _regras_da_linha(conexao, processamento_id)[4]
    assert cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)["ficam_de_fora"] == []


def test_mesmo_cpf_duas_vezes_e_ja_enviado_fica_de_fora_nas_duas_linhas(conexao, verdade):
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    ja_enviada = _linhas_ja_enviadas(_cpfs_do_envio(conexao, enviado))[0]
    processamento_id = _enviar_e_aceitar(conexao, _planilha([ja_enviada, _pessoa_nova(ja_enviada, 0), ja_enviada]),
                                         "dup_enviada.xlsx")
    regras = _regras_da_linha(conexao, processamento_id)
    assert regras[2] == {validador.JA_ENVIADO_AO_BANCO} and regras[4] == {validador.JA_ENVIADO_AO_BANCO}


def test_devolvido_pelo_banco_nao_e_barrado_como_ja_enviado(conexao, verdade):
    """Enquanto a pessoa está no envio de devolução, é a pendência de sempre (ela fica num envio só); tirada da
    devolução, ela pode voltar num arquivo novo (ADR-121)."""
    enviado = enviar_a_aurora_ao_banco(conexao, verdade)
    pessoas = avaliacao_do_banco.pessoas_do_envio(conexao, ESPECIALISTA, enviado)
    apontada = pessoas[0]
    avaliacao_do_banco.apontar(conexao, ESPECIALISTA, enviado, apontada["linha"], "salario",
                               "O salário parece alto demais para o cargo. Pode conferir?")
    devolucao = avaliacao_do_banco.avaliar(conexao, ESPECIALISTA, enviado,
                                           decisao="aprovar_e_devolver_marcados")["envio_de_devolucao"]
    cpf_devolvido = _so_digitos(apontada["cpf"])
    linha_devolvida = _linhas_ja_enviadas({cpf_devolvido})[0]
    # Ainda na devolução: não é "já enviada"; é a mesma pessoa em dois envios com a empresa
    primeiro = _enviar_e_aceitar(conexao, _planilha([linha_devolvida, _pessoa_nova(linha_devolvida, 1)]), "volta.xlsx")
    regras = _regras_da_linha(conexao, primeiro)[2]
    assert "PESSOA_EM_OUTRO_ENVIO" in regras and validador.JA_ENVIADO_AO_BANCO not in regras
    # A empresa tira a pessoa da devolução: no arquivo novo, ela vai
    retirada = correcoes.propor(conexao, devolucao, EMPRESA, apontada["linha"], correcoes.EXCLUIR, None,
                                "Vai no arquivo novo", LOGIN)
    assistente_na_tela.confirmar_retirada(conexao, EMPRESA, LOGIN, devolucao, retirada.correcao_id, True)
    validador.executar(conexao, primeiro, EMPRESA)
    assert not (_regras_da_linha(conexao, primeiro).get(2, set()) & {"PESSOA_EM_OUTRO_ENVIO",
                                                                     validador.JA_ENVIADO_AO_BANCO})


# ---------------- Um número só ----------------

def _arquivo_curto(pessoas: int, com_cpf_errado: bool, colunas_extras: str = "") -> bytes:
    """Um CSV curto (Nome;CPF e, se pedido, mais colunas), sem as outras informações obrigatórias do layout."""
    cabecalho = "Nome;CPF" + colunas_extras
    linhas = [cabecalho]
    for indice in range(pessoas):
        nome, nove_digitos, _ = PESSOAS_NOVAS[indice]
        cpf = cpf_valido(nove_digitos)
        if com_cpf_errado and indice == 0:
            cpf = cpf[:-1] + str((int(cpf[-1]) + 1) % 10)
        extra = ""
        if colunas_extras:
            extra = ";Casdo"
        linhas.append(f"{nome};{cpf}{extra}")
    return ("\n".join(linhas) + "\n").encode("utf-8")


@pytest.mark.parametrize("pessoas, com_cpf_errado, colunas_extras, escolhas", [
    (3, False, "", {"Nome": "nome_completo", "CPF": "cpf"}),
    (5, True, "", {"Nome": "nome_completo", "CPF": "cpf"}),
    (4, True, ";Estado civil", {"Nome": "nome_completo", "CPF": "cpf", "Estado civil": "estado_civil"}),
])
def test_um_numero_so_nas_tres_telas_contando_o_que_falta_no_arquivo(conexao, pessoas, com_cpf_errado, colunas_extras,
                                                                     escolhas):
    leitura = cadastro.enviar_arquivo(conexao, EMPRESA, LOGIN, _arquivo_curto(pessoas, com_cpf_errado, colunas_extras),
                                      "curto.csv", busca=busca_falsa)
    processamento_id = leitura["processamento_id"]
    leitura = cadastro.aceitar_mapeamento(conexao, EMPRESA, LOGIN, processamento_id, escolhas, busca=busca_falsa)
    resumo_do_painel = leitura["resumo_das_pendencias"]
    contagem = cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)["contagem"]
    cartoes = []
    for pendencia in acompanhamento.pendencias_da_empresa(conexao, EMPRESA):
        if pendencia["processamento_id"] == processamento_id:
            cartoes.append(pendencia)
    # O mesmo número nas três telas, e ele conta o que falta no arquivo inteiro
    assert resumo_do_painel == contagem
    assert contagem["para_revisar"] == len(cartoes)
    assert contagem["no_arquivo"] > 0
    assert contagem["para_revisar"] == contagem["corrigir"] + contagem["confirmar"] + contagem["no_arquivo"]
    # Falta informação para todos: ninguém está pronto ainda
    assert contagem["prontas"] == 0 and contagem["linhas"] == pessoas
    if com_cpf_errado:
        assert contagem["linhas_com_pendencia"] >= 1


def _achado(regra, severidade, linha, campo=None, resolvido=None):
    """Um achado do Validador para os testes do resumo."""
    return validador.Achado(regra, severidade, "mensagem", "ação", linha=linha, campo=campo, resolvido=resolvido)


@pytest.mark.parametrize("achados, esperado", [
    # Só o arquivo inteiro: 2 informações faltam; ninguém pronto
    ([_achado("OBRIGATORIO_SEM_COLUNA", "BLOQUEANTE", None, "cnpj_empregador"),
      _achado("OBRIGATORIO_SEM_COLUNA", "BLOQUEANTE", None, "data_admissao")],
     {"para_revisar": 2, "no_arquivo": 2, "prontas": 0, "linhas_com_pendencia": 0}),
    # Uma correção e a pergunta da IA sobre o mesmo campo: um cartão só; o alerta confirmado não conta
    ([_achado("OBRIGATORIO_VAZIO", "BLOQUEANTE", 2, "cpf"),
      _achado(validador.PREFIXO_DA_PERGUNTA_DA_IA + "cpf", "ALERTA", 2, "cpf"),
      _achado("RENDA_FORA_DO_CARGO", "ALERTA", 3, "valor_renda", resolvido="CONFIRMADO")],
     {"para_revisar": 1, "corrigir": 1, "confirmar": 0, "perguntas_da_ia": 1, "prontas": 2}),
    # Quem fica de fora não conta como pessoa nem como pendência; o aviso comum não conta
    ([_achado(validador.JA_ENVIADO_AO_BANCO, "AVISO", 4, "cpf"), _achado("ALGUM_AVISO", "AVISO", 2),
      _achado("RENDA_FORA_DO_CARGO", "ALERTA", 3, "valor_renda")],
     {"linhas": 2, "ficam_de_fora": 1, "para_revisar": 1, "confirmar": 1, "prontas": 1}),
])
def test_resumo_das_pendencias(achados, esperado):
    registros = [{"_linha": 2}, {"_linha": 3}, {"_linha": 4}]
    resumo = cadastro.resumo_das_pendencias(validador.RelatorioValidacao(achados=achados), registros)
    for chave, valor in esperado.items():
        assert resumo[chave] == valor, chave


def test_o_painel_nao_tem_numero_antes_da_conferencia(conexao):
    leitura = cadastro.enviar_arquivo(conexao, EMPRESA, LOGIN, _arquivo_curto(2, False), "curto.csv",
                                      busca=busca_falsa)
    assert leitura["etapa"] == "aprovar_mapeamento" and leitura["resumo_das_pendencias"] is None


# ---------------- A lista mostra o que vai para o banco ----------------

def _envio_curto(conexao) -> str:
    """Um CSV com Nome e CPF, aceito: todo o resto falta no arquivo. Devolve o envio."""
    leitura = cadastro.enviar_arquivo(conexao, EMPRESA, LOGIN, _arquivo_curto(3, False), "curto.csv",
                                      busca=busca_falsa)
    cadastro.aceitar_mapeamento(conexao, EMPRESA, LOGIN, leitura["processamento_id"],
                                {"Nome": "nome_completo", "CPF": "cpf"}, busca=busca_falsa)
    return leitura["processamento_id"]


def _valores_da_lista(conexao, processamento_id: str, campo: str) -> list[str]:
    """O valor de um campo em cada linha da lista para conferir (o campo precisa estar na lista)."""
    lista = cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)
    campos = []
    for item in lista["campos"]:
        campos.append(item["campo"])
    assert campo in campos
    valores = []
    for linha in lista["linhas"]:
        valores.append(linha["valores"][campo])
    return valores


def test_valor_informado_para_todos_aparece_na_lista(conexao):
    processamento_id = _envio_curto(conexao)
    cnpj = empresas.cnpjs_conhecidos(conexao, EMPRESA)["principal"]
    correcoes.preencher_para_todos(conexao, processamento_id, EMPRESA, "cnpj_empregador", cnpj, "Informado no cartão",
                                   LOGIN)
    assert _valores_da_lista(conexao, processamento_id, "cnpj_empregador") == [cnpj, cnpj, cnpj]


def test_valor_informado_pessoa_a_pessoa_aparece_na_lista(conexao):
    processamento_id = _envio_curto(conexao)
    linhas = []
    for linha in cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)["linhas"]:
        linhas.append(linha["linha"])
    datas = {linhas[0]: "03/02/2025", linhas[1]: "15/07/2024", linhas[2]: "01/09/2026"}
    correcoes.informar_por_pessoa(conexao, processamento_id, EMPRESA, "data_admissao", datas, "Informado no cartão",
                                  LOGIN)
    assert _valores_da_lista(conexao, processamento_id, "data_admissao") == ["2025-02-03", "2024-07-15", "2026-09-01"]


def test_campo_informado_em_parte_das_pessoas_e_o_da_coluna_continuam_na_lista(conexao):
    processamento_id = _envio_curto(conexao)
    primeira_linha = cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)["linhas"][0]["linha"]
    correcoes.informar_por_pessoa(conexao, processamento_id, EMPRESA, "data_admissao", {primeira_linha: "10/10/2023"},
                                  "Informado no cartão", LOGIN)
    # A pessoa sem o valor fica vazia (a tela mostra "Informação não encontrada"); a coluna do arquivo continua
    assert _valores_da_lista(conexao, processamento_id, "data_admissao") == ["2023-10-10", "", ""]
    assert len(_valores_da_lista(conexao, processamento_id, "nome_completo")) == 3


def test_campo_sem_coluna_e_sem_valor_nao_entra_na_lista(conexao):
    processamento_id = _envio_curto(conexao)
    campos = set()
    for item in cadastro.lista_para_conferir(conexao, EMPRESA, processamento_id)["campos"]:
        campos.add(item["campo"])
    # Só as duas colunas do arquivo: nenhum campo aparece sem valor
    assert campos == {"nome_completo", "cpf"}
    assert "cnpj_empregador" in parametros.campos_iguais_para_todos(conexao)
